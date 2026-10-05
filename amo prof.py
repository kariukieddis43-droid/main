import os
import random
import json
import re
from collections import Counter
from datetime import datetime
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from urllib.parse import quote_plus
from xml.etree.ElementTree import Element, SubElement, tostring

from flask import Flask, Response, render_template_string, request, redirect, url_for, send_from_directory, flash, session
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import inspect
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

app = Flask(__name__)
secret_key = os.environ.get('SECRET_KEY')
if os.environ.get('APP_ENV') == 'production' and not secret_key:
    raise RuntimeError('SECRET_KEY must be configured in production.')
app.config['SECRET_KEY'] = secret_key or 'dev-secret-key-change-this'
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.abspath(os.path.dirname(__file__)), 'uploads')
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///prof_amon.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024
app.config['PUBLIC_BASE_URL'] = os.environ.get('PUBLIC_BASE_URL', '').rstrip('/')
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SECURE'] = os.environ.get('SESSION_COOKIE_SECURE', '').lower() in {'1', 'true', 'yes'}
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.wsgi_app = ProxyFix(app.wsgi_app, x_proto=1, x_host=1)

ALLOWED_EXTENSIONS = {'pdf', 'txt', 'doc', 'docx'}

db = SQLAlchemy(app)


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    full_name = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(150), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    major = db.Column(db.String(120), default='Undecided')
    profile_title = db.Column(db.String(120), default='Student')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    results = db.relationship('QuizResult', backref='student', lazy=True)
    messages = db.relationship('ChatMessage', backref='student', lazy=True)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)


class QuizResult(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    unit_code = db.Column(db.String(50), nullable=False)
    score = db.Column(db.Integer, nullable=False)
    total = db.Column(db.Integer, nullable=False)
    percent = db.Column(db.Float, nullable=False)
    grade = db.Column(db.String(20), nullable=False)
    submitted_at = db.Column(db.DateTime, default=datetime.utcnow)


class ChatMessage(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    role = db.Column(db.String(20), nullable=False)
    content = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class AssessmentProgress(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    unit_code = db.Column(db.String(80), nullable=False)
    assessment_type = db.Column(db.String(20), nullable=False)
    used_question_ids = db.Column(db.Text, nullable=False, default='[]')
    attempt_count = db.Column(db.Integer, nullable=False, default=0)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    __table_args__ = (db.UniqueConstraint('user_id', 'unit_code', 'assessment_type'),)

UNIT_LIBRARY = {
    'SMA2104': {
        'code': 'SMA2104',
        'title': 'Mathematics for Science',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': 'Calculus, vectors, algebra, and problem-solving methods used in science, engineering, and quantitative disciplines.',
        'objectives': [
            'Understand limits and continuity.',
            'Differentiate functions and interpret rates of change.',
            'Apply vectors, matrices, and algebra to real problems.'
        ],
        'topics': [
            {'title': 'Functions and Limits', 'summary': 'A function models a relationship between variables; limits describe behavior near a point.', 'key_terms': ['limit', 'continuity', 'function', 'domain'], 'notes': ['A limit tells us what value a function approaches as the input approaches a number.', 'Continuity means there is no break, jump, or hole in the graph at the point.', 'Calculus begins with limits because derivatives are based on rates of change near a point.']},
            {'title': 'Differentiation', 'summary': 'Differentiation measures the instantaneous rate of change of a function.', 'key_terms': ['derivative', 'slope', 'tangent', 'rate of change'], 'notes': ['The derivative gives the slope of the tangent line to a curve at a point.', 'It is used to study motion, optimization, and change in physical systems.', 'The second derivative tells us whether a function is accelerating or decelerating.']},
            {'title': 'Vectors and Matrices', 'summary': 'Vectors show direction and magnitude; matrices help solve systems and transformations.', 'key_terms': ['vector', 'matrix', 'dot product', 'determinant'], 'notes': ['A vector has both magnitude and direction.', 'Matrices organize large systems of equations and linear transformations.', 'The determinant helps determine whether a matrix is invertible.']}
        ],
        'practice_questions': [
            'Evaluate the limit of (x^2 - 9)/(x - 3) as x approaches 3.',
            'Find the derivative of f(x) = 4x^3 - 5x + 9 and interpret it.',
            'Find the dot product of a = (2, 4) and b = (1, -3).'
        ],
        'cat_questions': [
            'Explain continuity in your own words and why it matters in calculus.',
            'A particle has displacement s(t) = t^2 + 3t. Find v(t) and explain its meaning.',
            'How can matrices be used to solve a system of equations?'
        ],
        'exam_questions': [
            'Determine the maximum and minimum values of f(x) = x^3 - 6x^2 + 9x + 1 on [0, 4].',
            'A particle moves along a line with s(t) = 2t^3 - 9t^2 + 12t. Find the times of rest and acceleration.',
            'Solve the system 2x + y = 7 and x - 3y = -4 using matrix methods.'
        ],
        'video_search': 'calculus derivatives vectors matrices university lecture',
        'quiz_questions': [
            {'question': 'Which expression best describes a derivative?', 'choices': ['A fixed value', 'The slope of a tangent line', 'The determinant of a matrix', 'A variable name'], 'answer': 'The slope of a tangent line'},
            {'question': 'A function is continuous at a point when:', 'choices': ['It has no jump, break, or hole there', 'It is always linear', 'It only has one input', 'It has zero variables'], 'answer': 'It has no jump, break, or hole there'},
            {'question': 'The dot product of (2,4) and (1,-3) is:', 'choices': ['-10', '10', '2', '-2'], 'answer': '-10'},
            {'question': 'Matrices are most useful for:', 'choices': ['Solving systems of equations', 'Counting letters in a list', 'Creating graphs only', 'Replacing differentiation'], 'answer': 'Solving systems of equations'},
            {'question': 'The second derivative helps us study:', 'choices': ['The domain only', 'Acceleration or curvature', 'The slope of a line only', 'The y-intercept only'], 'answer': 'Acceleration or curvature'}
        ]
    },
    'ICS1101': {
        'code': 'ICS1101',
        'title': 'Introduction to Computer Science and Programming',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': 'Programming fundamentals, logic, algorithms, and the organization of software systems for beginners.',
        'objectives': [
            'Write programs using variables, conditions, and loops.',
            'Design algorithms to solve problems clearly and efficiently.',
            'Understand data structures and software organization.'
        ],
        'topics': [
            {'title': 'Programming Fundamentals', 'summary': 'Programs are built from variables, conditions, loops, and functions.', 'key_terms': ['variable', 'loop', 'condition', 'function'], 'notes': ['Variables store data used during program execution.', 'Conditions allow programs to choose different actions based on inputs.', 'Loops enable repeated work without writing duplicate code.']},
            {'title': 'Algorithms and Problem Solving', 'summary': 'An algorithm is a step-by-step method for solving a problem.', 'key_terms': ['algorithm', 'pseudocode', 'efficiency', 'sequence'], 'notes': ['Algorithms require a clear sequence of instructions.', 'Pseudocode expresses logic before actual programming starts.', 'A good algorithm balances correctness, simplicity, and speed.']},
            {'title': 'Data Structures', 'summary': 'Data structures organize values so they can be stored, accessed, and manipulated efficiently.', 'key_terms': ['array', 'list', 'dictionary', 'abstraction'], 'notes': ['Lists store ordered data values.', 'Dictionaries map keys to values for quick lookup.', 'Abstraction hides internal complexity and makes code easier to use.']}
        ],
        'practice_questions': [
            'Write a program to check whether a number is odd or even.',
            'Explain the difference between a loop and a conditional statement.',
            'Design an algorithm to find the largest number in a list.'
        ],
        'cat_questions': [
            'What is the difference between a variable and a constant?',
            'Write pseudocode to add the first 10 natural numbers.',
            'Why is abstraction helpful when designing software?'
        ],
        'exam_questions': [
            'Write a Python program that calculates the average of student marks.',
            'Use a loop to print all even numbers from 1 to 100.',
            'Explain the difference between a list and a dictionary and give one real use case for each.'
        ],
        'video_search': 'python programming basics algorithms beginner tutorial',
        'quiz_questions': [
            {'question': 'A variable is best described as:', 'choices': ['A database row', 'A named storage location for data', 'A type of loop', 'A comment in code'], 'answer': 'A named storage location for data'},
            {'question': 'A loop is used to:', 'choices': ['Repeat actions until a condition changes', 'Store a single value', 'Remove all variables', 'Define a function'], 'answer': 'Repeat actions until a condition changes'},
            {'question': 'Which is the best description of an algorithm?', 'choices': ['A random number generator', 'A step-by-step method to solve a task', 'A type of browser', 'A comment in code'], 'answer': 'A step-by-step method to solve a task'},
            {'question': 'A dictionary is useful because it lets you:', 'choices': ['Store data as key-value pairs', 'Only hold numbers', 'Eliminate all loops', 'Replace functions'], 'answer': 'Store data as key-value pairs'},
            {'question': 'Abstraction is useful because it:', 'choices': ['Removes all complexity', 'Hides complexity and simplifies use', 'Runs code faster automatically', 'Creates variables by itself'], 'answer': 'Hides complexity and simplifies use'}
        ]
    },
    'SMA1101': {
        'code': 'SMA1101',
        'title': 'Basic Mathematics and Algebra',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': 'Foundational algebra, equations, functions, and mathematical reasoning used across other scientific disciplines.',
        'objectives': [
            'Solve equations and simplify expressions.',
            'Interpret graphs and algebraic models.',
            'Apply algebraic techniques to real-life and academic problem-solving.'
        ],
        'topics': [
            {'title': 'Algebraic Expressions', 'summary': 'Symbols and numbers are combined to form expressions and equations.', 'key_terms': ['expression', 'equation', 'variable', 'coefficient'], 'notes': ['An expression does not contain an equal sign.', 'An equation states that two expressions are equal.', 'Algebraic symbols let us model unknown values and relationships.']},
            {'title': 'Quadratic Equations', 'summary': 'Quadratic equations involve squared terms and appear in many areas of mathematics and science.', 'key_terms': ['quadratic', 'factorization', 'root', 'discriminant'], 'notes': ['A quadratic is of the form ax^2 + bx + c = 0.', 'Solutions can be found by factoring, completing the square, or using the quadratic formula.', 'The discriminant reveals whether roots are real and distinct, repeated, or imaginary.']},
            {'title': 'Functions and Graphing', 'summary': 'Functions show how one quantity depends on another and can be represented graphically.', 'key_terms': ['function', 'graph', 'intercept', 'slope'], 'notes': ['A function gives exactly one output for each input value.', 'Graphs help us visualize patterns and relationships.', 'The slope tells us how steep a line is and whether it increases or decreases.']}
        ],
        'practice_questions': [
            'Solve 3x + 7 = 22.',
            'Factor x^2 - 7x + 12.',
            'Find the gradient and y-intercept of y = 4x - 5.'
        ],
        'cat_questions': [
            'Explain the difference between an expression and an equation.',
            'Solve x^2 - 5x + 6 = 0 using factorisation.',
            'What does the slope of a graph represent?'
        ],
        'exam_questions': [
            'Solve the simultaneous equations 2x + y = 10 and x - y = 1.',
            'Find the roots of x^2 - 9x + 20 = 0 and sketch the parabola.',
            'A line passes through (2, 5) and (5, 11). Find its equation.'
        ],
        'video_search': 'basic algebra functions graphing quadratic equations tutorial',
        'quiz_questions': [
            {'question': 'An equation is different from an expression because it:', 'choices': ['Contains an equals sign', 'Has no variables', 'Always has two variables', 'Is never solved'], 'answer': 'Contains an equals sign'},
            {'question': 'The slope of a line shows:', 'choices': ['Its steepness and direction', 'Its exact value of x only', 'Only whether it is a parabola', 'The number of points on the graph'], 'answer': 'Its steepness and direction'},
            {'question': 'The discriminant tells us:', 'choices': ['Whether the roots are real or complex', 'The value of x only', 'The y-intercept', 'Only the sum of coefficients'], 'answer': 'Whether the roots are real or complex'},
            {'question': 'A function gives:', 'choices': ['Only one output for each input', 'A random output each time', 'Only text output', 'No relation between variables'], 'answer': 'Only one output for each input'},
            {'question': 'Factoring is useful because it helps us:', 'choices': ['Find roots and solve equations', 'Avoid all algebra', 'Remove variables permanently', 'Turn graphs into ratios'], 'answer': 'Find roots and solve equations'}
        ]
    },
    'PHY1201': {
        'code': 'PHY1201',
        'title': 'Physics Fundamentals',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': 'Mechanics, motion, force, work, energy, and the fundamentals of circuits and electricity.',
        'objectives': [
            'Explain motion using physical laws.',
            'Apply force, work, and energy concepts.',
            'Interpret simple electrical systems and circuits.'
        ],
        'topics': [
            {'title': 'Motion and Force', 'summary': 'The motion of objects is described by velocity, acceleration, and force.', 'key_terms': ['force', 'mass', 'velocity', 'acceleration'], 'notes': ['Velocity describes speed with direction.', 'Acceleration is the rate of change of velocity.', 'Newton’s second law shows how force changes the motion of an object.']},
            {'title': 'Energy and Work', 'summary': 'Energy is transferred by work and remains conserved in a closed system.', 'key_terms': ['work', 'energy', 'power', 'conservation'], 'notes': ['Work is done when a force causes displacement.', 'Kinetic energy is associated with motion while potential energy is stored energy.', 'Power measures how quickly energy is transferred.']},
            {'title': 'Electricity and Circuits', 'summary': 'Voltage, current, and resistance govern how circuits work.', 'key_terms': ['current', 'voltage', 'resistance', 'circuit'], 'notes': ['Voltage pushes charge through a circuit.', 'Current is the rate of flow of charge.', 'Ohm’s law relates voltage, current, and resistance.']}
        ],
        'practice_questions': [
            'A 5 kg object experiences a net force of 20 N. Find its acceleration.',
            'Explain the difference between kinetic and potential energy.',
            'Use Ohm’s law to find current if V = 12V and R = 4Ω.'
        ],
        'cat_questions': [
            'State Newton’s laws of motion in simple terms.',
            'How is work related to force and displacement?',
            'Explain why resistance affects current in a circuit.'
        ],
        'exam_questions': [
            'A ball is thrown vertically upward with speed 20 m/s. Find the maximum height reached.',
            'A circuit has a 12V source and a 3Ω resistor. Find the current and power.',
            'Compare elastic and inelastic collisions and discuss energy conservation.'
        ],
        'video_search': 'physics mechanics work energy electricity circuits tutorial',
        'quiz_questions': [
            {'question': 'Newton’s second law states force equals:', 'choices': ['Mass times acceleration', 'Velocity times time', 'Power times energy', 'Current times resistance'], 'answer': 'Mass times acceleration'},
            {'question': 'Which quantity measures how quickly work is done?', 'choices': ['Energy', 'Power', 'Mass', 'Velocity'], 'answer': 'Power'},
            {'question': 'Ohm’s law relates:', 'choices': ['Voltage, current, and resistance', 'Force, mass, and speed', 'Energy, work, and power', 'Distance and time'], 'answer': 'Voltage, current, and resistance'},
            {'question': 'Kinetic energy depends on:', 'choices': ['Mass and speed', 'Voltage only', 'Time only', 'Color of the object'], 'answer': 'Mass and speed'},
            {'question': 'A higher resistance in a circuit usually causes:', 'choices': ['More current', 'Less current at the same voltage', 'No effect on current', 'Higher mass'], 'answer': 'Less current at the same voltage'}
        ]
    },
    'ACC2201': {
        'code': 'ACC2201',
        'title': 'Financial Accounting',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': 'Accounting principles, journal entries, financial statements, and the interpretation of business information.',
        'objectives': [
            'Understand the purpose of accounting records.',
            'Prepare and interpret financial statements.',
            'Apply double-entry logic and accounting adjustments.'
        ],
        'topics': [
            {'title': 'Accounting Basics', 'summary': 'Accounting records business transactions so the financial position can be seen clearly.', 'key_terms': ['asset', 'liability', 'equity', 'journal'], 'notes': ['Assets are resources controlled by a business that have value.', 'Liabilities are obligations owed to outside parties.', 'The accounting equation is Assets = Liabilities + Equity.']},
            {'title': 'Financial Statements', 'summary': 'Financial statements show a business’s performance and financial health.', 'key_terms': ['income statement', 'balance sheet', 'cash flow'], 'notes': ['The income statement shows revenue and expenses over a time period.', 'The balance sheet displays what the business owns and owes at a specific date.', 'The cash flow statement explains why cash changed over time.']},
            {'title': 'Journal Entries and Adjustments', 'summary': 'Transactions must be recorded carefully and adjusted for accuracy at period end.', 'key_terms': ['debit', 'credit', 'adjustment', 'closing entry'], 'notes': ['Every transaction affects at least two accounts under double-entry bookkeeping.', 'Adjustments ensure revenues and expenses are matched to the correct period.', 'Closing entries transfer temporary account balances to retained earnings or equity.']}
        ],
        'practice_questions': [
            'Record a sale on credit in the journal.',
            'Differentiate between assets and liabilities using examples.',
            'Explain why adjusting entries are necessary before final accounts are prepared.'
        ],
        'cat_questions': [
            'What is the accounting equation?',
            'Explain how an income statement differs from a balance sheet.',
            'Why are debits and credits important in double-entry accounting?'
        ],
        'exam_questions': [
            'Prepare a simple trial balance from a set of transactions.',
            'Prepare a statement of profit or loss using given income and expense figures.',
            'Analyse how one transaction affects the accounting equation.'
        ],
        'video_search': 'financial accounting basics journal entries balance sheet tutorial',
        'quiz_questions': [
            {'question': 'The accounting equation is:', 'choices': ['Assets = Liabilities + Equity', 'Revenue = Cost + Profit', 'Cash = Sales - Expenses', 'Debits = Credits + Profit'], 'answer': 'Assets = Liabilities + Equity'},
            {'question': 'A debit usually:', 'choices': ['Increases assets', 'Increases liabilities', 'Decreases revenue', 'Removes equity'], 'answer': 'Increases assets'},
            {'question': 'The income statement is mainly used to show:', 'choices': ['Performance over a period', 'The company’s assets only', 'Owner capital only', 'Future projections'], 'answer': 'Performance over a period'},
            {'question': 'A liability is:', 'choices': ['An obligation owed to others', 'A business asset', 'A purchase only', 'Current cash'], 'answer': 'An obligation owed to others'},
            {'question': 'Adjusting entries are important because they:', 'choices': ['Ensure correct end-of-period reporting', 'Eliminate all liabilities', 'Remove all inventory', 'Increase debt automatically'], 'answer': 'Ensure correct end-of-period reporting'}
        ]
    },
    'SMA2160': {
        'code': 'SMA2160',
        'title': 'Mathematics: Indices, Logarithms and Quadratics',
        'university': 'Jomo Kenyatta University of Agriculture and Technology (JKUAT)',
        'level': 'Undergraduate',
        'description': 'Algebra topics associated with JKUAT SMA2160, including indices, logarithms, quadratic equations, polynomial methods, counting, and series.',
        'objectives': [
            'Apply index and logarithm laws to simplify expressions and solve equations.',
            'Solve and interpret quadratic equations and polynomial equations.',
            'Use counting principles and sequence and series methods in problems.'
        ],
        'topics': [
            {'title': 'Indices and Logarithms', 'summary': 'Index laws simplify repeated multiplication, while logarithms express exponents and help solve exponential equations.', 'key_terms': ['index laws', 'exponent', 'logarithm', 'change of base'], 'notes': ['For nonzero a, a^m × a^n = a^(m+n) and a^m / a^n = a^(m-n).', 'A logarithm answers the question: to what power must the base be raised?', 'Use logarithm laws to turn products into sums and powers into coefficients.', 'Check logarithm domains: the argument must be positive and the base positive but not equal to one.']},
            {'title': 'Quadratics and Polynomial Theorems', 'summary': 'Quadratic equations can be solved by factorisation, completing the square, or the quadratic formula; polynomial theorems help test and factor roots.', 'key_terms': ['quadratic equation', 'discriminant', 'factor theorem', 'remainder theorem'], 'notes': ['For ax^2 + bx + c = 0, the discriminant is b^2 - 4ac.', 'A positive discriminant gives two distinct real roots, zero gives a repeated real root, and a negative value gives complex roots.', 'The factor theorem says (x - a) is a factor of f(x) exactly when f(a) = 0.', 'The remainder theorem says the remainder after division by (x - a) is f(a).']},
            {'title': 'Counting, Sequences and Series', 'summary': 'Permutations count ordered arrangements; combinations count selections where order does not matter. Series add terms from a sequence.', 'key_terms': ['permutation', 'combination', 'arithmetic series', 'geometric series'], 'notes': ['Use permutations when changing the order creates a different outcome.', 'Use combinations when only the selected group matters.', 'An arithmetic sequence has a constant difference between consecutive terms.', 'A geometric sequence has a constant ratio between consecutive terms.']}
        ],
        'practice_questions': [
            'Simplify (a^5 × a^2) / a^3 and state the index law used.',
            'Solve log_2(x) + log_2(x - 2) = 3, stating the domain restrictions.',
            'Solve 2x^2 - 5x - 3 = 0 and classify the roots using the discriminant.',
            'Use the factor theorem to test whether (x - 2) is a factor of x^3 - 3x^2 + 4.',
            'Find the sum of the first 10 terms of the arithmetic sequence 3, 7, 11, ... .'
        ],
        'cat_questions': [
            'State and apply three laws of indices to simplify an algebraic expression.',
            'Solve a logarithmic equation and explain how you check that each solution is in the domain.',
            'Explain when to use permutations rather than combinations.'
        ],
        'exam_questions': [
            'Solve a logarithmic equation involving two logarithms, showing domain checks and rejecting extraneous roots.',
            'Solve a parameterized quadratic and describe how its discriminant determines the nature of its roots.',
            'Given a cubic polynomial, use the remainder and factor theorems to find a factorization.',
            'Solve a counting problem that requires distinguishing ordered arrangements from selections.',
            'Derive and use the sum formula for an arithmetic or geometric series in a contextual problem.'
        ],
        'video_search': 'indices logarithms quadratic equations factor theorem permutations series JKUAT SMA2160',
        'quiz_questions': [
            {'question': 'Simplify a^4 × a^3 for a nonzero value of a.', 'choices': ['a^12', 'a^7', 'a', '2a^7'], 'answer': 'a^7'},
            {'question': 'log_b(x) means:', 'choices': ['The number b multiplied by x', 'The power to which b must be raised to obtain x', 'The square root of x', 'The reciprocal of b'], 'answer': 'The power to which b must be raised to obtain x'},
            {'question': 'If a quadratic has discriminant zero, its roots are:', 'choices': ['Two distinct real roots', 'One repeated real root', 'No real roots', 'Always irrational'], 'answer': 'One repeated real root'},
            {'question': 'If f(2) = 0, the factor theorem says:', 'choices': ['(x + 2) is a factor', '(x - 2) is a factor', '2 is the only root', 'f has degree two'], 'answer': '(x - 2) is a factor'},
            {'question': 'A combination is used when:', 'choices': ['Order matters', 'Order does not matter', 'There is only one choice', 'All outcomes are repeated'], 'answer': 'Order does not matter'}
        ]
    },
    'SMA2100-IT': {
        'code': 'SMA2100-IT',
        'title': 'Discrete Mathematics',
        'university': 'Jomo Kenyatta University of Agriculture and Technology (JKUAT)',
        'level': 'Undergraduate',
        'description': 'Discrete mathematical structures and reasoning used in information technology and computer science. This is the IT programme unit; the Agribusiness programme uses SMA2100 for Mathematics for Agriculture.',
        'objectives': ['Work with logical statements and proof methods.', 'Represent relations, sets, and functions precisely.', 'Understand counting, graphs, and discrete structures.'],
        'topics': [
            {'title': 'Logic and Proof', 'summary': 'Propositional logic gives precise ways to evaluate statements and construct arguments.', 'key_terms': ['proposition', 'truth table', 'implication', 'proof'], 'notes': ['A proposition is a statement that is either true or false.', 'Truth tables evaluate compound statements for every possible truth assignment.', 'An implication p → q is false only when p is true and q is false.', 'A proof explains why a statement must be true from accepted assumptions.']},
            {'title': 'Sets, Relations and Functions', 'summary': 'Sets collect objects, relations describe connections, and functions assign outputs to inputs.', 'key_terms': ['set', 'relation', 'equivalence relation', 'function'], 'notes': ['Set operations include union, intersection, and complement.', 'A relation is a subset of a Cartesian product.', 'An equivalence relation is reflexive, symmetric, and transitive.', 'A function assigns exactly one output to each element of its domain.']},
            {'title': 'Counting and Graphs', 'summary': 'Counting principles and graph models help solve finite arrangement and network problems.', 'key_terms': ['permutation', 'combination', 'graph', 'degree'], 'notes': ['The product rule counts sequential choices.', 'Permutations account for order; combinations do not.', 'A graph consists of vertices and edges.', 'The degree of a vertex is the number of edges incident to it.']}
        ],
        'practice_questions': ['Construct a truth table for (p AND q) OR NOT p.', 'Determine whether a relation is reflexive, symmetric, and transitive.', 'How many ways can a committee of 3 be chosen from 8 people?', 'Draw a graph for a small network and find each vertex degree.'],
        'cat_questions': ['Use a truth table to test whether two propositions are logically equivalent.', 'Give an example of an equivalence relation and justify its properties.', 'Explain the difference between a permutation and a combination.'],
        'exam_questions': ['Prove a statement using direct proof and proof by contradiction.', 'Analyze a relation on a finite set and determine whether it is an equivalence relation.', 'Solve a counting problem using both the product rule and combinations.', 'Represent a network as a graph and analyze vertex degrees and paths.'],
        'video_search': 'discrete mathematics logic sets relations graph theory university',
        'quiz_questions': [
            {'question': 'A proposition is a statement that:', 'choices': ['Is always a question', 'Has a truth value', 'Must contain a variable', 'Cannot be evaluated'], 'answer': 'Has a truth value'},
            {'question': 'An implication p → q is false when:', 'choices': ['p and q are both true', 'p is true and q is false', 'p is false and q is true', 'p and q are both false'], 'answer': 'p is true and q is false'},
            {'question': 'An equivalence relation must be:', 'choices': ['Reflexive, symmetric, and transitive', 'Only symmetric', 'Only transitive', 'Antisymmetric and irreflexive'], 'answer': 'Reflexive, symmetric, and transitive'},
            {'question': 'A combination counts selections where:', 'choices': ['Order matters', 'Order does not matter', 'Every item repeats', 'There are no choices'], 'answer': 'Order does not matter'},
            {'question': 'The degree of a graph vertex is:', 'choices': ['The number of vertices', 'The number of incident edges', 'The number of graph components', 'The length of the graph'], 'answer': 'The number of incident edges'}
        ]
    },
    'DIT0106': {
        'code': 'DIT0106',
        'title': 'Basic Mathematics for IT',
        'university': 'Jomo Kenyatta University of Agriculture and Technology (JKUAT)',
        'level': 'Diploma',
        'description': 'Foundational algebra and quantitative methods for information technology students.',
        'objectives': ['Manipulate numbers, fractions, ratios, and algebraic expressions.', 'Solve linear and quadratic equations.', 'Apply basic mathematical reasoning to IT examples.'],
        'topics': [
            {'title': 'Number Systems and Algebra', 'summary': 'Number properties and algebraic notation provide a foundation for technical problem solving.', 'key_terms': ['integer', 'fraction', 'ratio', 'expression'], 'notes': ['Use order of operations consistently when simplifying expressions.', 'Ratios compare quantities and can be scaled by multiplying both terms by the same factor.', 'An expression has no equality claim, while an equation states that two quantities are equal.']},
            {'title': 'Equations and Functions', 'summary': 'Equations find unknown values; functions describe how one quantity depends on another.', 'key_terms': ['linear equation', 'quadratic', 'function', 'graph'], 'notes': ['Maintain equality by applying the same operation to both sides.', 'A linear function has a constant rate of change.', 'Quadratic equations can be solved by factorisation or the quadratic formula.']},
            {'title': 'Data and Quantitative Reasoning', 'summary': 'Basic statistics and proportional reasoning help interpret technical data.', 'key_terms': ['mean', 'percentage', 'proportion', 'data'], 'notes': ['The mean is the sum of observations divided by the number of observations.', 'A percentage is a ratio expressed per hundred.', 'Check units and scale before interpreting a numerical result.']}
        ],
        'practice_questions': ['Simplify 3(2x - 4) + 5.', 'Solve 4x - 7 = 21.', 'A file transfer moves 240 MB in 30 seconds. Find the average rate in MB/s.', 'Find the mean of 12, 15, 9, and 20.'],
        'cat_questions': ['Solve a linear equation and show how you preserve equality.', 'Convert a ratio into a percentage and explain the method.', 'Explain how a function can model data usage over time.'],
        'exam_questions': ['Solve a pair of simultaneous linear equations and verify the solution.', 'Model a simple data-transfer problem with a linear function.', 'Solve a quadratic equation and interpret the roots in context.', 'Summarize a small dataset using mean, range, and percentage change.'],
        'video_search': 'basic mathematics for information technology algebra equations statistics',
        'quiz_questions': [
            {'question': 'Solve 4x - 7 = 21.', 'choices': ['x = 5', 'x = 7', 'x = 3.5', 'x = 28'], 'answer': 'x = 7'},
            {'question': 'A percentage is a ratio expressed per:', 'choices': ['Ten', 'Hundred', 'Thousand', 'One'], 'answer': 'Hundred'},
            {'question': 'The mean of a dataset is found by:', 'choices': ['Adding values and dividing by their count', 'Subtracting the smallest value', 'Multiplying all values', 'Choosing the most frequent value'], 'answer': 'Adding values and dividing by their count'},
            {'question': 'A linear function has:', 'choices': ['A constant rate of change', 'Only one possible input', 'No variables', 'Always a squared term'], 'answer': 'A constant rate of change'},
            {'question': 'When solving an equation, equality is preserved by:', 'choices': ['Changing only one side', 'Doing the same operation to both sides', 'Removing a term at random', 'Dividing by zero'], 'answer': 'Doing the same operation to both sides'}
        ]
    }
}


def register_jkuat_unit(code, title, discipline, description, topics, practice, cat, exam, quiz):
    UNIT_LIBRARY[code] = {
        'code': code,
        'title': title,
        'university': 'Jomo Kenyatta University of Agriculture and Technology (JKUAT)',
        'level': 'Undergraduate',
        'description': description,
        'objectives': [
            f'Explain core {discipline} concepts using accurate terminology.',
            f'Apply {discipline} methods to structured academic problems.',
            'Evaluate results and communicate a justified conclusion.'
        ],
        'topics': topics,
        'practice_questions': practice,
        'cat_questions': cat,
        'exam_questions': exam,
        'video_search': f'{title} JKUAT lecture tutorial',
        'quiz_questions': quiz,
    }


register_jkuat_unit(
    'SMA2216',
    'Applied Linear Algebra',
    'linear algebra',
    'Matrix methods and vector-space ideas for representing and solving applied mathematical systems.',
    [
        {'title': 'Linear Systems and Matrix Methods', 'summary': 'A system of linear equations can be represented compactly as Ax = b and solved using elimination or matrix methods.', 'key_terms': ['augmented matrix', 'row reduction', 'pivot', 'rank'], 'notes': ['Write each equation in the same variable order before forming the augmented matrix.', 'Elementary row operations preserve the solution set: swap rows, scale a row by a nonzero value, or add a multiple of one row to another.', 'A pivot identifies a leading variable; a free variable indicates infinitely many solutions when the system is consistent.', 'An inconsistent row such as 0 = nonzero means the system has no solution.', 'Verify a computed solution by substituting it into the original equations.']},
        {'title': 'Vector Spaces and Linear Transformations', 'summary': 'Vector spaces describe sets where vector addition and scalar multiplication follow consistent rules; linear maps preserve those operations.', 'key_terms': ['span', 'linear independence', 'basis', 'kernel'], 'notes': ['A linear combination of vectors has the form c1v1 + ... + cnvn.', 'A set spans a space if every vector in that space can be written as a linear combination of the set.', 'Linear independence means only the all-zero coefficients produce the zero vector.', 'A basis is both linearly independent and spanning; its size is the dimension.', 'For a linear transformation T, the kernel contains inputs sent to zero and the image contains all outputs it can produce.']},
        {'title': 'Determinants and Eigenvalues', 'summary': 'Determinants reveal invertibility and scaling, while eigenvectors describe directions preserved by a transformation.', 'key_terms': ['determinant', 'eigenvalue', 'eigenvector', 'characteristic equation'], 'notes': ['A square matrix is invertible exactly when its determinant is nonzero.', 'For a 2 by 2 matrix, det([[a,b],[c,d]]) = ad - bc.', 'Eigenpairs satisfy Av = lambda v for a nonzero vector v.', 'Find eigenvalues by solving det(A - lambda I) = 0, then solve (A - lambda I)v = 0 for eigenvectors.', 'Check eigenvectors by multiplying A by v and comparing the result with lambda v.']}
    ],
    ['Solve a 3 by 3 linear system by row reduction and classify its solution set.', 'Test a set of vectors for linear independence and find a basis for its span.', 'Find the eigenvalues and eigenvectors of a 2 by 2 matrix.'],
    ['Explain how pivot columns and rank help classify a linear system.', 'Determine whether a given vector set is a basis and justify the result.', 'Interpret an eigenvector in terms of a transformation.'],
    ['Solve a parameterized linear system and identify parameter values yielding zero, one, or infinitely many solutions.', 'Find the kernel, image, rank, and nullity of a linear transformation and verify the rank-nullity relationship.', 'Diagonalize a suitable matrix and explain when diagonalization is not possible.'],
    [
        {'question': 'A square matrix is invertible when its determinant is:', 'choices': ['Zero', 'Nonzero', 'Always one', 'Undefined'], 'answer': 'Nonzero'},
        {'question': 'A basis for a vector space must be:', 'choices': ['Only spanning', 'Only linearly independent', 'Spanning and linearly independent', 'A square matrix'], 'answer': 'Spanning and linearly independent'},
        {'question': 'An eigenvector v satisfies:', 'choices': ['Av = v + lambda', 'Av = lambda v', 'A + v = lambda', 'A v = 0 for every v'], 'answer': 'Av = lambda v'},
        {'question': 'A row of zeros equal to a nonzero constant indicates:', 'choices': ['A unique solution', 'An inconsistent system', 'A basis', 'A zero determinant only'], 'answer': 'An inconsistent system'},
        {'question': 'A free variable in a consistent linear system generally means:', 'choices': ['No solutions', 'Infinitely many solutions', 'The matrix is identity', 'The system is nonlinear'], 'answer': 'Infinitely many solutions'}
    ]
)

register_jkuat_unit(
    'HBC2110',
    'Introduction to Business Statistics',
    'business statistics',
    'Statistical tools for summarizing business data, reasoning about uncertainty, and supporting evidence-based decisions.',
    [
        {'title': 'Data, Tables and Descriptive Statistics', 'summary': 'Descriptive statistics organize observations and summarize their centre, spread, and shape.', 'key_terms': ['population', 'sample', 'mean', 'standard deviation'], 'notes': ['Distinguish the target population from the sample actually observed.', 'The mean uses every observation but is sensitive to extreme values; the median is more resistant.', 'Range, variance, and standard deviation describe spread, with standard deviation expressed in the data units.', 'For grouped data, a class midpoint is an approximation used in summary calculations.', 'Graphs should have clear labels, units, and scales that do not distort the pattern.']},
        {'title': 'Probability and Distributions', 'summary': 'Probability models quantify uncertainty and distributions describe the possible values of a random variable.', 'key_terms': ['event', 'conditional probability', 'expected value', 'normal distribution'], 'notes': ['Probability lies between zero and one; complementary event probabilities sum to one.', 'For independent events, P(A and B) = P(A)P(B); do not assume independence without evidence.', 'Conditional probability updates the chance of A given B: P(A|B) = P(A and B)/P(B).', 'Expected value is a long-run weighted average, not a prediction that every single outcome will equal it.', 'A normal model is symmetric; standard scores indicate distance from the mean in standard-deviation units.']},
        {'title': 'Sampling, Estimation and Tests', 'summary': 'Inferential statistics use sample evidence to estimate population quantities and assess specific claims.', 'key_terms': ['standard error', 'confidence interval', 'null hypothesis', 'p-value'], 'notes': ['A standard error measures the sampling variability of an estimator, not the spread of individual observations.', 'A confidence interval gives a range produced by a procedure with a stated long-run coverage rate.', 'A p-value is the probability, assuming the null model, of data at least as extreme as the observed data.', 'A small p-value is evidence against the null, not the probability that the null hypothesis is true.', 'Statistical significance does not automatically imply practical or commercial importance.']}
    ],
    ['For a small sales dataset, calculate the mean, median, range, and standard deviation.', 'Use a conditional probability table to calculate a customer segment probability.', 'Interpret a confidence interval and p-value in a business decision context.'],
    ['Explain when the median is a better measure of centre than the mean.', 'Distinguish population standard deviation from the standard error of a sample mean.', 'State a null hypothesis and alternative hypothesis for a business question.'],
    ['Analyze a supplied sales dataset, select suitable descriptive measures, and justify your choice.', 'Construct and interpret a confidence interval for a population mean, stating assumptions.', 'Conduct and report a hypothesis test with hypotheses, test statistic, p-value, decision, and practical interpretation.'],
    [
        {'question': 'Which measure of centre is generally least affected by an extreme outlier?', 'choices': ['Mean', 'Median', 'Range', 'Variance'], 'answer': 'Median'},
        {'question': 'A p-value is calculated under the assumption that:', 'choices': ['The alternative hypothesis is true', 'The null hypothesis is true', 'The sample has no variation', 'The population mean is zero always'], 'answer': 'The null hypothesis is true'},
        {'question': 'For independent events A and B:', 'choices': ['P(A and B) = P(A) + P(B)', 'P(A and B) = P(A)P(B)', 'P(A|B) = 1', 'P(A) = P(B) always'], 'answer': 'P(A and B) = P(A)P(B)'},
        {'question': 'Standard error describes:', 'choices': ['Sampling variability of an estimator', 'The range of raw data only', 'A percentage error in every measurement', 'The sample size itself'], 'answer': 'Sampling variability of an estimator'},
        {'question': 'A statistically significant result:', 'choices': ['Must be commercially important', 'Is evidence against the null under the test model', 'Proves the alternative with certainty', 'Means there was no sampling error'], 'answer': 'Is evidence against the null under the test model'}
    ]
)

register_jkuat_unit(
    'HBC2101',
    'Introduction to Accounting I',
    'financial accounting',
    'Core accounting records, double-entry bookkeeping, adjustments, and the preparation of introductory financial statements.',
    [
        {'title': 'The Accounting Equation and Double Entry', 'summary': 'The accounting equation connects resources, obligations, and owners’ claims; double entry records both sides of each transaction.', 'key_terms': ['asset', 'liability', 'equity', 'double entry'], 'notes': ['Assets = Liabilities + Equity is the foundation of the statement of financial position.', 'A transaction changes at least two accounts while keeping the equation balanced.', 'Debits are left-side entries and credits are right-side entries; their effect depends on the account type.', 'Asset and expense increases are normally debits; liability, equity, and income increases are normally credits.', 'Separate the business entity from its owners when recording transactions.']},
        {'title': 'Journals, Ledgers and Trial Balance', 'summary': 'Transactions move from source documents to journals and ledgers, then balances are checked in a trial balance.', 'key_terms': ['journal', 'ledger', 'posting', 'trial balance'], 'notes': ['A journal records transactions chronologically with accounts and explanations.', 'Posting transfers journal entries into the relevant ledger accounts.', 'A trial balance checks whether total debit balances equal total credit balances.', 'Agreement of a trial balance cannot detect every error, such as a transaction omitted from both sides.', 'Retain an audit trail from source evidence through to the final account.']},
        {'title': 'Adjustments and Financial Statements', 'summary': 'Period-end adjustments apply accrual accounting before performance and financial position are reported.', 'key_terms': ['accrual', 'prepayment', 'depreciation', 'statement of profit or loss'], 'notes': ['Recognize income when earned and expenses when incurred, rather than only when cash changes hands.', 'A prepayment is an asset for a benefit paid for but not yet consumed.', 'Accrued expenses are recognized even when the supplier has not yet been paid.', 'Depreciation allocates a depreciable asset’s cost over its useful life; it is not a direct market valuation.', 'The statement of profit or loss reports performance over a period; the statement of financial position reports balances at a date.']}
    ],
    ['Record a credit sale and subsequent customer payment using double-entry accounts.', 'Prepare a trial balance from a list of ledger balances.', 'Adjust a prepaid expense and calculate its effect on profit and assets.'],
    ['Explain why every transaction needs a debit and a credit.', 'Give two errors that a balanced trial balance cannot reveal.', 'Distinguish an accrual from a prepayment.'],
    ['Prepare journal entries, ledger accounts, and a trial balance from a set of business transactions.', 'Post year-end accrual, prepayment, and depreciation adjustments and explain each effect.', 'Prepare a statement of profit or loss and statement of financial position from adjusted balances.'],
    [
        {'question': 'The accounting equation is:', 'choices': ['Assets = Liabilities + Equity', 'Assets = Income - Expenses', 'Cash = Capital + Sales', 'Liabilities = Assets + Expenses'], 'answer': 'Assets = Liabilities + Equity'},
        {'question': 'A trial balance primarily checks whether:', 'choices': ['Every transaction is correct', 'Total debits equal total credits', 'Cash equals profit', 'All assets are valued at market price'], 'answer': 'Total debits equal total credits'},
        {'question': 'An expense incurred but not yet paid is usually a/an:', 'choices': ['Prepayment', 'Accrual', 'Capital contribution', 'Inventory sale'], 'answer': 'Accrual'},
        {'question': 'Depreciation is best described as:', 'choices': ['A cash payment to replace an asset', 'Allocation of depreciable cost over useful life', 'A market revaluation', 'A liability repayment'], 'answer': 'Allocation of depreciable cost over useful life'},
        {'question': 'A credit sale normally increases:', 'choices': ['Accounts receivable and revenue', 'Cash and accounts payable', 'Expenses and inventory only', 'Owner drawings and cash'], 'answer': 'Accounts receivable and revenue'}
    ]
)

UNIT_LIBRARY['HPS2103'] = {
    **UNIT_LIBRARY['HBC2101'],
    'code': 'HPS2103',
    'title': 'Fundamentals of Accounting',
    'description': 'Introductory accounting unit listed by JKUAT resources as HPS 2103; covers accounting records, double entry, and statements.'
}

register_jkuat_unit(
    'HBC2112',
    'Principles of Marketing',
    'marketing',
    'Marketing fundamentals covering customer value, market research, segmentation, positioning, and coordinated marketing decisions.',
    [
        {'title': 'Markets, Needs and Customer Value', 'summary': 'Marketing identifies customer needs and creates, communicates, and delivers value in exchange.', 'key_terms': ['need', 'value proposition', 'exchange', 'market orientation'], 'notes': ['A need is a felt requirement; a want is shaped by culture and individual preference.', 'A value proposition explains why a chosen customer should prefer an offer.', 'Marketing is broader than promotion: it includes product, price, place, and communication decisions.', 'Customer value compares perceived benefits with the full costs and effort of obtaining an offer.', 'A market-oriented organization uses customer and competitor evidence to guide decisions.']},
        {'title': 'Research, Segmentation and Positioning', 'summary': 'Research reduces uncertainty; segmentation groups customers, and positioning creates a distinct place in their minds.', 'key_terms': ['market research', 'segmentation', 'target market', 'positioning'], 'notes': ['Define the decision question before gathering data.', 'Primary data is collected for the current question; secondary data already exists.', 'Useful segments are measurable, reachable, substantial, differentiable, and actionable.', 'Targeting evaluates segment fit against organizational objectives and capabilities.', 'A positioning statement connects a target audience, category, point of difference, and reason to believe.']},
        {'title': 'Marketing Mix and Performance', 'summary': 'The marketing mix coordinates controllable offer, price, channel, and promotion choices to support the position.', 'key_terms': ['product life cycle', 'pricing', 'distribution', 'promotion'], 'notes': ['A product includes the core benefit and supporting features or services.', 'Pricing considers customer value, costs, competitors, and strategic objectives.', 'Distribution decisions balance coverage, convenience, cost, and channel control.', 'Promotion should be consistent across advertising, sales promotion, public relations, and direct channels.', 'Track outcomes such as conversion, retention, contribution, and customer satisfaction against defined objectives.']}
    ],
    ['Write a positioning statement for a new student-focused mobile service.', 'Choose a research method for estimating demand and justify the choice.', 'Recommend a coherent product, price, place, and promotion mix for a defined segment.'],
    ['Distinguish a customer need from a value proposition.', 'Explain the difference between segmentation and targeting.', 'Why should promotion match the product’s intended positioning?'],
    ['Analyze a case, identify customer segments, select a target market, and justify a positioning strategy.', 'Develop a complete marketing mix for a new product with evidence-based decisions.', 'Propose marketing performance measures and explain how results should inform a revised strategy.'],
    [
        {'question': 'Market segmentation is the process of:', 'choices': ['Setting one price for all firms', 'Grouping customers with similar needs or characteristics', 'Advertising without research', 'Choosing a distribution warehouse'], 'answer': 'Grouping customers with similar needs or characteristics'},
        {'question': 'Primary research is data that is:', 'choices': ['Collected first-hand for the current question', 'Always free', 'Published in an old report', 'Only numerical'], 'answer': 'Collected first-hand for the current question'},
        {'question': 'The traditional marketing mix is commonly summarized as:', 'choices': ['Product, price, place, promotion', 'People, profit, policy, planning', 'Product, process, purpose, proof', 'Price, profit, packaging, people'], 'answer': 'Product, price, place, promotion'},
        {'question': 'A target market is:', 'choices': ['Every possible customer', 'The selected segment or segments an organization chooses to serve', 'A competitor’s sales area', 'A promotional channel'], 'answer': 'The selected segment or segments an organization chooses to serve'},
        {'question': 'Positioning aims to:', 'choices': ['Make a brand distinct in the target customer’s mind', 'Change the legal ownership of a product', 'Remove all competition', 'Set only the product cost'], 'answer': 'Make a brand distinct in the target customer’s mind'}
    ]
)

register_jkuat_unit(
    'HBC2206',
    'Business Law',
    'business law',
    'Introductory legal principles affecting business transactions, contracts, agency, and commercial responsibility in Kenya.',
    [
        {'title': 'Legal Foundations and Business Obligations', 'summary': 'Business law defines enforceable rights and duties and provides processes for resolving disputes.', 'key_terms': ['statute', 'precedent', 'jurisdiction', 'remedy'], 'notes': ['Distinguish a legal rule from an ethical preference or a business custom.', 'In Kenya, legislation and judicial decisions are important sources of law; the applicable rule depends on the issue and jurisdiction.', 'A remedy is the legal response to a proven wrong, and may include damages or another court order.', 'Document the facts, dates, communications, and governing agreement before analyzing a dispute.', 'This learning page is an introduction, not legal advice; consult current Kenyan law and a qualified professional for real cases.']},
        {'title': 'Contract Formation and Terms', 'summary': 'A contract is assessed through its formation, terms, capacity, consent, legality, and available remedies.', 'key_terms': ['offer', 'acceptance', 'consideration', 'breach'], 'notes': ['An offer must be sufficiently definite and communicated; invitations to treat are generally not offers.', 'Acceptance must correspond to the offer and be communicated according to the applicable rules.', 'Consideration describes something of value exchanged in common-law contract analysis.', 'Separate express terms from terms implied by law or necessary to make the agreement workable.', 'Analyze breach by identifying the term, the conduct, causation, loss, and any contractual limitation.']},
        {'title': 'Agency and Commercial Transactions', 'summary': 'Agency allows an agent to act for a principal, while commercial laws regulate transactions and remedies.', 'key_terms': ['principal', 'agent', 'authority', 'consumer protection'], 'notes': ['Actual authority may be express or implied from the role assigned.', 'Apparent authority concerns what a principal’s representations lead a third party reasonably to believe.', 'An agent should act within authority, follow lawful instructions, and disclose relevant conflicts.', 'Commercial transactions may involve statutory obligations beyond the written contract.', 'Always check the current statute and facts; classroom hypotheticals cannot replace professional advice.']}
    ],
    ['Identify the possible offer, acceptance, and consideration in a short business scenario.', 'Distinguish actual authority from apparent authority in an agency example.', 'List the facts and documents needed to analyze a commercial contract dispute.'],
    ['Explain the difference between an offer and an invitation to treat.', 'How can an agent bind a principal in a transaction?', 'Why must a business check current statutory requirements?'],
    ['Analyze a contract scenario by addressing formation, terms, breach, loss, and possible remedies.', 'Evaluate an agency dispute involving actual and apparent authority.', 'Apply a relevant Kenyan commercial-law principle to a case and state where legal advice is required.'],
    [
        {'question': 'Contract acceptance generally must:', 'choices': ['Change the offer terms', 'Correspond to the offer and be communicated as required', 'Occur after litigation begins', 'Be inferred from silence in every case'], 'answer': 'Correspond to the offer and be communicated as required'},
        {'question': 'An agent acts on behalf of a:', 'choices': ['Principal', 'Court only', 'Competitor', 'Witness'], 'answer': 'Principal'},
        {'question': 'Apparent authority is primarily concerned with:', 'choices': ['The principal’s representations to third parties', 'The agent’s private intentions only', 'A court’s budget', 'The agent’s personal property'], 'answer': 'The principal’s representations to third parties'},
        {'question': 'A remedy is:', 'choices': ['A legal response to a proven wrong', 'A marketing slogan', 'A type of company asset', 'An informal opinion'], 'answer': 'A legal response to a proven wrong'},
        {'question': 'For a real legal dispute, a student should:', 'choices': ['Rely only on a generic study note', 'Check current law and seek qualified legal advice', 'Ignore documents', 'Assume every case has the same result'], 'answer': 'Check current law and seek qualified legal advice'}
    ]
)

register_jkuat_unit(
    'HBC2303',
    'Strategic Management',
    'strategic management',
    'Frameworks for analyzing an organization’s environment and capabilities, choosing a direction, implementing plans, and evaluating results.',
    [
        {'title': 'Purpose, Governance and Strategy', 'summary': 'Strategy links organizational purpose and long-term choices to the capabilities and outcomes needed to compete.', 'key_terms': ['mission', 'vision', 'stakeholder', 'competitive advantage'], 'notes': ['A mission states what an organization does and for whom; a vision describes its desired future.', 'Stakeholders can influence or be affected by strategy and may have competing interests.', 'A strategic objective should be specific enough to guide resource allocation and measurement.', 'Competitive advantage depends on delivering distinctive value that competitors cannot easily reproduce.', 'Strategy is a set of connected choices, not simply a list of goals.']},
        {'title': 'External and Internal Analysis', 'summary': 'Environmental and capability analysis identifies opportunities, threats, strengths, and weaknesses relevant to strategic choice.', 'key_terms': ['PESTEL', 'Five Forces', 'value chain', 'SWOT'], 'notes': ['PESTEL scans political, economic, social, technological, environmental, and legal macro factors.', 'Five Forces analyzes industry rivalry, entry, substitutes, buyer power, and supplier power.', 'A value-chain analysis examines how activities create value and incur cost.', 'SWOT is useful only when its points are evidence-based and linked to strategic action.', 'Separate an attractive industry from an organization’s ability to compete successfully within it.']},
        {'title': 'Strategic Choice, Implementation and Control', 'summary': 'Strategic options must fit the organization, be feasible with its resources, and be implemented and reviewed.', 'key_terms': ['strategic fit', 'implementation', 'KPI', 'scenario planning'], 'notes': ['Evaluate options for suitability, feasibility, and acceptability to key stakeholders.', 'Implementation translates strategy into responsibilities, budgets, capabilities, and timelines.', 'KPIs should measure outcomes that matter, not merely activity that is easy to count.', 'Risk registers assign owners and responses to significant uncertainties.', 'Review assumptions and measures regularly so strategy can adapt to new evidence.']}
    ],
    ['Complete a concise PESTEL scan for a Kenyan organization entering a new market.', 'Use Five Forces to explain the competitive pressure in an industry.', 'Create three strategic objectives and one measurable KPI for each.'],
    ['Distinguish mission from vision.', 'How does a value-chain analysis support strategic decisions?', 'Why is SWOT insufficient unless it leads to choices and actions?'],
    ['Analyze an organization using external and internal evidence and identify its central strategic issue.', 'Compare two strategic options using suitability, feasibility, and acceptability.', 'Prepare an implementation and control plan with owners, resources, milestones, KPIs, and risks.'],
    [
        {'question': 'PESTEL is used to scan:', 'choices': ['The macro-environment', 'Only internal staff', 'An organization’s bank account', 'A product’s packaging'], 'answer': 'The macro-environment'},
        {'question': 'Which is one of Porter’s Five Forces?', 'choices': ['Threat of substitutes', 'Employee attendance', 'Exchange-rate accounting entry', 'Brand color'], 'answer': 'Threat of substitutes'},
        {'question': 'A KPI should primarily measure:', 'choices': ['A strategically important outcome or driver', 'Any activity that is easy to count', 'Only the number of meetings', 'A competitor’s mission'], 'answer': 'A strategically important outcome or driver'},
        {'question': 'A value chain helps analyze:', 'choices': ['Activities that create value and cost', 'Only macroeconomic inflation', 'A list of shareholders', 'The legal definition of a contract'], 'answer': 'Activities that create value and cost'},
        {'question': 'Strategy implementation includes:', 'choices': ['Responsibilities, resources, and timelines', 'Only writing a vision statement', 'Ignoring risks', 'Removing performance reviews'], 'answer': 'Responsibilities, resources, and timelines'}
    ]
)

register_jkuat_unit(
    'HPS2101',
    'Principles of Procurement',
    'procurement',
    'Procurement planning, sourcing, supplier evaluation, contract administration, and ethical value-for-money decisions.',
    [
        {'title': 'Procurement Cycle and Planning', 'summary': 'The procurement cycle converts a validated need into a competitively sourced, received, and reviewed supply or service.', 'key_terms': ['needs assessment', 'specification', 'procurement plan', 'total cost'], 'notes': ['Confirm the need, budget, timing, and approval before approaching suppliers.', 'A specification should state outcomes and required performance without unfairly favoring a supplier.', 'Planning considers lead time, market capacity, risk, and the cost of the full lifecycle.', 'Aggregate demand where appropriate, while avoiding improper splitting or unnecessary over-ordering.', 'Keep decision records so each step can be reviewed and audited.']},
        {'title': 'Sourcing, Tendering and Evaluation', 'summary': 'Sourcing methods and evaluation criteria should fit the purchase and preserve fairness, transparency, and competition.', 'key_terms': ['request for quotation', 'tender', 'evaluation criteria', 'conflict of interest'], 'notes': ['Choose a sourcing method according to value, complexity, market conditions, and applicable rules.', 'Publish clear criteria before evaluation and apply them consistently to every eligible bid.', 'Separate mandatory compliance checks from scored quality and price evaluation.', 'Declare and manage conflicts of interest; protect confidential supplier information.', 'Lowest price is not automatically best value when quality, risk, lifecycle cost, or delivery differ.']},
        {'title': 'Contract and Supplier Performance', 'summary': 'Effective contract management confirms delivery, manages change, and evaluates supplier performance after award.', 'key_terms': ['purchase order', 'service-level agreement', 'delivery acceptance', 'supplier performance'], 'notes': ['A purchase order should correspond to approved requirements and agreed commercial terms.', 'Inspect and document goods or services against acceptance criteria before approval for payment.', 'Monitor key performance indicators such as quality, delivery, safety, and responsiveness.', 'Changes to scope, price, or schedule should be authorized and documented before implementation.', 'Close contracts formally, reconcile obligations, retain records, and capture lessons for future sourcing.']}
    ],
    ['Write a simple procurement plan for a defined need, including specification, budget, and schedule.', 'Design fair evaluation criteria for a purchase where quality and price both matter.', 'Propose supplier KPIs and evidence for verifying delivery.'],
    ['Why should a specification describe needs rather than favor a named supplier?', 'Explain why lowest purchase price can differ from best value.', 'Name two controls that support transparent supplier evaluation.'],
    ['Prepare an end-to-end procurement strategy for a complex requirement with risks and controls.', 'Evaluate competing bids using pre-published criteria and justify the outcome.', 'Develop a contract-management scorecard and explain how it handles underperformance and authorized change.'],
    [
        {'question': 'Evaluation criteria should be:', 'choices': ['Defined before bids are evaluated', 'Changed for each supplier', 'Kept secret from the evaluation team', 'Based only on personal preference'], 'answer': 'Defined before bids are evaluated'},
        {'question': 'Total cost of ownership considers:', 'choices': ['Purchase price only', 'Lifecycle costs and value over the period of use', 'Only supplier advertising', 'Only delivery distance'], 'answer': 'Lifecycle costs and value over the period of use'},
        {'question': 'A conflict of interest should be:', 'choices': ['Hidden', 'Declared and appropriately managed', 'Ignored if the supplier is familiar', 'Used as a scoring advantage'], 'answer': 'Declared and appropriately managed'},
        {'question': 'Goods should be accepted after:', 'choices': ['They are checked against agreed requirements', 'The supplier requests payment', 'The purchase is advertised', 'The contract is drafted'], 'answer': 'They are checked against agreed requirements'},
        {'question': 'A procurement specification should:', 'choices': ['Describe required outcomes and performance', 'Name a preferred supplier without reason', 'Omit acceptance criteria', 'Change after bids arrive'], 'answer': 'Describe required outcomes and performance'}
    ]
)

register_jkuat_unit(
    'HBC2209',
    'Organizational Behaviour',
    'organizational behaviour',
    'How individuals, groups, leadership, and organizational systems shape workplace behavior and performance.',
    [
        {'title': 'Individual Behaviour and Motivation', 'summary': 'Personality, perception, attitudes, and motivation influence how people respond to work and organizational conditions.', 'key_terms': ['perception', 'attitude', 'motivation', 'job satisfaction'], 'notes': ['Perception is how people select and interpret information; it can differ between observers.', 'Attitudes include beliefs, feelings, and behavioral intentions toward work or an organization.', 'Motivation concerns the direction, intensity, and persistence of effort.', 'Different employees may value achievement, security, recognition, autonomy, or belonging differently.', 'Avoid assuming one motivational theory explains every person or workplace situation.']},
        {'title': 'Groups, Teams and Communication', 'summary': 'Groups develop norms and roles; effective teams coordinate complementary contributions toward shared outcomes.', 'key_terms': ['group norm', 'team role', 'communication channel', 'conflict'], 'notes': ['Group norms are informal expectations that influence member behavior.', 'Clear roles reduce duplication and gaps, but overly rigid roles can block adaptation.', 'Choose communication channels based on urgency, ambiguity, sensitivity, and the need for a record.', 'Task conflict can improve decisions when respectfully managed; relationship conflict often harms collaboration.', 'Psychological safety supports speaking up, asking questions, and reporting mistakes.']},
        {'title': 'Leadership, Culture and Change', 'summary': 'Leadership and organizational culture shape how people coordinate, make decisions, and respond to change.', 'key_terms': ['leadership', 'organizational culture', 'change management', 'equity'], 'notes': ['Leadership effectiveness depends on people, task, context, and organizational conditions.', 'Culture includes shared assumptions and norms, not just written values or slogans.', 'Change plans should explain why change is needed, support affected people, and provide feedback channels.', 'Fairness perceptions influence trust, commitment, and willingness to contribute.', 'Assess outcomes and unintended effects rather than judging a change only by whether it was announced.']}
    ],
    ['Analyze how a workplace policy might affect employee motivation and job satisfaction.', 'Design team norms that improve communication and constructive disagreement.', 'Plan communication and support for a small organizational change.'],
    ['Distinguish motivation from job satisfaction.', 'How can task conflict differ from relationship conflict?', 'Why can stated organizational values differ from actual culture?'],
    ['Analyze a workplace case at individual, group, and organizational levels and recommend evidence-based interventions.', 'Evaluate a leadership approach in a specific context, including risks and employee effects.', 'Prepare a change plan with stakeholder analysis, communication, support, feedback, and outcome measures.'],
    [
        {'question': 'Psychological safety primarily supports employees to:', 'choices': ['Speak up and raise concerns without undue interpersonal fear', 'Avoid all performance feedback', 'Ignore team goals', 'Prevent any disagreement'], 'answer': 'Speak up and raise concerns without undue interpersonal fear'},
        {'question': 'Group norms are:', 'choices': ['Shared informal expectations about behavior', 'Only written employment contracts', 'Financial statements', 'Individual job titles'], 'answer': 'Shared informal expectations about behavior'},
        {'question': 'Motivation is commonly described through:', 'choices': ['Direction, intensity, and persistence of effort', 'Salary only', 'Personality alone', 'Communication channel only'], 'answer': 'Direction, intensity, and persistence of effort'},
        {'question': 'Task conflict can be useful when it is:', 'choices': ['Respectfully managed around ideas and evidence', 'Personal and hostile', 'Hidden from everyone', 'Used to stop decisions'], 'answer': 'Respectfully managed around ideas and evidence'},
        {'question': 'Organizational culture includes:', 'choices': ['Shared assumptions and norms', 'Only the company logo', 'Only published financial data', 'A list of office locations'], 'answer': 'Shared assumptions and norms'}
    ]
)


AGRIBUSINESS_TOPIC_PACKS = {
    'communication': [
        {'title': 'Communication Foundations', 'summary': 'Communication is a purposeful process of creating, sending, receiving, and interpreting messages in a specific context.', 'key_terms': ['sender', 'audience', 'purpose', 'feedback'], 'notes': ['Identify the communication purpose and audience before selecting content or a channel.', 'A message is affected by context, prior knowledge, culture, and noise.', 'Feedback checks whether the intended meaning was understood; it is not merely a reply.', 'Use plain language, accurate evidence, and an appropriate level of detail.']},
        {'title': 'Academic and Professional Writing', 'summary': 'Strong academic communication organizes a claim, supporting evidence, and a clear line of reasoning.', 'key_terms': ['thesis', 'paragraph', 'evidence', 'citation'], 'notes': ['A focused thesis gives a response a clear direction.', 'Each paragraph should develop one main point using explanation and relevant evidence.', 'Paraphrase ideas accurately and cite sources to distinguish evidence from your own analysis.', 'Revise structure and meaning before proofreading grammar and presentation.']},
        {'title': 'Presentations and Interpersonal Communication', 'summary': 'Presentations and conversations work best when verbal, visual, and listening skills fit the audience and goal.', 'key_terms': ['active listening', 'nonverbal cue', 'visual aid', 'persuasion'], 'notes': ['Active listening includes attention, clarification, and fair restatement of another person’s point.', 'Use visuals to clarify a message rather than decorate or repeat every spoken word.', 'Support persuasive claims with evidence and address reasonable counterarguments.', 'Adapt tone and detail while maintaining accuracy and respect.']}
    ],
    'social_ethics': [
        {'title': 'Development Concepts and Indicators', 'summary': 'Development is multidimensional change in wellbeing, capabilities, institutions, and opportunity.', 'key_terms': ['development', 'capability', 'poverty', 'indicator'], 'notes': ['Income is important but does not capture health, education, security, or voice on its own.', 'Indicators summarize selected conditions and should be interpreted with their definitions and limits.', 'Disaggregate data to reveal who benefits and who may be left behind.', 'Distinguish economic growth from broader and more equitable development.']},
        {'title': 'Institutions, Inequality and Participation', 'summary': 'Institutions, power, gender, environment, and participation shape how development outcomes are distributed.', 'key_terms': ['institution', 'inequality', 'participation', 'sustainability'], 'notes': ['Institutions include formal rules and informal norms that shape incentives and access.', 'Ask who makes decisions, who bears costs, and who receives benefits.', 'Meaningful participation involves influence over decisions, not just attendance.', 'Sustainable development considers effects across communities, generations, and ecosystems.']},
        {'title': 'Social Ethics and Public Decisions', 'summary': 'Ethical analysis makes values, duties, consequences, rights, and fairness explicit in development choices.', 'key_terms': ['rights', 'justice', 'consequence', 'accountability'], 'notes': ['Compare likely consequences while also respecting rights and duties.', 'Fairness may concern equal treatment, equitable support, procedure, or distribution.', 'Transparency and accountability make decision makers answerable for impacts.', 'An ethical recommendation should name affected groups, trade-offs, evidence, and safeguards.']}
    ],
    'agri_calculus': [
        {'title': 'Functions and Agricultural Models', 'summary': 'Functions model relationships such as yield, cost, water use, and output as conditions change.', 'key_terms': ['function', 'domain', 'model', 'rate'], 'notes': ['Define what each variable represents and use consistent units.', 'The domain should reflect realistic limits such as nonnegative land, time, or quantity.', 'A mathematical model simplifies reality, so state its assumptions and scope.', 'Check whether the model’s predictions are reasonable for the context.']},
        {'title': 'Limits and Differentiation', 'summary': 'Limits describe local behavior, and derivatives measure instantaneous rates of change.', 'key_terms': ['limit', 'derivative', 'marginal change', 'continuity'], 'notes': ['A derivative estimates how output changes when an input changes slightly.', 'Marginal product and marginal cost are interpreted as rates, with units.', 'Use differentiation rules carefully and show intermediate steps.', 'Continuity assumptions matter when interpreting a model across a range.']},
        {'title': 'Optimization and Integration', 'summary': 'Calculus supports optimization of farm decisions and accumulation of quantities over time or area.', 'key_terms': ['critical point', 'maximum', 'minimum', 'integral'], 'notes': ['Find candidate optima using derivative zero points and relevant endpoints.', 'Use a second derivative or sign changes to classify a candidate maximum or minimum.', 'An integral can represent accumulated output, cost, or resource use.', 'Translate the mathematical result into a practical recommendation and note constraints.']}
    ],
    'hiv_prevention': [
        {'title': 'HIV Science and Transmission', 'summary': 'Evidence-based HIV education explains the virus, transmission routes, testing, and prevention without stigma.', 'key_terms': ['HIV', 'transmission', 'viral load', 'testing'], 'notes': ['HIV is transmitted through specific body fluids; everyday contact such as sharing utensils does not transmit HIV.', 'Testing is the only way to know HIV status, and recommended timing depends on the test and possible exposure.', 'Effective antiretroviral treatment can suppress viral load and protect health.', 'Use current national and WHO guidance because clinical recommendations can change.']},
        {'title': 'Prevention and Treatment', 'summary': 'Prevention combines informed choices, testing, condoms, PrEP or PEP access, and effective treatment services.', 'key_terms': ['PrEP', 'PEP', 'condom', 'antiretroviral therapy'], 'notes': ['PrEP is preventive medicine for people who may be exposed; suitability requires a qualified provider.', 'PEP is time-sensitive after a possible exposure and should be sought urgently from a health professional.', 'Antiretroviral therapy supports health and viral suppression; adherence and follow-up matter.', 'No educational note replaces confidential clinical advice or local service guidance.']},
        {'title': 'Stigma, Ethics and Community Support', 'summary': 'Ethical HIV programmes protect dignity, confidentiality, informed consent, and access to support.', 'key_terms': ['confidentiality', 'consent', 'stigma', 'support'], 'notes': ['HIV status is private health information and should not be disclosed without lawful consent.', 'Use person-first, nonjudgmental language and challenge misinformation.', 'Prevention education should be inclusive, accurate, and linked to accessible services.', 'Support can include counseling, treatment access, peer support, and protection from discrimination.']}
    ],
    'agribusiness_intro': [
        {'title': 'Agribusiness Systems and Participants', 'summary': 'Agribusiness connects input supply, production, processing, transport, finance, markets, and consumers.', 'key_terms': ['agribusiness', 'value chain', 'stakeholder', 'transaction'], 'notes': ['Map activities from inputs through production and handling to final consumers.', 'Farmers, input dealers, processors, traders, lenders, regulators, and buyers have linked incentives.', 'Bottlenecks may arise from information gaps, logistics, quality, finance, or coordination.', 'A business decision should consider the whole chain rather than a single enterprise in isolation.']},
        {'title': 'Enterprise Models and Value Creation', 'summary': 'Agribusiness enterprises create value by solving customer problems while managing costs, risks, and resources.', 'key_terms': ['business model', 'value proposition', 'revenue', 'margin'], 'notes': ['Specify the customer, need, offer, delivery method, and revenue source.', 'Gross margin compares revenue with variable costs, but does not include every business expense.', 'Seasonality, perishability, and production uncertainty affect cash flow and capacity.', 'Validate customer demand and operational feasibility before scaling.']},
        {'title': 'Institutions, Sustainability and Decisions', 'summary': 'Agribusiness operates within policy, infrastructure, standards, climate, and community systems.', 'key_terms': ['market institution', 'standard', 'externality', 'resilience'], 'notes': ['Standards and contracts coordinate quality and reduce uncertainty when applied fairly.', 'Externalities are costs or benefits borne by people outside a transaction.', 'Sustainable practice considers soil, water, biodiversity, labor, and long-run economic viability.', 'Resilience comes from anticipating shocks, adapting operations, and maintaining options.']}
    ],
    'microeconomics': [
        {'title': 'Demand, Supply and Equilibrium', 'summary': 'Demand and supply describe buyer and seller behavior and determine market price and traded quantity.', 'key_terms': ['demand', 'supply', 'equilibrium', 'elasticity'], 'notes': ['A movement along a curve follows a price change; a shift follows another determinant.', 'Equilibrium occurs where quantity demanded equals quantity supplied.', 'Elasticity measures percentage responsiveness and helps predict how revenue changes.', 'Weather, income, substitutes, input costs, and expectations can shift agricultural markets.']},
        {'title': 'Consumer and Producer Choice', 'summary': 'Consumers allocate budgets across goods; firms choose inputs and output under constraints.', 'key_terms': ['utility', 'marginal utility', 'production function', 'marginal cost'], 'notes': ['Marginal reasoning compares the additional benefit and additional cost of one more unit.', 'A production function maps inputs into possible output under specified technology.', 'Diminishing marginal returns can arise when one input increases while others are fixed.', 'Separate accounting costs from opportunity costs when evaluating a decision.']},
        {'title': 'Market Structure and Welfare', 'summary': 'Competition, market power, and externalities influence prices, output, efficiency, and distribution.', 'key_terms': ['competition', 'market power', 'surplus', 'externality'], 'notes': ['Market structure depends on entry barriers, product differentiation, and number of buyers and sellers.', 'Consumer and producer surplus are gains relative to reservation values and costs.', 'Externalities can make private incentives diverge from social costs or benefits.', 'Policy analysis should compare intended gains, implementation cost, and possible unintended effects.']}
    ],
    'information_technology': [
        {'title': 'Computer Systems and Digital Practice', 'summary': 'IT systems combine hardware, software, networks, data, and people to perform useful work.', 'key_terms': ['hardware', 'software', 'operating system', 'network'], 'notes': ['Hardware provides physical input, processing, storage, and output components.', 'An operating system manages resources and provides services to applications.', 'Networks enable communication and sharing but introduce security and reliability concerns.', 'Choose tools according to the task, accessibility, compatibility, and data sensitivity.']},
        {'title': 'Productivity and Information Management', 'summary': 'Digital applications support document creation, spreadsheets, presentations, databases, and collaboration.', 'key_terms': ['spreadsheet', 'database', 'formula', 'data validation'], 'notes': ['Keep raw data separate from calculated values and document important assumptions.', 'Use formulas consistently and validate inputs to reduce spreadsheet errors.', 'A database organizes records with defined fields and relationships.', 'Use clear filenames, versioning, backups, and access controls for shared work.']},
        {'title': 'Cybersecurity and Responsible Use', 'summary': 'Safe IT practice protects accounts, devices, information, and people from preventable harm.', 'key_terms': ['authentication', 'phishing', 'backup', 'privacy'], 'notes': ['Use unique strong passwords and multi-factor authentication where available.', 'Verify links and requests before providing credentials or sensitive information.', 'Back up important data and test that it can be restored.', 'Collect and share only the personal data necessary for a legitimate purpose.']}
    ],
    'macroeconomics': [
        {'title': 'National Output and Income', 'summary': 'Macroeconomics studies aggregate production, income, employment, prices, and growth.', 'key_terms': ['GDP', 'real output', 'per capita', 'business cycle'], 'notes': ['GDP measures market value of final goods and services produced in a period.', 'Real measures adjust for price changes; nominal measures use current prices.', 'GDP per capita is an average and does not show distribution or all aspects of wellbeing.', 'Agriculture contributes through production, jobs, exports, food supply, and input demand.']},
        {'title': 'Inflation, Employment and Money', 'summary': 'Price levels, labor markets, and monetary conditions affect purchasing power and investment.', 'key_terms': ['inflation', 'unemployment', 'interest rate', 'exchange rate'], 'notes': ['Inflation is a sustained increase in the general price level, not a rise in every single price.', 'Interest rates affect borrowing, saving, investment, and working capital.', 'Exchange-rate movements change import costs and export competitiveness.', 'Interpret indicators with their measurement period, base, and limitations.']},
        {'title': 'Fiscal, Monetary and Trade Policy', 'summary': 'Public policy can stabilize the economy and shape agricultural production, trade, and household welfare.', 'key_terms': ['fiscal policy', 'monetary policy', 'budget', 'trade balance'], 'notes': ['Fiscal policy uses public spending and taxation; monetary policy influences money and credit conditions.', 'Policy effects vary with timing, expectations, financing, and implementation.', 'Trade balances summarize exports and imports but do not alone determine welfare.', 'Trace policy effects through prices, input costs, employment, public services, and distribution.']}
    ],
    'production_economics': [
        {'title': 'Production Functions and Input Decisions', 'summary': 'Production economics analyzes how land, labor, capital, and management combine to produce output.', 'key_terms': ['production function', 'marginal product', 'returns to scale', 'opportunity cost'], 'notes': ['A production function represents technically feasible input-output relationships.', 'Marginal product is the additional output associated with one more unit of an input, holding others fixed.', 'Distinguish diminishing marginal returns from returns to scale, which change all inputs together.', 'Include opportunity costs when comparing alternative uses of scarce resources.']},
        {'title': 'Costs, Revenue and Profit', 'summary': 'Farm decisions compare expected revenue with economic costs under uncertain yields and prices.', 'key_terms': ['fixed cost', 'variable cost', 'marginal cost', 'gross margin'], 'notes': ['Variable costs change with production activity; fixed costs do not change over the short decision period.', 'Economic profit subtracts explicit and opportunity costs from revenue.', 'Break-even analysis identifies the output or price needed to cover specified costs.', 'Sensitivity analysis reveals which assumptions most affect the decision.']},
        {'title': 'Enterprise Choice and Risk', 'summary': 'Enterprise selection allocates resources across activities while balancing returns, constraints, and risk.', 'key_terms': ['enterprise budget', 'contribution margin', 'risk', 'resource constraint'], 'notes': ['An enterprise budget lists expected returns and costs for a defined activity and time period.', 'Contribution margin helps evaluate an additional activity before shared fixed costs.', 'Compare enterprises using consistent land, labor, capital, and calendar assumptions.', 'Risk-adjusted decisions consider variability, downside exposure, liquidity, and diversification.']}
    ],
    'business_mathematics': [
        {'title': 'Percentages, Ratios and Commercial Arithmetic', 'summary': 'Business mathematics uses ratios, rates, percentages, and unit conversions to support sound decisions.', 'key_terms': ['ratio', 'percentage change', 'unit cost', 'proportion'], 'notes': ['Percentage change is change divided by the original value, multiplied by 100.', 'Label units in every calculation and convert quantities before comparing them.', 'A unit cost divides total relevant cost by the number of units produced or purchased.', 'A percentage point change is not the same as a percentage change.']},
        {'title': 'Interest, Discounts and Cash Flow', 'summary': 'Interest and discount calculations compare money received or paid at different times.', 'key_terms': ['simple interest', 'compound interest', 'discount', 'present value'], 'notes': ['Simple interest is calculated on principal; compound interest includes accumulated interest.', 'Match the rate period to the time period before calculating interest.', 'Present value discounts future cash flows using a stated rate and timing assumption.', 'Read loan and discount terms carefully, including fees and repayment schedule.']},
        {'title': 'Break-Even and Decision Models', 'summary': 'Algebraic models translate prices, costs, volume, and targets into usable commercial decisions.', 'key_terms': ['revenue', 'contribution', 'break-even', 'margin'], 'notes': ['Revenue is price multiplied by quantity under the stated assumptions.', 'Break-even quantity divides fixed cost by unit contribution when unit contribution is positive.', 'Contribution per unit equals selling price less variable cost per unit.', 'Use scenario analysis when price, yield, or cost estimates are uncertain.']}
    ],
    'accounting': [
        {'title': 'Accounting Equation and Transaction Records', 'summary': 'Accounting records transactions using a consistent system so enterprise resources and claims can be measured.', 'key_terms': ['asset', 'liability', 'equity', 'double entry'], 'notes': ['Assets = Liabilities + Equity is the core accounting equation.', 'Each transaction affects at least two accounts and preserves the equation.', 'Separate business transactions from the owner’s personal activity.', 'Retain source documents and explanations to support the audit trail.']},
        {'title': 'Ledgers, Trial Balance and Adjustments', 'summary': 'Journal entries are posted to ledgers, checked, and adjusted to report the correct period.', 'key_terms': ['journal', 'ledger', 'trial balance', 'accrual'], 'notes': ['A trial balance checks debit-credit arithmetic but does not detect every error.', 'Accrual accounting recognizes income when earned and expenses when incurred.', 'Prepayments and accruals assign costs to the period that receives the benefit.', 'Inventory and receivable records need consistent policies and supporting evidence.']},
        {'title': 'Agribusiness Financial Statements', 'summary': 'Financial statements summarize performance, financial position, and cash movements for enterprise decisions.', 'key_terms': ['profit or loss', 'balance sheet', 'cash flow', 'working capital'], 'notes': ['The statement of profit or loss reports performance over a period.', 'The statement of financial position lists assets, liabilities, and equity at a date.', 'Cash flow distinguishes operating, investing, and financing movements.', 'Seasonality can make working-capital needs different from reported profitability.']}
    ],
    'enterprise_development': [
        {'title': 'Opportunity and Customer Discovery', 'summary': 'Enterprise development starts by identifying a real customer problem and validating demand.', 'key_terms': ['customer segment', 'problem', 'value proposition', 'validation'], 'notes': ['Interview likely customers about current behavior and unresolved problems.', 'Separate evidence of demand from compliments or hypothetical interest.', 'Describe the value proposition in terms of a specific customer outcome.', 'Test assumptions with a small, low-cost experiment before investing heavily.']},
        {'title': 'Business Model and Operations', 'summary': 'A business model explains how an enterprise creates, delivers, and captures value.', 'key_terms': ['revenue model', 'channel', 'key resource', 'unit economics'], 'notes': ['Identify customers, offer, channels, partners, activities, costs, and revenue sources.', 'Unit economics estimate contribution and cost per customer, product, or transaction.', 'Operations should account for seasonality, quality, storage, transport, and working capital.', 'A business plan should be internally consistent rather than optimistic in every assumption.']},
        {'title': 'Launch, Learning and Growth', 'summary': 'A new venture needs practical milestones, cash controls, learning loops, and responsible growth.', 'key_terms': ['minimum viable offer', 'cash runway', 'milestone', 'scaling'], 'notes': ['A pilot tests critical assumptions with bounded cost and risk.', 'Cash runway depends on available cash and the rate of net cash use.', 'Track customer retention, contribution, quality, and cash as well as sales growth.', 'Scale only when the business can maintain quality and meet obligations.']}
    ],
    'plant_animal': [
        {'title': 'Plant Growth and Crop Production', 'summary': 'Crop production depends on genetics, environment, management, soil, water, and pest conditions.', 'key_terms': ['photosynthesis', 'soil fertility', 'water balance', 'integrated pest management'], 'notes': ['Plant growth converts light, water, carbon dioxide, and nutrients into biomass and yield.', 'Soil structure, pH, organic matter, and nutrients influence crop performance.', 'Water stress at sensitive growth stages can reduce yield and quality.', 'Integrated pest management combines monitoring, prevention, thresholds, and suitable controls.']},
        {'title': 'Animal Systems and Health', 'summary': 'Animal production combines nutrition, genetics, reproduction, welfare, housing, and disease prevention.', 'key_terms': ['ration', 'biosecurity', 'breeding', 'animal welfare'], 'notes': ['Balanced rations match energy, protein, minerals, water, and life-stage needs.', 'Biosecurity reduces introduction and spread of infectious disease.', 'Breeding decisions should consider health, adaptation, productivity, and diversity.', 'Good welfare includes access to feed and water, comfort, health, and appropriate behavior.']},
        {'title': 'Integrated Farm Production', 'summary': 'Integrated plant and animal systems coordinate resources, by-products, labor, and environmental stewardship.', 'key_terms': ['crop-livestock integration', 'manure management', 'farm calendar', 'resilience'], 'notes': ['Crop residues may feed livestock, while managed manure can return nutrients to fields.', 'Plan production calendars around labor, rainfall, inputs, and market windows.', 'Monitor productivity alongside soil health, water use, and animal health.', 'Local climate, regulations, and extension advice shape appropriate production practice.']}
    ],
    'intermediate_micro': [
        {'title': 'Consumer and Producer Optimization', 'summary': 'Intermediate microeconomics formalizes choice under preferences, technology, and resource constraints.', 'key_terms': ['utility maximization', 'budget constraint', 'cost minimization', 'duality'], 'notes': ['An optimum balances marginal benefits with marginal costs subject to constraints.', 'A budget constraint describes affordable bundles at given prices and income.', 'Cost minimization chooses the least-cost input mix for a target output.', 'Interpret mathematical conditions in words and check feasibility.']},
        {'title': 'Market Equilibrium and Comparative Statics', 'summary': 'Equilibrium models show how changes in fundamentals affect prices, quantities, and welfare.', 'key_terms': ['equilibrium', 'comparative statics', 'elasticity', 'incidence'], 'notes': ['Comparative statics compares equilibria before and after a specified change.', 'Elasticities measure percentage responses and are unit-free.', 'Policy incidence depends on relative elasticities, not only on who legally pays a tax.', 'State ceteris paribus assumptions when isolating a single change.']},
        {'title': 'Information, Externalities and Policy', 'summary': 'Information failures, market power, and externalities can change incentives and justify careful policy analysis.', 'key_terms': ['asymmetric information', 'externality', 'public good', 'deadweight loss'], 'notes': ['Asymmetric information occurs when transaction participants have different relevant information.', 'Externalities create spillover costs or benefits outside market prices.', 'Policy can improve or worsen outcomes depending on design, information, and implementation.', 'Compare efficiency, equity, administrative cost, and unintended effects.']}
    ],
    'agri_marketing': [
        {'title': 'Agricultural Markets and Price Formation', 'summary': 'Agricultural prices reflect demand, supply, quality, seasonality, storage, transport, and market information.', 'key_terms': ['market margin', 'price spread', 'seasonality', 'arbitrage'], 'notes': ['Prices can vary by location, grade, time, and transaction terms.', 'A price spread may include transport, handling, storage, losses, and services.', 'Seasonal supply patterns can influence farm-gate prices and storage decisions.', 'Compare like products and terms before interpreting a price difference.']},
        {'title': 'Customer, Product and Channel Strategy', 'summary': 'Agricultural marketing matches product quality, quantity, timing, and channel to customer requirements.', 'key_terms': ['market segment', 'grading', 'branding', 'channel'], 'notes': ['Identify buyer specifications for quality, volume, delivery, traceability, and packaging.', 'Grading and standards reduce uncertainty when measurements are consistent.', 'Channel choice affects reach, cost, control, and payment timing.', 'Value addition should be evaluated against investment, market demand, and compliance needs.']},
        {'title': 'Market Intelligence and Pricing', 'summary': 'Market intelligence supports pricing and selling decisions using reliable, comparable evidence.', 'key_terms': ['market information', 'price analysis', 'negotiation', 'contract'], 'notes': ['Check source date, location, grade, and sample when using price data.', 'Negotiation preparation includes alternatives, costs, quality evidence, and delivery capacity.', 'Contracts specify quantity, quality, price, timing, payment, and dispute arrangements.', 'Use current rules and qualified advice for regulated or high-value agreements.']}
    ],
    'financial_management': [
        {'title': 'Financial Planning and Time Value', 'summary': 'Financial management evaluates funding, investment, cash timing, and returns under uncertainty.', 'key_terms': ['time value of money', 'discount rate', 'capital budgeting', 'cash flow'], 'notes': ['Money available now can earn returns, so timing affects value.', 'Net present value discounts expected cash flows and subtracts initial investment.', 'Use a consistent discount rate and period when comparing alternatives.', 'Cash-flow forecasts should include working capital, taxes where relevant, and realistic timing.']},
        {'title': 'Working Capital and Financing', 'summary': 'Working-capital management coordinates cash, inventory, receivables, payables, and short-term finance.', 'key_terms': ['working capital', 'cash conversion cycle', 'liquidity', 'credit'], 'notes': ['Working capital is current assets minus current liabilities.', 'Seasonality and payment delays can create liquidity pressure even in a profitable enterprise.', 'The cash conversion cycle tracks time between paying for inputs and receiving customer cash.', 'Compare financing cost, repayment timing, collateral, and currency exposure.']},
        {'title': 'Investment Risk and Performance', 'summary': 'Financial decisions balance expected return, risk, affordability, and strategic fit.', 'key_terms': ['risk-return trade-off', 'sensitivity analysis', 'leverage', 'diversification'], 'notes': ['Sensitivity analysis shows which uncertain assumptions drive project results.', 'Leverage can magnify returns and losses and increases fixed obligations.', 'Diversification spreads exposure but cannot eliminate all market or systemic risk.', 'Report assumptions and downside scenarios alongside a central forecast.']}
    ],
    'business_statistics': [
        {'title': 'Data Description and Probability', 'summary': 'Business statistics summarizes data and models uncertainty in enterprise decisions.', 'key_terms': ['sample', 'mean', 'variance', 'probability'], 'notes': ['Distinguish the population of interest from the sample observed.', 'Mean and median measure different aspects of center; outliers affect them differently.', 'Variance and standard deviation describe spread in squared and original units respectively.', 'Probability assumptions should be explicit and checked against the context.']},
        {'title': 'Sampling and Estimation', 'summary': 'Sampling distributions explain how estimates vary from sample to sample.', 'key_terms': ['standard error', 'confidence interval', 'sampling bias', 'margin of error'], 'notes': ['Standard error describes estimator variability, not variation among individuals.', 'A confidence interval reflects a procedure with long-run coverage under assumptions.', 'Selection bias can undermine an estimate even with a large sample.', 'State the population, method, time, and uncertainty when reporting results.']},
        {'title': 'Hypothesis Tests and Regression', 'summary': 'Tests and regression evaluate evidence and relationships while requiring careful interpretation.', 'key_terms': ['null hypothesis', 'p-value', 'correlation', 'regression'], 'notes': ['A p-value is calculated assuming the null model and is not the probability the null is true.', 'Correlation does not by itself establish a causal effect.', 'Regression coefficients depend on variable definitions, units, model specification, and data quality.', 'Statistical significance and practical importance are distinct.']}
    ],
    'value_chain': [
        {'title': 'Mapping the Agrifood Value Chain', 'summary': 'Value-chain analysis maps actors, functions, products, services, and relationships from input to consumer.', 'key_terms': ['value chain', 'actor', 'governance', 'value addition'], 'notes': ['Map product, information, finance, and service flows separately.', 'Identify who performs each function and where costs, margins, and risks occur.', 'Chain governance describes how coordination and standards are set.', 'Include smallholders, workers, women, youth, and informal actors in analysis.']},
        {'title': 'Costs, Upgrading and Coordination', 'summary': 'Value-chain improvement may involve process, product, functional, or market upgrading.', 'key_terms': ['process upgrading', 'product upgrading', 'coordination', 'margin'], 'notes': ['Process upgrading improves efficiency or reduces losses.', 'Product upgrading changes quality, attributes, or processing level to meet demand.', 'Coordination may use contracts, cooperatives, aggregation, or shared information.', 'Assess who pays for upgrading and who captures the resulting value.']},
        {'title': 'Inclusion, Resilience and Governance', 'summary': 'Sustainable chains balance competitiveness with inclusion, resilience, safety, and fair governance.', 'key_terms': ['inclusion', 'traceability', 'resilience', 'standard'], 'notes': ['Standards can improve trust but may impose costly compliance requirements.', 'Traceability links products to relevant production and handling records.', 'Resilience depends on alternatives, information, logistics, and adaptive capacity.', 'Analyze power, risk transfer, environmental impact, and benefit distribution.']}
    ],
    'entrepreneurship': [
        {'title': 'Opportunity Recognition and Customer Testing', 'summary': 'Entrepreneurship turns a validated problem into a value proposition for a defined market.', 'key_terms': ['opportunity', 'customer discovery', 'value proposition', 'assumption'], 'notes': ['Separate a problem worth solving from an idea that merely sounds attractive.', 'Interview users about actual behavior, existing alternatives, and willingness to pay.', 'List assumptions and test the riskiest ones early.', 'Use evidence to refine or reject an opportunity.']},
        {'title': 'Venture Model and Resource Mobilization', 'summary': 'A venture model connects customers, value delivery, revenue, cost, partners, and resources.', 'key_terms': ['business model', 'revenue stream', 'startup cost', 'partnership'], 'notes': ['Estimate startup and operating costs separately.', 'Choose revenue mechanisms that fit customer purchasing patterns and seasonality.', 'Partnerships can provide distribution, finance, technical knowledge, or market access.', 'Keep financial forecasts traceable to operational assumptions.']},
        {'title': 'Pitching, Growth and Responsible Practice', 'summary': 'A venture pitch explains the problem, solution, market, model, evidence, team, and request.', 'key_terms': ['pitch', 'milestone', 'runway', 'responsible growth'], 'notes': ['A credible pitch presents evidence and uncertainty, not only optimistic projections.', 'Milestones should measure progress toward tested assumptions.', 'Protect customer data, worker safety, product quality, and transparent claims.', 'Growth should be paced to preserve service, cash control, and compliance.']}
    ],
    'organizational_hr': [
        {'title': 'People, Motivation and Work Design', 'summary': 'Organizational behavior examines how individual differences and job design influence workplace outcomes.', 'key_terms': ['motivation', 'job design', 'equity', 'wellbeing'], 'notes': ['Motivation includes the direction, intensity, and persistence of effort.', 'Clear roles and manageable workloads support performance and wellbeing.', 'Fairness perceptions influence trust, effort, and retention.', 'Avoid applying one motivational theory as if it fits every person.']},
        {'title': 'Teams, Leadership and Communication', 'summary': 'Teams coordinate interdependent work through roles, communication, leadership, and shared goals.', 'key_terms': ['team norm', 'leadership', 'conflict', 'feedback'], 'notes': ['Team norms shape everyday behavior and should be made explicit for important work.', 'Constructive task disagreement can improve decisions when respectfully managed.', 'Feedback should be timely, specific, and connected to observable behavior.', 'Leadership effectiveness depends on task, people, and context.']},
        {'title': 'Human Resource Systems and Employment Practice', 'summary': 'HR systems recruit, develop, support, evaluate, and retain people within legal and ethical requirements.', 'key_terms': ['recruitment', 'performance management', 'training', 'employee relations'], 'notes': ['Define job requirements before recruitment and use fair selection criteria.', 'Training should address assessed capability needs and be evaluated for application.', 'Performance management combines expectations, support, feedback, and accountability.', 'Employment decisions should comply with current Kenyan law and organizational policy.']}
    ],
    'supply_chain': [
        {'title': 'Supply Chain Flows and Network Design', 'summary': 'Supply-chain management coordinates product, information, and financial flows from suppliers to customers.', 'key_terms': ['procurement', 'inventory', 'lead time', 'distribution'], 'notes': ['Map suppliers, facilities, transport, storage, and customers before redesigning a chain.', 'Lead time and variability affect inventory, freshness, and service levels.', 'Perishable products require temperature, handling, and timing controls.', 'Information visibility helps coordinate orders, stock, and delivery.']},
        {'title': 'Logistics, Inventory and Service', 'summary': 'Logistics balances cost, speed, quality, availability, and risk across movement and storage.', 'key_terms': ['reorder point', 'safety stock', 'cold chain', 'service level'], 'notes': ['A reorder point reflects expected demand during lead time and chosen uncertainty protection.', 'Safety stock protects against variability but ties up cash and can lead to spoilage.', 'Cold-chain management controls time and temperature for sensitive products.', 'Service measures should include on-time delivery, order accuracy, and product condition.']},
        {'title': 'Risk, Procurement and Sustainability', 'summary': 'Resilient supply chains anticipate disruption, manage supplier relationships, and account for sustainability.', 'key_terms': ['supplier risk', 'traceability', 'resilience', 'ethical sourcing'], 'notes': ['Assess supplier concentration, transport interruption, weather, quality, and payment risk.', 'Alternative suppliers and routes can improve resilience but may add cost.', 'Traceability supports safety, recall, and origin verification.', 'Evaluate environmental and labor effects throughout the supply chain.']}
    ],
    'farm_management': [
        {'title': 'Farm Planning and Resource Allocation', 'summary': 'Farm management allocates land, labor, capital, and management across activities over time.', 'key_terms': ['farm plan', 'enterprise budget', 'resource constraint', 'calendar'], 'notes': ['Inventory available land, labor, equipment, finance, and timing constraints.', 'Prepare enterprise budgets using consistent prices, yields, and cost assumptions.', 'Farm calendars coordinate planting, input, labor, harvest, and sales decisions.', 'Compare alternatives using both cash and opportunity costs.']},
        {'title': 'Operations Research for Farm Decisions', 'summary': 'Operations research formulates constrained farm problems to identify feasible and efficient plans.', 'key_terms': ['decision variable', 'objective function', 'constraint', 'linear programming'], 'notes': ['Define decision variables, objective, and constraints in words before writing equations.', 'A linear program assumes linear relationships within the modeled range.', 'Check units, signs, and feasibility of every constraint.', 'A model recommendation must be interpreted in light of omitted realities and uncertainty.']},
        {'title': 'Control, Performance and Risk', 'summary': 'Farm control compares actual performance with plans and adapts decisions to changing conditions.', 'key_terms': ['variance', 'performance indicator', 'risk', 'contingency'], 'notes': ['Compare actual and budgeted yield, price, cost, and timing to identify drivers of variance.', 'Use indicators that relate to decisions, such as cost per unit or yield per area.', 'Plan for weather, pests, price, liquidity, and labor disruptions.', 'Update forecasts when new information changes the decision environment.']}
    ],
    'econometrics': [
        {'title': 'Regression Foundations and Data', 'summary': 'Econometrics uses statistical models to estimate economic relationships from observed data.', 'key_terms': ['dependent variable', 'regressor', 'OLS', 'residual'], 'notes': ['Define outcome and explanatory variables with units and time period.', 'Ordinary least squares estimates coefficients by minimizing squared residuals under its setup.', 'A fitted relationship is not automatically causal.', 'Plot and inspect data for outliers, missingness, and measurement issues.']},
        {'title': 'Inference and Model Assumptions', 'summary': 'Statistical inference depends on model assumptions, sampling design, and uncertainty estimation.', 'key_terms': ['standard error', 'heteroskedasticity', 'multicollinearity', 'confidence interval'], 'notes': ['Standard errors quantify estimator uncertainty under assumptions.', 'Heteroskedasticity concerns nonconstant error variance and may affect inference.', 'Multicollinearity makes separate coefficient effects harder to estimate precisely.', 'Report assumptions and use suitable diagnostics or robust methods where justified.']},
        {'title': 'Causality and Applied Interpretation', 'summary': 'Causal interpretation needs a defensible identification strategy beyond association.', 'key_terms': ['confounding', 'endogeneity', 'instrument', 'treatment effect'], 'notes': ['Confounders influence both treatment and outcome and can bias comparisons.', 'Endogeneity arises when regressors relate to unobserved determinants in the error.', 'Instruments require relevance and a credible exclusion restriction.', 'Explain practical magnitude and limits instead of reporting a coefficient alone.']}
    ],
    'business_law': [
        {'title': 'Legal Framework and Business Responsibility', 'summary': 'Business law establishes rights, obligations, processes, and remedies for commercial activity.', 'key_terms': ['statute', 'contract', 'liability', 'remedy'], 'notes': ['Identify jurisdiction, governing rule, parties, dates, and source documents.', 'Distinguish a legal obligation from a voluntary code or ethical preference.', 'Business operators should keep records that support decisions and transactions.', 'This study material is introductory, not legal advice.']},
        {'title': 'Contracts, Agency and Commercial Terms', 'summary': 'Contracts and agency rules determine how agreements are formed and when actions bind a principal.', 'key_terms': ['offer', 'acceptance', 'consideration', 'authority'], 'notes': ['Analyze offer, acceptance, terms, capacity, consent, legality, breach, and remedy.', 'Actual authority can be express or implied; apparent authority depends on the principal’s representations.', 'Write quantity, quality, delivery, price, payment, and dispute terms clearly.', 'Check current Kenyan statutes and obtain legal advice for real disputes.']},
        {'title': 'Agricultural Contracts and Dispute Prevention', 'summary': 'Agricultural contracts allocate quality, delivery, price, risk, and responsibility across uncertain production cycles.', 'key_terms': ['quality specification', 'force majeure', 'dispute resolution', 'compliance'], 'notes': ['Define objective grading, sampling, rejection, and inspection procedures.', 'State responsibilities for transport, loss, insurance, and delayed delivery.', 'Dispute clauses should identify notice, negotiation, mediation, or other agreed procedures.', 'Do not assume a clause is enforceable without current legal review.']}
    ],
    'agri_trade': [
        {'title': 'Trade Theory and Agricultural Competitiveness', 'summary': 'Trade analysis examines specialization, comparative advantage, market access, and distributional effects.', 'key_terms': ['comparative advantage', 'export', 'import', 'trade balance'], 'notes': ['Comparative advantage depends on relative opportunity cost, not absolute productivity alone.', 'Trade can create gains while also imposing adjustment costs on particular groups.', 'Agricultural competitiveness depends on quality, logistics, standards, finance, and information.', 'Trade data should identify units, value basis, product classification, and reporting period.']},
        {'title': 'Trade Policy, Standards and Market Access', 'summary': 'Tariffs, quotas, standards, and trade agreements shape costs and opportunities across borders.', 'key_terms': ['tariff', 'non-tariff measure', 'SPS standard', 'rules of origin'], 'notes': ['Tariffs tax imports; non-tariff measures include technical and sanitary requirements.', 'Sanitary and phytosanitary controls aim to protect human, animal, or plant health.', 'Exporters must understand destination-market standards, certification, and documentation.', 'Check current trade agreements and official market-access rules before acting.']},
        {'title': 'Cross-Border Operations and Risk', 'summary': 'International agribusiness manages logistics, payments, documentation, currency, and political risks.', 'key_terms': ['Incoterms', 'foreign exchange', 'letter of credit', 'trade finance'], 'notes': ['Incoterms allocate delivery responsibilities, costs, and risk at specified points.', 'Currency movements affect local-currency revenue and imported input costs.', 'Trade finance can reduce payment risk but carries terms, fees, and documentation requirements.', 'Plan shipping, customs, insurance, traceability, and contingency options.']}
    ],
    'project_management': [
        {'title': 'Project Identification and Design', 'summary': 'Project planning defines a problem, beneficiaries, objectives, activities, outputs, outcomes, and assumptions.', 'key_terms': ['problem tree', 'objective', 'output', 'assumption'], 'notes': ['Separate root causes from symptoms before choosing interventions.', 'Objectives describe desired changes; outputs are direct products of activities.', 'A theory of change explains how outputs are expected to contribute to outcomes.', 'Stakeholder analysis reveals interests, influence, and possible unintended effects.']},
        {'title': 'Appraisal, Budget and Implementation', 'summary': 'Appraisal compares alternatives and checks feasibility, affordability, economic value, and risk.', 'key_terms': ['feasibility', 'NPV', 'work plan', 'risk register'], 'notes': ['Appraise technical, financial, economic, social, environmental, and institutional feasibility.', 'Net present value compares discounted benefits and costs under explicit assumptions.', 'Work plans assign activities, owners, dependencies, and timing.', 'A risk register records probability, impact, owner, and response.']},
        {'title': 'Monitoring, Evaluation and Closure', 'summary': 'Monitoring and evaluation track implementation and assess whether outcomes occurred and why.', 'key_terms': ['indicator', 'baseline', 'evaluation', 'lesson learned'], 'notes': ['Indicators need a definition, data source, frequency, and responsible owner.', 'A baseline establishes conditions before or at the start of an intervention.', 'Evaluation asks about relevance, effectiveness, efficiency, impact, and sustainability.', 'Close projects by reconciling resources, documenting results, and transferring responsibilities.']}
    ],
    'credit_finance': [
        {'title': 'Agricultural Finance and Credit Needs', 'summary': 'Agricultural finance matches funding products to seasonal cash flow, assets, production risk, and market timing.', 'key_terms': ['creditworthiness', 'collateral', 'working capital', 'seasonality'], 'notes': ['Farm cash flows often vary by planting, input, harvest, and payment cycle.', 'Credit assessment considers repayment capacity, cash flow, history, purpose, and risk.', 'Collateral can support lending but does not replace a viable repayment plan.', 'Borrowers should understand fees, interest basis, timing, security, and penalties.']},
        {'title': 'Financial Institutions and Products', 'summary': 'Institutions provide savings, payments, credit, insurance, and investment services under different rules and costs.', 'key_terms': ['commercial bank', 'cooperative finance', 'microfinance', 'insurance'], 'notes': ['Compare products by total cost, access, eligibility, repayment schedule, and consumer protection.', 'Savings and payment systems can help manage timing and transaction risk.', 'Cooperatives can aggregate demand or supply finance, subject to governance and capacity.', 'Use regulated institutions and verify terms through official sources.']},
        {'title': 'Credit Risk and Responsible Borrowing', 'summary': 'Credit risk management protects borrowers and lenders through transparent assessment and repayment planning.', 'key_terms': ['default risk', 'debt service', 'credit record', 'responsible lending'], 'notes': ['Debt service should be compared with realistic cash flow under downside scenarios.', 'Borrowing for a productive use still carries price, weather, and operational risks.', 'Maintain records and communicate early if repayment difficulties arise.', 'Responsible lending includes clear terms, fair treatment, and appropriate affordability checks.']}
    ],
    'research_methods': [
        {'title': 'Research Questions and Design', 'summary': 'Research begins with a focused question and a design that can produce credible evidence.', 'key_terms': ['research problem', 'objective', 'methodology', 'sampling frame'], 'notes': ['A research question should be focused, answerable, and relevant.', 'Align objectives, data, methods, and analysis before collecting information.', 'Choose qualitative, quantitative, or mixed methods based on the question.', 'Define the population and sampling frame to clarify coverage and limitations.']},
        {'title': 'Ethics, Instruments and Fieldwork', 'summary': 'Ethical fieldwork protects participants and supports reliable, respectful data collection.', 'key_terms': ['informed consent', 'confidentiality', 'questionnaire', 'reliability'], 'notes': ['Explain purpose, risks, voluntary participation, and withdrawal rights during consent.', 'Collect only necessary personal data and protect it appropriately.', 'Pilot questionnaires to identify ambiguous or leading items.', 'Train enumerators and document field procedures consistently.']},
        {'title': 'Analysis, Reporting and Reproducibility', 'summary': 'Analysis should transparently connect evidence to conclusions and acknowledge uncertainty.', 'key_terms': ['coding', 'descriptive analysis', 'triangulation', 'limitation'], 'notes': ['Clean and document data before analysis while preserving the raw source.', 'Use methods appropriate to measurement level and research design.', 'Triangulation compares multiple evidence sources or methods.', 'Report limitations and avoid claims beyond the study’s design and sample.']}
    ],
    'innovation': [
        {'title': 'Agricultural Innovation Systems', 'summary': 'Innovation emerges through interaction among producers, researchers, firms, public agencies, and support organizations.', 'key_terms': ['innovation system', 'diffusion', 'co-creation', 'adoption'], 'notes': ['Innovation includes organizational, process, market, and institutional changes, not only technology.', 'Adoption depends on perceived value, compatibility, skills, trust, and access.', 'Co-design with intended users can reveal practical constraints and improve fit.', 'Knowledge flows through formal research, extension, peers, markets, and digital channels.']},
        {'title': 'Incubation and Venture Support', 'summary': 'Incubation helps early ventures test assumptions and access mentorship, networks, facilities, and finance.', 'key_terms': ['incubator', 'mentor', 'prototype', 'validation'], 'notes': ['A prototype tests functionality or user response before full-scale launch.', 'Incubation support should match a venture’s stage and bottlenecks.', 'Mentorship can improve networks and decision quality but does not replace customer evidence.', 'Define milestones and exit criteria for support programmes.']},
        {'title': 'Scaling Responsible Innovation', 'summary': 'Scaling expands beneficial innovations while maintaining quality, inclusion, and financial viability.', 'key_terms': ['scaling pathway', 'replication', 'institutionalization', 'impact'], 'notes': ['Scaling out reaches more users; scaling up influences institutions and policy.', 'Assess affordability, service capacity, maintenance, and local adaptation.', 'Monitor who benefits and who faces new costs or exclusions.', 'Impact claims should be supported by credible evidence and comparison.']}
    ],
    'risk_insurance': [
        {'title': 'Agricultural Risk and Exposure', 'summary': 'Agribusiness faces production, price, operational, financial, policy, and climate risks.', 'key_terms': ['hazard', 'exposure', 'vulnerability', 'risk matrix'], 'notes': ['Risk combines an uncertain event with exposure and potential consequences.', 'Map risks by likelihood, severity, time horizon, and affected stakeholders.', 'Differentiate controllable operational risks from external shocks.', 'Historical patterns inform but do not fully determine future risk.']},
        {'title': 'Insurance and Risk Transfer', 'summary': 'Insurance transfers specified financial risks under policy terms, exclusions, limits, and evidence requirements.', 'key_terms': ['premium', 'deductible', 'indemnity', 'index insurance'], 'notes': ['A premium is paid for specified coverage; deductibles and exclusions shape the protection.', 'Indemnity insurance generally relates compensation to documented loss within policy limits.', 'Index insurance pays based on a defined index trigger and may have basis risk.', 'Read definitions, waiting periods, claims rules, and exclusions before purchase.']},
        {'title': 'Mitigation, Diversification and Continuity', 'summary': 'Risk management combines prevention, mitigation, transfer, acceptance, and contingency planning.', 'key_terms': ['mitigation', 'diversification', 'contingency', 'basis risk'], 'notes': ['Diversification reduces dependence on one crop, buyer, route, or source of income.', 'Mitigation reduces probability or impact through practices, infrastructure, or information.', 'Business continuity plans prioritize critical functions and recovery steps.', 'Insurance is one layer and cannot remove operational or livelihood risk entirely.']}
    ],
    'strategy': [
        {'title': 'Strategic Position and Environment', 'summary': 'Strategic management aligns an enterprise’s goals and capabilities with a changing external environment.', 'key_terms': ['mission', 'PESTEL', 'Five Forces', 'SWOT'], 'notes': ['A mission explains purpose and customers; strategic objectives make priorities actionable.', 'PESTEL scans political, economic, social, technological, environmental, and legal conditions.', 'Industry analysis examines rivalry, entry, substitutes, and buyer and supplier power.', 'SWOT should synthesize evidence and lead to choices, not remain a disconnected list.']},
        {'title': 'Competitive Advantage and Choice', 'summary': 'Strategic choices define where to compete and how to deliver distinctive customer value.', 'key_terms': ['competitive advantage', 'value chain', 'strategic fit', 'portfolio'], 'notes': ['Advantage may come from cost, differentiation, focus, relationships, or hard-to-copy capability.', 'Value-chain analysis identifies activities that create value and cost.', 'Evaluate options for suitability, feasibility, and acceptability.', 'Portfolio decisions allocate limited resources among enterprises and opportunities.']},
        {'title': 'Implementation, Governance and Review', 'summary': 'Strategy succeeds through execution, accountable governance, measurable review, and adaptation.', 'key_terms': ['implementation', 'KPI', 'governance', 'scenario'], 'notes': ['Translate strategy into owners, budgets, capabilities, milestones, and communication.', 'KPIs should measure strategic outcomes and leading drivers.', 'Governance clarifies decision rights, oversight, accountability, and ethics.', 'Review assumptions and scenarios as markets, climate, and policy change.']}
    ],
    'food_policy': [
        {'title': 'Food Security and Nutrition', 'summary': 'Food security includes availability, access, utilization, and stability over time.', 'key_terms': ['availability', 'access', 'utilization', 'stability'], 'notes': ['Food availability concerns production, stocks, and trade.', 'Access depends on income, prices, distribution, and social conditions.', 'Utilization includes nutrition, safety, health, and preparation.', 'Stability asks whether the other dimensions persist through shocks and seasons.']},
        {'title': 'Policy Instruments and Institutions', 'summary': 'Food and agricultural policy uses regulations, investment, services, trade, and safety nets to pursue public goals.', 'key_terms': ['policy instrument', 'subsidy', 'public good', 'safety net'], 'notes': ['Policy objectives can include productivity, affordability, nutrition, resilience, and inclusion.', 'Instruments have fiscal cost, administrative requirements, and distributional effects.', 'Public goods such as research and surveillance may be underprovided by private markets.', 'Policy coherence matters across agriculture, health, trade, environment, and finance.']},
        {'title': 'Ethics, Distribution and Evaluation', 'summary': 'Ethical policy analysis considers rights, fairness, evidence, and effects across affected groups.', 'key_terms': ['equity', 'accountability', 'trade-off', 'evaluation'], 'notes': ['Identify groups who gain, lose, pay, or face implementation burdens.', 'Transparent criteria and grievance mechanisms support accountable delivery.', 'Evaluate intended outcomes and unintended effects using credible measures.', 'Recommendations should explain evidence, uncertainty, trade-offs, and safeguards.']}
    ],
    'data_analytics': [
        {'title': 'Agribusiness Information Systems', 'summary': 'Information systems collect, process, store, and communicate data to support enterprise and value-chain decisions.', 'key_terms': ['information system', 'database', 'workflow', 'data quality'], 'notes': ['Define the decision and users before choosing software or collecting data.', 'A database uses structured records and relationships to support retrieval and updates.', 'Data quality includes accuracy, completeness, consistency, timeliness, and relevance.', 'Access controls protect confidential commercial and personal information.']},
        {'title': 'Descriptive Analytics and Visualization', 'summary': 'Descriptive analytics summarizes historical data so patterns, outliers, and operational issues can be examined.', 'key_terms': ['dashboard', 'aggregation', 'outlier', 'visualization'], 'notes': ['Check definitions, missing values, units, and time periods before aggregation.', 'A dashboard should answer a small set of operational questions.', 'Choose chart types that represent comparisons and trends without distortion.', 'Outliers may be errors or meaningful events and need investigation.']},
        {'title': 'Forecasting, Models and Governance', 'summary': 'Analytical models can support prediction and planning but require validation and responsible governance.', 'key_terms': ['forecast', 'model validation', 'bias', 'privacy'], 'notes': ['Separate training and evaluation data when assessing predictive performance.', 'Compare forecasts with simple baselines and track error over time.', 'Models can encode bias or fail when conditions change.', 'Document data provenance, user permissions, retention, and limitations.']}
    ],
    'field_attachment': [
        {'title': 'Placement Readiness and Learning Plan', 'summary': 'Field attachment connects academic concepts to supervised workplace practice and professional development.', 'key_terms': ['learning objective', 'supervisor', 'workplace safety', 'reflection'], 'notes': ['Agree on learning objectives, responsibilities, supervision, and assessment evidence.', 'Follow workplace safety, confidentiality, and professional conduct requirements.', 'Keep an approved log of tasks and learning without disclosing confidential information.', 'Ask for feedback and connect practice to course concepts.']},
        {'title': 'Field Observation and Evidence', 'summary': 'Structured observation helps students understand actual processes, constraints, and stakeholder roles.', 'key_terms': ['observation', 'process map', 'stakeholder', 'evidence'], 'notes': ['Record date, context, activity, and observation rather than unsupported judgments.', 'Map process steps, handoffs, delays, and quality controls.', 'Ask permission before collecting photographs, data, or interviews.', 'Distinguish observed facts from interpretation and recommendations.']},
        {'title': 'Attachment Reporting and Ethics', 'summary': 'A professional attachment report summarizes learning, evidence, reflection, and recommendations responsibly.', 'key_terms': ['report', 'reflection', 'confidentiality', 'recommendation'], 'notes': ['Follow institutional format and supervisor approval requirements.', 'Use anonymized or authorized data and respect confidentiality.', 'Reflect on skills gained, difficulties, and how theory relates to practice.', 'Make feasible recommendations supported by observations and evidence.']}
    ],
    'research_project': [
        {'title': 'Proposal, Question and Literature', 'summary': 'A research proposal makes the problem, question, objectives, evidence base, and method coherent.', 'key_terms': ['proposal', 'research question', 'literature review', 'conceptual framework'], 'notes': ['A defensible problem statement shows what is known, what is uncertain, and why it matters.', 'Objectives should be specific, aligned to the question, and feasible.', 'A literature review synthesizes evidence rather than listing unrelated summaries.', 'A conceptual framework makes proposed relationships and assumptions explicit.']},
        {'title': 'Data, Analysis and Interpretation', 'summary': 'Project analysis applies appropriate methods to collected evidence and interprets results against the question.', 'key_terms': ['dataset', 'method', 'validity', 'interpretation'], 'notes': ['Use the approved method and document departures from the plan.', 'Protect participants and data throughout collection, storage, and analysis.', 'Check assumptions and data quality before choosing a statistical or qualitative technique.', 'Separate results from interpretation and relate findings to prior evidence.']},
        {'title': 'Thesis, Seminar and Defense', 'summary': 'A final research report communicates a reproducible study, limitations, conclusions, and contribution.', 'key_terms': ['thesis', 'citation', 'limitation', 'defense'], 'notes': ['Organize the thesis so each chapter advances the research question.', 'Cite sources consistently and distinguish quotation, paraphrase, and original analysis.', 'State limitations honestly and avoid claims beyond the design.', 'A defense should explain rationale, evidence, choices, contribution, and responses to critique.']}
    ],
    'feasibility': [
        {'title': 'Customer, Market and Technical Feasibility', 'summary': 'A feasibility study tests whether a proposed agrienterprise has a real market and can deliver its offer.', 'key_terms': ['market size', 'customer validation', 'capacity', 'feasibility'], 'notes': ['Estimate addressable demand using traceable evidence and realistic customer reach.', 'Define customer segment, competition, pricing, and route to market.', 'Assess production technology, inputs, labor, location, utilities, and quality requirements.', 'Identify critical assumptions that require field validation.']},
        {'title': 'Financial and Economic Appraisal', 'summary': 'Financial analysis evaluates cash viability, while economic appraisal considers wider costs and benefits.', 'key_terms': ['capital cost', 'operating cost', 'NPV', 'sensitivity'], 'notes': ['Separate capital expenditure, operating costs, revenues, and working capital.', 'Use consistent timing, currency, tax treatment, and discount assumptions.', 'Calculate break-even and cash needs, not only accounting profit.', 'Test downside, central, and upside scenarios for key prices and yields.']},
        {'title': 'Business Plan, Risk and Decision', 'summary': 'A business plan turns feasibility evidence into an actionable, financeable, and risk-aware operating strategy.', 'key_terms': ['business plan', 'milestone', 'risk register', 'go/no-go'], 'notes': ['Link operations, marketing, staffing, finance, and compliance into one realistic plan.', 'Assign owners, milestones, and performance measures.', 'Identify permits, contracts, environmental issues, and contingency actions.', 'A go/no-go recommendation should explain remaining uncertainty and conditions.']}
    ],
    'natural_resource_economics': [
        {'title': 'Resource Scarcity and Allocation', 'summary': 'Natural resource economics studies allocation of land, water, forests, fisheries, and ecosystem services over time.', 'key_terms': ['scarcity', 'property rights', 'common pool', 'opportunity cost'], 'notes': ['Resource decisions involve opportunity costs and often affect future users.', 'Property rights and access rules shape incentives and stewardship.', 'Common-pool resources are difficult to exclude users from and can be depleted.', 'Intergenerational effects require attention to long-run values and uncertainty.']},
        {'title': 'Externalities and Valuation', 'summary': 'Environmental valuation estimates benefits and costs that may not be reflected in market prices.', 'key_terms': ['externality', 'ecosystem service', 'contingent valuation', 'social cost'], 'notes': ['Ecosystem services include provisioning, regulating, cultural, and supporting functions.', 'Market prices can omit pollution, soil loss, biodiversity, and health effects.', 'Valuation methods depend on data, assumptions, and the decision context.', 'Avoid interpreting a monetary estimate as the complete value of nature.']},
        {'title': 'Policy, Conservation and Climate', 'summary': 'Resource policy balances livelihoods, production, conservation, resilience, and equity.', 'key_terms': ['conservation', 'payment for ecosystem services', 'climate adaptation', 'discounting'], 'notes': ['Compare regulation, incentives, collective governance, and information measures.', 'Discount rates influence how future environmental costs and benefits are weighted.', 'Climate adaptation reduces vulnerability while accounting for local knowledge and constraints.', 'Assess who bears conservation costs and who receives benefits.']}
    ],
    'policy_regulation': [
        {'title': 'Agricultural Policy and Governance', 'summary': 'Policy and regulatory frameworks shape agricultural incentives, public services, market access, and accountability.', 'key_terms': ['policy cycle', 'regulation', 'devolution', 'stakeholder'], 'notes': ['The policy cycle includes agenda setting, design, adoption, implementation, and review.', 'Map mandates and coordination across national and county institutions.', 'Regulation should have a clear objective, authority, process, and compliance mechanism.', 'Consult affected producers, firms, consumers, and communities.']},
        {'title': 'Standards, Compliance and Enforcement', 'summary': 'Standards and enforcement protect safety, quality, fair competition, and environmental objectives.', 'key_terms': ['standard', 'licensing', 'inspection', 'compliance cost'], 'notes': ['Distinguish voluntary standards from binding legal requirements.', 'Compliance systems include documentation, inspection, corrective action, and appeal.', 'Unclear or inconsistent enforcement can increase cost and reduce trust.', 'Use current official legal texts and regulator guidance for compliance decisions.']},
        {'title': 'Policy Analysis and Reform', 'summary': 'Policy analysis examines evidence, alternatives, effects, implementation capacity, and reform trade-offs.', 'key_terms': ['impact assessment', 'cost-benefit', 'incidence', 'accountability'], 'notes': ['Define the problem and establish a baseline before recommending intervention.', 'Compare alternatives by effectiveness, cost, feasibility, equity, and risk.', 'Trace effects on farm incentives, consumers, firms, public budgets, and ecosystems.', 'Specify monitoring indicators, review dates, and accountability channels.']}
    ],
    'extension': [
        {'title': 'Extension Principles and Adult Learning', 'summary': 'Agricultural extension facilitates learning, problem solving, and access to knowledge and services.', 'key_terms': ['extension', 'adult learning', 'participatory approach', 'facilitation'], 'notes': ['Start from farmers’ priorities, existing knowledge, and local context.', 'Adults learn through relevant problems, practice, discussion, and reflection.', 'Extension is facilitation and feedback, not one-way transmission alone.', 'Respect local knowledge while checking claims against evidence.']},
        {'title': 'Technology Transfer and Adoption', 'summary': 'Technology transfer supports adaptation and use of innovations that fit user needs and conditions.', 'key_terms': ['technology transfer', 'adoption', 'demonstration', 'diffusion'], 'notes': ['Adoption depends on expected benefit, compatibility, affordability, skills, and risk.', 'Demonstrations should be credible, locally relevant, and transparent about results.', 'Feedback from users can reveal maintenance, access, or design barriers.', 'Monitor adoption and outcomes instead of counting attendance alone.']},
        {'title': 'Communication, Inclusion and Evaluation', 'summary': 'Effective extension reaches diverse groups through trusted channels and evaluates learning and outcomes.', 'key_terms': ['pluralistic extension', 'inclusion', 'feedback', 'impact'], 'notes': ['Use multiple channels and local languages where appropriate.', 'Check whether women, youth, remote producers, and marginalized groups can participate.', 'Feedback loops help adjust advice to changing conditions.', 'Evaluate knowledge, practice, productivity, income, and unintended effects as appropriate.']}
    ],
    'derivatives': [
        {'title': 'Financial Markets and Derivative Contracts', 'summary': 'Derivatives derive value from an underlying asset, rate, index, or event and can transfer or reshape risk.', 'key_terms': ['forward', 'futures', 'option', 'underlying'], 'notes': ['A forward is a customized agreement; exchange-traded futures are standardized and typically margined.', 'An option gives the buyer a right, not an obligation, subject to its contract terms.', 'The underlying may be a commodity price, currency, interest rate, or index.', 'Contract size, expiry, settlement, and counterparty rules matter.']},
        {'title': 'Hedging, Pricing and Basis Risk', 'summary': 'Hedging uses an offsetting position to reduce exposure but does not eliminate every risk.', 'key_terms': ['hedge', 'basis', 'margin', 'price risk'], 'notes': ['A hedge should correspond to an identifiable exposure and time horizon.', 'Basis risk arises when the hedge instrument price does not move exactly with the exposure.', 'Margin requirements can create liquidity needs even when the hedge is economically sound.', 'Hedging can reduce adverse price risk while also limiting favorable outcomes.']},
        {'title': 'Options, Scenarios and Governance', 'summary': 'Options and other instruments require payoff analysis, suitability checks, controls, and clear governance.', 'key_terms': ['call option', 'put option', 'strike price', 'market conduct'], 'notes': ['A call provides the right to buy at a strike; a put provides the right to sell, subject to terms.', 'Payoff diagrams clarify outcomes across possible market prices.', 'Leverage can produce losses greater than expected for some positions and products.', 'Use only regulated, suitable arrangements and understand fees, liquidity, and legal obligations.']}
    ],
    'leadership_governance': [
        {'title': 'Leadership and Organizational Direction', 'summary': 'Leadership aligns people around purpose, supports sound decisions, and builds capacity for collective action.', 'key_terms': ['leadership', 'vision', 'delegation', 'accountability'], 'notes': ['Leadership behaviors should fit the task, team capability, and context.', 'Clear purpose helps coordinate decisions across an enterprise.', 'Delegation requires authority, resources, and accountability to be aligned.', 'Ethical leaders explain decisions and invite responsible challenge.']},
        {'title': 'Governance and Oversight', 'summary': 'Governance structures define decision rights, oversight, transparency, and accountability.', 'key_terms': ['board', 'fiduciary duty', 'transparency', 'conflict of interest'], 'notes': ['Clarify roles of owners, board, management, members, and regulators.', 'Controls protect assets and support reliable reporting and compliance.', 'Conflicts of interest should be disclosed and managed.', 'Transparent records support member trust and informed oversight.']},
        {'title': 'Ethics, Succession and Resilience', 'summary': 'Responsible governance promotes integrity, succession, inclusion, and continuity through change.', 'key_terms': ['ethics', 'succession', 'inclusion', 'resilience'], 'notes': ['Codes of conduct need reporting channels and fair enforcement.', 'Succession planning develops future leaders and preserves institutional knowledge.', 'Inclusive decisions consider voices affected by strategy and operations.', 'Review governance practices after shocks, disputes, and significant change.']}
    ],
    'postharvest': [
        {'title': 'Harvest Quality and Loss Prevention', 'summary': 'Post-harvest management preserves quality and reduces losses from field to market.', 'key_terms': ['maturity index', 'handling', 'loss assessment', 'quality'], 'notes': ['Harvest at suitable maturity for intended market and transport duration.', 'Rough handling causes bruising, breakage, and microbial entry.', 'Measure losses by quantity and quality at specific chain stages.', 'Use clean equipment and safe handling practices.']},
        {'title': 'Storage, Processing and Food Safety', 'summary': 'Storage and processing manage moisture, temperature, contamination, and product transformation.', 'key_terms': ['moisture content', 'cold storage', 'processing', 'food safety'], 'notes': ['Control moisture and temperature according to commodity requirements.', 'Prevent contamination through cleaning, separation, and pest management.', 'Processing can extend shelf life and add value but requires quality and safety controls.', 'Follow current food safety and labeling requirements.']},
        {'title': 'Value Addition and Marketing', 'summary': 'Value addition links product transformation, packaging, standards, and customer willingness to pay.', 'key_terms': ['value addition', 'packaging', 'traceability', 'market specification'], 'notes': ['Identify customer requirements before investing in processing equipment.', 'Packaging protects the product and communicates required information.', 'Traceability supports recall, quality assurance, and buyer confidence.', 'Compare added revenue with processing, energy, packaging, compliance, and capital costs.']}
    ],
    'venture_scaling': [
        {'title': 'Scaling Readiness and Replicable Operations', 'summary': 'Scaling a venture requires a validated model that can deliver consistent value beyond the original market.', 'key_terms': ['scaling readiness', 'replication', 'standard operating procedure', 'capacity'], 'notes': ['Confirm repeat demand, unit economics, reliable supply, and customer retention.', 'Document critical processes before adding locations or partners.', 'Assess leadership, staffing, technology, quality, and working-capital capacity.', 'Pilot new regions or channels before committing large resources.']},
        {'title': 'Franchising and Partnership Models', 'summary': 'Franchising and partnerships distribute a model while assigning rights, responsibilities, standards, and support.', 'key_terms': ['franchise', 'license', 'royalty', 'brand standard'], 'notes': ['Franchise agreements define territory, fees, training, quality, reporting, and termination.', 'Brand standards protect customer expectations but should be practical to implement.', 'Partnerships need aligned incentives, data sharing, dispute processes, and accountability.', 'Review legal and competition obligations before deployment.']},
        {'title': 'Growth Governance and Impact', 'summary': 'Sustainable scaling balances growth, cash, service quality, risk, and social and environmental outcomes.', 'key_terms': ['growth capital', 'governance', 'quality assurance', 'impact measure'], 'notes': ['Growth increases coordination and control demands as well as revenue potential.', 'Monitor quality, complaints, working capital, and staff capacity at each new stage.', 'Protect farmers and customers from unfair risk transfer or misleading claims.', 'Set impact measures and corrective triggers alongside commercial targets.']}
    ]
}


AGRIBUSINESS_COURSES = [
    ('HRD2101', 'Communication Skills', 1, 1, 'communication'),
    ('HRD2102', 'Development Studies and Social Ethics', 1, 1, 'social_ethics'),
    ('SMA2100', 'Mathematics for Agriculture / Calculus', 1, 1, 'agri_calculus'),
    ('SZL2111', 'HIV/AIDS Prevention and Management', 1, 1, 'hiv_prevention'),
    ('EAG2101', 'Introduction to Agribusiness Management', 1, 1, 'agribusiness_intro'),
    ('EAG2102', 'Principles of Microeconomics', 1, 1, 'microeconomics'),
    ('ICS2110', 'Information Technology and Computer Applications', 1, 1, 'information_technology'),
    ('EAG2103', 'Principles of Macroeconomics', 1, 2, 'macroeconomics'),
    ('EAG2104', 'Agricultural Production Economics', 1, 2, 'production_economics'),
    ('EAG2105', 'Business Mathematics for Agribusiness', 1, 2, 'business_mathematics'),
    ('EAG2106', 'Principles of Accounting for Agribusiness', 1, 2, 'accounting'),
    ('EAG2107', 'Introduction to Enterprise Development', 1, 2, 'enterprise_development'),
    ('EAG2108', 'Plant and Animal Production Principles', 1, 2, 'plant_animal'),
    ('EAG2201', 'Intermediate Microeconomics', 2, 1, 'intermediate_micro'),
    ('EAG2202', 'Agricultural Marketing and Price Analysis', 2, 1, 'agri_marketing'),
    ('EAG2203', 'Financial Management in Agribusiness', 2, 1, 'financial_management'),
    ('EAG2204', 'Business Statistics for Agribusiness', 2, 1, 'business_statistics'),
    ('EAG2205', 'Agribusiness Value Chain Analysis', 2, 1, 'value_chain'),
    ('EAG2206', 'Entrepreneurship Skills and New Venture Creation', 2, 1, 'entrepreneurship'),
    ('EAG2207', 'Intermediate Macroeconomics', 2, 2, 'macroeconomics'),
    ('EAG2208', 'Cost and Management Accounting', 2, 2, 'accounting'),
    ('EAG2209', 'Organizational Behaviour and Human Resource Management', 2, 2, 'organizational_hr'),
    ('EAG2210', 'Agricultural Supply Chain and Logistics Management', 2, 2, 'supply_chain'),
    ('EAG2211', 'Farm Management and Operations Research', 2, 2, 'farm_management'),
    ('EAG2212', 'Commercial Farm Field Attachment / Practical Studies', 2, 2, 'field_attachment'),
    ('EAG2301', 'Econometrics for Agribusiness', 3, 1, 'econometrics'),
    ('EAG2302', 'Agribusiness Business Law and Contracts', 3, 1, 'business_law'),
    ('EAG2303', 'Agricultural Trade and International Business', 3, 1, 'agri_trade'),
    ('EAG2304', 'Project Planning, Appraisal, and Management', 3, 1, 'project_management'),
    ('EAG2305', 'Agribusiness Credit and Financial Institutions', 3, 1, 'credit_finance'),
    ('EAG2306', 'Research Methods in Agribusiness', 3, 1, 'research_methods'),
    ('EAG2307', 'Agribusiness Innovation and Incubations', 3, 2, 'innovation'),
    ('EAG2308', 'Risk Management and Insurance in Agribusiness', 3, 2, 'risk_insurance'),
    ('EAG2309', 'Strategic Management in Agribusiness', 3, 2, 'strategy'),
    ('EAG2310', 'Food Security, Policy, and Agribusiness Ethics', 3, 2, 'food_policy'),
    ('EAG2311', 'Agribusiness Information Systems and Data Analytics', 3, 2, 'data_analytics'),
    ('EAG2312', 'Industrial / Field Attachment (8 to 12 weeks)', 3, 2, 'field_attachment'),
    ('EAG2401', 'Research Project I: Proposal Writing and Seminar', 4, 1, 'research_project'),
    ('EAG2402', 'Agri-Enterprise Feasibility Study and Business Plan', 4, 1, 'feasibility'),
    ('EAG2403', 'Natural Resource and Environmental Economics', 4, 1, 'natural_resource_economics'),
    ('EAG2404', 'Agribusiness Policy and Regulatory Frameworks', 4, 1, 'policy_regulation'),
    ('EAG-ELECTIVE-I', 'Elective I: Agricultural Extension and Technology Transfer', 4, 1, 'extension'),
    ('EAG-ELECTIVE-II', 'Elective II: International Trade / Futures and Options', 4, 1, 'derivatives'),
    ('EAG2405', 'Research Project II: Data Analysis and Thesis Defense', 4, 2, 'research_project'),
    ('EAG2406', 'Leadership, Governance, and Ethics in Agribusiness', 4, 2, 'leadership_governance'),
    ('EAG2407', 'Post-Harvest Management and Value Addition Marketing', 4, 2, 'postharvest'),
    ('EAG2408', 'Agribusiness Venture Scaling and Franchise Management', 4, 2, 'venture_scaling'),
]


for code, title, year, semester, topic_pack in AGRIBUSINESS_COURSES:
    topics = AGRIBUSINESS_TOPIC_PACKS[topic_pack]
    discipline = topic_pack.replace('_', ' ')
    register_jkuat_unit(
        code,
        title,
        discipline,
        f'JKUAT Agribusiness programme unit: {title}. Structured study from introductory foundations through applied analysis and professional practice.',
        topics,
        [f"Define and explain a key idea from {topic['title']}." for topic in topics],
        [f"Explain the main principles of {topics[0]['title']} and apply them to an agribusiness example.", f"Interpret evidence or a decision related to {topics[1]['title']}.", f"Evaluate an applied case using {topics[2]['title']} concepts."],
        [f"Analyze an agribusiness case using the methods and evidence in {topic['title']}." for topic in topics] + [f"Develop a justified recommendation that integrates all three topics in {title}."],
        [],
    )
    UNIT_LIBRARY[code]['year'] = year
    UNIT_LIBRARY[code]['semester'] = semester
    UNIT_LIBRARY[code]['programme'] = 'Agribusiness Management'
    UNIT_LIBRARY[code]['level'] = f'Year {year}, Semester {semester}'
    if not UNIT_LIBRARY[code].get('quiz_questions'):
        quiz_questions = []
        all_notes = [note for topic in topics for note in topic['notes']]
        for index, topic in enumerate(topics):
            correct = topic['notes'][0]
            distractors = [note for note in all_notes if note != correct]
            quiz_questions.append({
                'question': f"Which statement best describes {topic['title']}?",
                'choices': [correct] + distractors[:3],
                'answer': correct,
            })
        UNIT_LIBRARY[code]['quiz_questions'] = quiz_questions


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def normalize_unit_code(value):
    return ''.join(character for character in (value or '').upper() if character.isalnum())


def unit_matches_query(query, code, unit):
    raw_query = (query or '').strip().lower()
    compact_query = normalize_unit_code(query)
    return (
        (len(compact_query) >= 3 and compact_query in normalize_unit_code(code))
        or raw_query in unit['title'].lower()
        or raw_query in unit['description'].lower()
        or raw_query in unit.get('university', '').lower()
    )


def find_unit(query=''):
    q = (query or '').strip().lower()
    if not q:
        return UNIT_LIBRARY.get('ICS1101')
    for key, unit in UNIT_LIBRARY.items():
        if normalize_unit_code(key) == normalize_unit_code(query) or unit_matches_query(query, key, unit):
            return unit
    generic = q.replace('-', ' ').title()
    return {
        'code': q.upper()[:8] or 'GEN',
        'title': generic or 'Custom Unit',
        'university': 'Any University',
        'level': 'Undergraduate',
        'description': f'{generic} introduces the core ideas, principles, and methods needed for independent study in this subject area.',
        'objectives': ['Understand the foundations of the subject.', 'Develop practical problem-solving ability.', 'Apply concepts to assignments and exams.'],
        'topics': [
            {'title': 'Foundations', 'summary': 'Build the core foundations and understand the central ideas first.', 'key_terms': ['concept', 'definition', 'principle'], 'notes': ['Begin with key definitions and terminology.', 'Understand how ideas connect to each other before studying details.', 'Create a summary of the main concepts.']},
            {'title': 'Application', 'summary': 'Use the concepts in simple worked examples and case studies.', 'key_terms': ['application', 'example', 'method'], 'notes': ['Practise one example at a time with a clear method.', 'Explain the steps in your own words.', 'Connect the theory to a real or classroom example.']},
            {'title': 'Revision and Assessment', 'summary': 'Prepare for tests by revising what you know and addressing weak points.', 'key_terms': ['revision', 'assessment', 'practice'], 'notes': ['Review short notes regularly rather than waiting for the last minute.', 'Complete past questions under timed conditions.', 'Focus on understanding the method, not just the answer.']}
        ],
        'practice_questions': ['Explain the main idea of this unit in your own words.', 'List three important concepts and why they matter.', 'Solve one practice question and explain each step clearly.'],
        'cat_questions': ['Define the most important concept in this subject.', 'Give an everyday example of how this unit is used.', 'Explain one common mistake students make in this topic.'],
        'exam_questions': ['Write a detailed answer on one major topic from the unit.', 'Solve a full exam-like problem and show all steps.', 'Compare two related concepts and outline their differences.'],
        'video_search': f'{generic} university lecture tutorial study',
        'quiz_questions': [
            {'question': f'What is the best first step when studying {generic}?', 'choices': ['Memorize without understanding', 'Learn the key concepts and definitions', 'Skip examples', 'Avoid revision'], 'answer': 'Learn the key concepts and definitions'},
            {'question': 'Which approach best prepares you for an assessment?', 'choices': ['Practice questions and review mistakes', 'Read only the headings', 'Avoid difficult topics', 'Study only once'], 'answer': 'Practice questions and review mistakes'},
            {'question': 'Why should you explain a concept in your own words?', 'choices': ['It checks whether you understand it', 'It replaces all practice', 'It makes the topic shorter', 'It removes the need for revision'], 'answer': 'It checks whether you understand it'}
        ]
    }


def generate_questions(unit, count=5):
    questions = []
    for topic in unit['topics']:
        for term in topic['key_terms'][:2]:
            questions.append(f"Explain the meaning of '{term}' in the context of {unit['title']}.")
    extra = unit.get('practice_questions', [])
    for q in extra:
        questions.append(q)
    random.shuffle(questions)
    return questions[:count]


def generate_cat(unit):
    cat = unit.get('cat_questions', [])
    if not cat:
        return generate_questions(unit, 3)
    return cat[:3]


def generate_exam(unit):
    exam = unit.get('exam_questions', [])
    if not exam:
        return generate_questions(unit, 4)
    return exam[:4]


def build_video_links(unit):
    title = (unit.get('title') or '').strip()
    code = (unit.get('code') or '').strip()
    if code and normalize_unit_code(title) == normalize_unit_code(code):
        title = next((topic['title'] for topic in unit.get('topics', []) if topic.get('title')), 'university course topic')
    query = f'"{title}" lecture tutorial'
    return [f'https://www.youtube.com/results?search_query={quote_plus(query)}']


def build_external_resources(unit):
    search_terms = quote_plus(f"{unit['code']} {unit['title']}")
    return [
        {
            'name': 'JKUAT Library e-books',
            'description': 'Search JKUAT-listed e-book databases. Some services require an active JKUAT login.',
            'url': 'https://www.jkuat.ac.ke/department/library/?page_id=17437',
        },
        {
            'name': 'ProQuest Ebook Central',
            'description': 'JKUAT-linked scholarly e-book collection; access may depend on institutional authentication.',
            'url': 'https://ebookcentral.proquest.com/lib/jkuat-ebooks/home.action',
        },
        {
            'name': 'JKUAT Institutional Repository',
            'description': 'Search JKUAT research outputs and repository records related to this unit.',
            'url': f'https://ir.jkuat.ac.ke/simple-search?query={search_terms}',
        },
        {
            'name': 'Kuizz JKUAT resources',
            'description': 'Browse JKUAT course-unit past papers and study resources on Kuizz.',
            'url': 'https://www.kuizz.co.ke/questions-by-university/jomo-kenyatta-university-of-agriculture-and-technology',
        },
        {
            'name': 'Stuvia unit search',
            'description': 'Find student-uploaded study documents. Check the author, course version, access terms, and academic integrity rules.',
            'url': f'https://www.stuvia.com/search?s={search_terms}',
        },
        {
            'name': 'Studocu JKUAT resources',
            'description': 'Browse JKUAT course pages and student-contributed materials; availability and access terms vary.',
            'url': 'https://www.studocu.com/row/institution/jomo-kenyatta-university-of-agriculture-and-technology/6215',
        },
    ]


def public_url(endpoint, **values):
    base_url = app.config['PUBLIC_BASE_URL'] or request.url_root.rstrip('/')
    return f"{base_url}{url_for(endpoint, **values)}"


def answer_question(unit, user_question):
    q = (user_question or '').lower()
    if not q:
        return 'Please ask a question about the unit and I will explain it clearly.'

    for topic in unit['topics']:
        for term in topic['key_terms']:
            if term.lower() in q:
                notes = '\n'.join(f'- {n}' for n in topic['notes'])
                return f"Here is the idea for {topic['title']}:\n\n{topic['summary']}\n\n{notes}\n\nStudy tip: explain the concept in your own words and then solve one short problem using it."

    if 'what' in q and 'means' in q or 'define' in q:
        return f"In {unit['title']}, the best way to understand the concept is to start with its definition, then identify where it is used, and finally try a worked example. This builds both memory and confidence."

    if 'difference' in q or 'compare' in q:
        return f"To compare ideas in {unit['title']}, list their definitions, conditions of use, and a worked example of each. Then explain how they are similar and how they differ."

    if 'example' in q or 'show' in q:
        return f"Use the method: state the formula, substitute numbers, simplify, explain each step, and interpret the final answer in plain language."

    return f"The best way to master {unit['title']} is to revise the definitions, learn one topic at a time, and then solve CAT and exam questions without notes."


def score_quiz(unit, submitted_answers, questions=None):
    questions = questions if questions is not None else unit.get('quiz_questions', [])
    if not questions:
        return 0, 0, 0.0, 'F', []

    total = len(questions)
    score = 0
    details = []

    for index, question in enumerate(questions):
        selected = (submitted_answers.get(f'q{index}') or '').strip()
        is_correct = selected == question['answer']
        if is_correct:
            score += 1
        details.append({
            'question': question['question'],
            'selected': selected,
            'correct': question['answer'],
            'is_correct': is_correct,
            'topic_title': question.get('topic_title', ''),
            'stage': question.get('stage', ''),
            'explanation': question.get('explanation', question['answer']),
        })

    percent = round((score / total) * 100, 2) if total else 0.0
    if percent >= 80:
        grade = 'A'
    elif percent >= 70:
        grade = 'B'
    elif percent >= 60:
        grade = 'C'
    elif percent >= 50:
        grade = 'D'
    else:
        grade = 'F'
    return score, total, percent, grade, details


ASSESSMENT_SIZES = {'cat': 10, 'exam': 15}
ASSESSMENT_STAGES = ('Introduction', 'Core Concepts', 'Advanced Application')


def build_assessment_bank(unit):
    stage_questions = {stage: [] for stage in ASSESSMENT_STAGES}
    fact_bank = []
    topics = unit.get('topics', [])
    for topic_index, topic in enumerate(topics):
        if len(topics) <= 1:
            stage_index = 0
        else:
            stage_index = min(2, (topic_index * 3) // len(topics))
        stage = ASSESSMENT_STAGES[stage_index]
        notes = topic.get('notes', [])
        terms = topic.get('key_terms', []) or [topic['title']]

        for note_index, note in enumerate(notes):
            fact = {'text': note, 'topic_index': topic_index, 'topic': topic['title'], 'stage': stage}
            if unit.get('document_source'):
                fact['concept'] = extract_document_concept(note, terms)
            fact_bank.append(fact)
            if unit.get('document_source'):
                continue

            matching_terms = [term for term in terms if term.lower() in note.lower()]
            if matching_terms:
                question = f"Which statement about {matching_terms[0]} is supported by {topic['title']}?"
            else:
                question = f"Which key point belongs to {topic['title']}?"
            stage_questions[stage].append({
                'id': f"note-{topic_index}-{note_index}",
                'question': question,
                'answer': note,
                'topic_title': topic['title'],
                'stage': stage,
                'explanation': note,
            })

        summary = topic.get('summary', '').strip()
        other_summaries = [
            other_topic['summary']
            for other_index, other_topic in enumerate(topics)
            if other_index != topic_index and other_topic.get('summary', '').strip()
        ]
        if summary and other_summaries and not unit.get('document_source'):
            stage_questions[stage].append({
                'id': f'topic-summary-{topic_index}',
                'question': f"Which description best matches {topic['title']}?",
                'answer': summary,
                'choices': [summary] + random.sample(other_summaries, min(3, len(other_summaries))),
                'topic_title': topic['title'],
                'stage': stage,
                'explanation': summary,
            })

    if unit.get('document_source'):
        generated_questions = unit.get('assessment_questions') or []
        if generated_questions:
            for index, item in enumerate(generated_questions):
                stage = item.get('stage')
                if stage not in ASSESSMENT_STAGES:
                    stage = ASSESSMENT_STAGES[min(2, index * 3 // len(generated_questions))]
                stage_questions[stage].append({
                    'id': item.get('id') or f'gemini-{index}',
                    'question': item['question'],
                    'answer': item['answer'],
                    'choices': list(item['choices']),
                    'topic_title': item.get('topic_title') or topics[min(2, index * 3 // len(generated_questions))]['title'],
                    'stage': stage,
                    'explanation': item.get('explanation') or item['answer'],
                })
        else:
            concepts = list(dict.fromkeys(fact['concept'] for fact in fact_bank))
            known_concepts = {concept.casefold() for concept in concepts}
            for topic in topics:
                for term in topic.get('key_terms', []):
                    if len(term) >= 4 and term.casefold() not in known_concepts:
                        concepts.append(term)
                        known_concepts.add(term.casefold())
            question_stems = (
                'Which concept completes this statement: "{statement}"?',
                'Choose the missing concept in this statement: "{statement}"?',
                'Fill in the blank with the correct concept: "{statement}"?',
            )
            for fact_index, fact in enumerate(fact_bank):
                topic = topics[fact['topic_index']]
                concept = fact['concept']
                statement = re.sub(
                    re.escape(concept), '_____', fact['text'], count=1, flags=re.IGNORECASE
                )
                distractors = [item for item in concepts if item.casefold() != concept.casefold()]
                for variant, stem in enumerate(question_stems):
                    offset = (fact_index + variant * 2) % max(1, len(distractors))
                    choices = [concept] + [
                        distractors[(offset + choice_index) % len(distractors)]
                        for choice_index in range(min(3, len(distractors)))
                    ] if distractors else [concept]
                    stage_questions[fact['stage']].append({
                        'id': f"note-{fact['topic_index']}-{fact_index}-{variant}",
                        'question': stem.format(statement=statement),
                        'answer': concept,
                        'choices': choices,
                        'topic_title': topic['title'],
                        'stage': fact['stage'],
                        'explanation': fact['text'],
                    })

    for index, item in enumerate(unit.get('quiz_questions', [])):
        stage_index = min(2, (index * 3) // max(1, len(unit.get('quiz_questions', []))))
        stage = ASSESSMENT_STAGES[stage_index]
        topic = topics[min(stage_index, len(topics) - 1)] if topics else {'title': unit['title']}
        stage_questions[stage].append({
            'id': f'bank-{index}',
            'question': item['question'],
            'choices': list(item['choices']),
            'answer': item['answer'],
            'topic_title': topic['title'],
            'stage': stage,
            'explanation': item['answer'],
        })

    for stage, questions in stage_questions.items():
        for question in questions:
            if 'choices' not in question:
                correct = question['answer']
                distractors = [
                    fact['text'] for fact in fact_bank
                    if fact['text'] != correct and fact['topic'] != question['topic_title']
                ]
                if len(distractors) < 3:
                    distractors.extend(
                        fact['text'] for fact in fact_bank
                        if fact['text'] != correct and fact['text'] not in distractors
                    )
                question['choices'] = [correct] + random.sample(distractors, min(3, len(distractors)))
            random.shuffle(question['choices'])

    return stage_questions


def generate_assessment_questions(unit, assessment_type, used_question_ids=None):
    if assessment_type not in ASSESSMENT_SIZES:
        raise ValueError('Assessment type must be cat or exam.')

    bank = build_assessment_bank(unit)
    used_question_ids = set(used_question_ids or ())
    quotas = (4, 3, 3) if assessment_type == 'cat' else (5, 5, 5)
    selected = []

    for stage, quota in zip(ASSESSMENT_STAGES, quotas):
        questions = bank[stage]
        fresh_questions = [item for item in questions if item['id'] not in used_question_ids]
        stage_selection = random.sample(fresh_questions, min(quota, len(fresh_questions)))
        selected.extend(stage_selection)

    selected_ids = {question['id'] for question in selected}
    remaining = ASSESSMENT_SIZES[assessment_type] - len(selected)
    if remaining:
        unused_questions = [
            question
            for stage in ASSESSMENT_STAGES
            for question in bank[stage]
            if question['id'] not in used_question_ids and question['id'] not in selected_ids
        ]
        extra_questions = random.sample(unused_questions, min(remaining, len(unused_questions)))
        selected.extend(extra_questions)
        selected_ids.update(question['id'] for question in extra_questions)

    remaining = ASSESSMENT_SIZES[assessment_type] - len(selected)
    if remaining:
        unused_questions = [
            question
            for stage in ASSESSMENT_STAGES
            for question in bank[stage]
            if question['id'] not in selected_ids
        ]
        selected.extend(random.sample(unused_questions, min(remaining, len(unused_questions))))

    for index, question in enumerate(selected):
        question = dict(question)
        question['choices'] = list(question['choices'])
        random.shuffle(question['choices'])
        selected[index] = question

    return selected


def extract_document_text(filepath):
    extension = os.path.splitext(filepath)[1].lower()
    if extension == '.txt':
        with open(filepath, 'r', encoding='utf-8-sig', errors='replace') as source_file:
            return source_file.read()
    if extension == '.pdf':
        from pypdf import PdfReader

        reader = PdfReader(filepath)
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    if extension == '.docx':
        from docx import Document as WordDocument

        document = WordDocument(filepath)
        return '\n'.join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    raise ValueError('Text extraction supports PDF, DOCX, and TXT files. Convert legacy DOC files to DOCX first.')


def extract_document_facts(text):
    cleaned_text = re.sub(r'\s+', ' ', text or '').strip()
    sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z0-9])', cleaned_text)
    facts = []
    seen = set()
    for sentence in sentences:
        sentence = sentence.strip(' \t\r\n-•')
        normalized = sentence.lower()
        if len(sentence) >= 35 and normalized not in seen:
            facts.append(sentence[:700])
            seen.add(normalized)
    if not facts and cleaned_text:
        facts = [cleaned_text[:700]]
    return facts[:240]


def extract_document_concept(fact, key_terms):
    predicates = (
        'introduces', 'organizes', 'organises', 'answers', 'describes', 'represents',
        'satisfies', 'contains', 'provides', 'requires', 'supports', 'measures',
        'solves', 'defines', 'shows', 'refers', 'means', 'uses', 'is', 'are',
        'was', 'were', 'has', 'have', 'does', 'do', 'did', 'can', 'could',
        'may', 'might', 'must', 'should', 'will', 'would',
    )
    predicate_pattern = r'\b(?:' + '|'.join(predicates) + r')\b'
    predicate_match = re.search(predicate_pattern, fact, flags=re.IGNORECASE)
    if predicate_match:
        subject = fact[:predicate_match.start()].strip(' \t\r\n,;:')
        subject = re.sub(r'^(?:a|an|the)\s+', '', subject, flags=re.IGNORECASE)
        if subject and len(subject.split()) <= 5:
            return subject.strip(' .!?')

    term_matches = [
        (fact.lower().find(term.lower()), -len(term), term)
        for term in key_terms
        if term and term.lower() in fact.lower()
    ]
    if term_matches:
        return min(term_matches)[2]
    return ' '.join(fact.split()[:2]).strip(' .!?')


def build_local_document_pack(text, purpose):
    facts = extract_document_facts(text)
    if not facts:
        raise ValueError('No readable text was found in this file. Scanned PDFs need OCR before they can be analyzed.')

    sections = []
    section_size = max(1, (len(facts) + 2) // 3)
    section_names = ('Introduction and Foundations', 'Core Concepts and Methods', 'Application and Revision')
    for index, name in enumerate(section_names):
        section_facts = facts[index * section_size:(index + 1) * section_size]
        if section_facts:
            sections.append({'title': name, 'summary': section_facts[0], 'notes': section_facts[:10]})

    note_sections = ['# Generated study notes', 'These notes are organized from the text extracted from your upload. Check them against the original source.']
    for index, section in enumerate(sections, start=1):
        note_sections.extend([f'## {index}. {section["title"]}', section['summary']])
        note_sections.extend(f'- {fact}' for fact in section['notes'][1:])

    answers = []
    if purpose == 'assignment':
        lines = [line.strip(' \t\r\n-•') for line in (text or '').splitlines() if line.strip()]
        prompts = [line for line in lines if '?' in line or re.match(r'^(?:question|q\.?\s*\d+|task\s*\d+|\d+[.)])\s*[:.)-]?\s*\S', line, re.IGNORECASE)]
        prompts = list(dict.fromkeys(prompts))[:30]
        if not prompts and facts:
            prompts = ['Summarize the central ideas and explain how the main concepts relate.']
        stop_words = {'about', 'after', 'again', 'also', 'because', 'between', 'could', 'does', 'from', 'into', 'more', 'that', 'their', 'there', 'these', 'this', 'through', 'using', 'what', 'when', 'where', 'which', 'while', 'with', 'would', 'your', 'explain', 'describe', 'discuss', 'compare', 'outline', 'state', 'identify'}
        for prompt in prompts:
            tokens = {word.lower() for word in re.findall(r'[A-Za-z]{4,}', prompt)} - stop_words
            ranked_facts = sorted(
                facts,
                key=lambda fact: sum(token in fact.lower() for token in tokens),
                reverse=True,
            )
            evidence = ranked_facts[:3]
            if not evidence or not any(sum(token in fact.lower() for token in tokens) for fact in evidence):
                evidence = facts[:3]
            answers.append({
                'question': prompt,
                'answer': 'Draft answer using only evidence found in the uploaded file: ' + ' '.join(evidence),
                'review': 'Verify this draft, add your own reasoning, and cite the original source where required.',
            })

    return {
        'facts': facts,
        'sections': sections,
        'notes': '\n'.join(note_sections),
        'answers': answers,
        'source_word_count': len(re.findall(r'\b\w+\b', text or '')),
    }


def build_document_unit(doc):
    facts = extract_document_facts(doc.extracted_text or '')
    if not facts:
        raise ValueError('This uploaded document has no extracted study text.')
    try:
        assessment_questions = json.loads(getattr(doc, 'generated_questions', None) or '[]')
    except (TypeError, json.JSONDecodeError):
        assessment_questions = []

    topics = []
    stop_words = {'about', 'after', 'also', 'been', 'from', 'have', 'into', 'more', 'only', 'that', 'their', 'there', 'these', 'this', 'those', 'through', 'using', 'what', 'when', 'where', 'which', 'with'}
    for index, title in enumerate(('Introduction and Key Terms', 'Core Ideas and Methods', 'Application and Revision')):
        topic_facts = facts[index::3] or facts
        word_counts = Counter(
            word.lower() for fact in topic_facts for word in re.findall(r'[A-Za-z]{4,}', fact)
            if word.lower() not in stop_words
        )
        terms = [word for word, _ in word_counts.most_common(4)] or [f'concept {index + 1}']
        topics.append({
            'title': title,
            'summary': topic_facts[0],
            'key_terms': terms,
            'notes': topic_facts[:20],
        })

    title = doc.title
    return {
        'code': f'DOC{doc.id}',
        'title': title,
        'document_source': True,
        'university': doc.university,
        'level': 'Uploaded study material',
        'description': f'Assessment generated from {doc.title}, using {doc.source_word_count} words extracted from your uploaded file.',
        'objectives': ['Review the source document in order.', 'Check understanding of key ideas from the upload.', 'Apply the extracted concepts to practice questions.'],
        'topics': topics,
        'practice_questions': [f"Explain this idea in your own words: {topic['summary']}" for topic in topics],
        'cat_questions': [],
        'exam_questions': [],
        'video_search': f'{doc.course} {doc.title} study review',
        'quiz_questions': [],
        'assessment_questions': assessment_questions,
    }


def resolve_assessment_unit(unit_code):
    document_match = re.fullmatch(r'DOC(\d+)', (unit_code or '').strip().upper())
    if document_match:
        doc = Document.query.get_or_404(int(document_match.group(1)))
        if doc.document_type == 'notes' and not getattr(doc, 'generated_questions', None):
            generated_questions = request_gemini_assessment_questions(doc.extracted_text or '', doc.title)
            if generated_questions:
                doc.generated_questions = json.dumps(generated_questions, ensure_ascii=False)
                db.session.commit()
        return build_document_unit(doc)
    return UNIT_LIBRARY.get(unit_code.strip().upper()) or find_unit(unit_code)


def request_gemini_study_pack(text, title, purpose):
    api_key = os.environ.get('GEMINI_API_KEY')
    if not api_key:
        return None

    task = 'Create thorough, well-organized study notes from introduction to advanced application.'
    if purpose == 'assignment':
        task = 'Identify each assignment question and draft a worked answer grounded only in the uploaded material. Clearly mark unsupported parts.'
    prompt = (
        f"{task}\nDocument title: {title}\n"
        'Use your own wording. Do not invent citations or claim unsupported facts. Organize the response with clear headings, '
        'definitions, step-by-step methods where relevant, examples drawn from the source, and a short revision checklist. '
        'For assignments, explain reasoning and label the result as a draft for student review.\n\n'
        f'Uploaded document text:\n{text[:80000]}'
    )
    payload = json.dumps({'contents': [{'parts': [{'text': prompt}]}]}).encode('utf-8')
    api_request = Request(
        'https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent',
        data=payload,
        headers={'Content-Type': 'application/json', 'x-goog-api-key': api_key},
        method='POST',
    )
    try:
        with urlopen(api_request, timeout=35) as response:
            result = json.loads(response.read().decode('utf-8'))
        return result['candidates'][0]['content']['parts'][0]['text'].strip()
    except (HTTPError, URLError, TimeoutError, KeyError, IndexError, ValueError):
        return None


def request_gemini_assessment_questions(text, title):
    api_key = os.environ.get('GEMINI_API_KEY')
    if not api_key:
        return None

    prompt = (
        'Create exactly 30 high-quality multiple-choice questions for a university CAT and exam practice bank. '
        'Use only the source material below. Write 10 questions at each stage: Introduction, Core Concepts, and '
        'Advanced Application. Test understanding, interpretation, reasoning, and application; do not ask students '
        'what a document or notes say, and do not make simple sentence-recognition or fill-in-the-blank questions. '
        'Use short scenarios or calculations when supported by the source. Each question must stand alone and have '
        'exactly four concise, plausible, distinct choices with one unambiguously correct answer. Distractors should '
        'reflect realistic misunderstandings of the same concepts. Explain why the answer is correct using the source. '
        'Do not invent facts or rely on outside knowledge. Return only a JSON object with a "questions" array; each '
        'item must contain question, choices, answer, explanation, and stage fields.\n'
        f'Unit: {title}\nSource material:\n{text[:80000]}'
    )
    payload = json.dumps({
        'contents': [{'parts': [{'text': prompt}]}],
        'generationConfig': {'responseMimeType': 'application/json'},
    }).encode('utf-8')
    api_request = Request(
        'https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent',
        data=payload,
        headers={'Content-Type': 'application/json', 'x-goog-api-key': api_key},
        method='POST',
    )
    try:
        with urlopen(api_request, timeout=35) as response:
            result = json.loads(response.read().decode('utf-8'))
        content = result['candidates'][0]['content']['parts'][0]['text']
        parsed = json.loads(content)
        questions = parsed.get('questions', []) if isinstance(parsed, dict) else parsed
        if not isinstance(questions, list):
            return None

        valid_questions = []
        seen_questions = set()
        forbidden_phrases = ('uploaded document', 'uploaded notes', 'according to the notes', 'the document says')
        stages = ('Introduction', 'Core Concepts', 'Advanced Application')
        for item in questions:
            if not isinstance(item, dict):
                continue
            question = str(item.get('question', '')).strip()
            choices = item.get('choices')
            answer = str(item.get('answer', '')).strip()
            explanation = str(item.get('explanation', '')).strip()
            if not question or not explanation or not isinstance(choices, list):
                continue
            choices = [str(choice).strip() for choice in choices]
            normalized_choices = {choice.casefold() for choice in choices}
            normalized_question = re.sub(r'\W+', ' ', question.casefold()).strip()
            if (len(choices) != 4 or len(normalized_choices) != 4 or not answer
                    or answer.casefold() not in normalized_choices
                    or normalized_question in seen_questions
                    or any(phrase in question.casefold() for phrase in forbidden_phrases)):
                continue
            answer = next(choice for choice in choices if choice.casefold() == answer.casefold())
            stage_index = min(2, len(valid_questions) // 10)
            valid_questions.append({
                'id': f'gemini-{len(valid_questions)}',
                'question': question,
                'choices': choices,
                'answer': answer,
                'explanation': explanation,
                'stage': item.get('stage') if item.get('stage') in stages else stages[stage_index],
            })
            seen_questions.add(normalized_question)
            if len(valid_questions) == 30:
                break

        return valid_questions if len(valid_questions) >= 15 else None
    except (HTTPError, URLError, TimeoutError, KeyError, IndexError, TypeError, ValueError):
        return None


def get_current_user():
    user_id = session.get('user_id')
    if not user_id:
        return None
    return User.query.get(user_id)


@app.context_processor
def inject_auth_navigation():
    return {'current_user': get_current_user()}


INDEX_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>PROF AMON by EDIS | JKUAT Units, University Notes and Online CATs</title>
    <meta name="description" content="Study JKUAT and university units with structured notes, practice questions, CATs, exams, and links to trusted library and repository resources.">
    <meta name="author" content="PETER EDIS">
    <meta name="robots" content="index, follow">
    <link rel="canonical" href="{{ canonical_url }}">
    <meta property="og:title" content="PROF AMON by EDIS | University Study Platform">
    <meta property="og:description" content="Structured university study notes, unit search, and online assessment practice.">
    <meta property="og:type" content="website">
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        body { background: linear-gradient(180deg, #f4f8ff 0%, #edf3ff 100%); }
        .hero { background: linear-gradient(135deg, #0d1b2a, #1b4d8c); border-radius: 22px; color: white; }
        .feature-card { border-radius: 18px; }
        .unit-card { border-radius: 18px; }
    </style>
</head>
<body>
    <nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm sticky-top">
        <div class="container">
            <a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a>
            <div class="d-flex flex-wrap gap-2">
                <a class="btn btn-light" href="/study">Study</a>
                <a class="btn btn-outline-light" href="/#edis-guide">Guide</a>
                <a class="btn btn-outline-light" href="/upload">Upload</a>
                {% if current_user %}
                <a class="btn btn-outline-light" href="/profile">Profile</a>
                <a class="btn btn-light" href="/logout">Log out</a>
                {% else %}
                <a class="btn btn-outline-light" href="/login">Sign in</a>
                <a class="btn btn-light" href="/register">Create account</a>
                {% endif %}
            </div>
        </div>
    </nav>

    <div class="container py-4">
        <div class="hero p-5 mb-4">
            <div class="row align-items-center">
                <div class="col-lg-8">
                    <span class="badge bg-light text-primary mb-3">Interactive tutor system</span>
                    <h1 class="fw-bold display-6">Learn deeper, practise often, and prepare for exams with confidence.</h1>
                    <p class="lead text-light">Choose a university unit, study the full lesson content, ask questions, and practise with CATs and exam-style tasks.</p>
                    <form action="/" method="GET" class="d-flex gap-2 mt-4">
                        <input type="text" name="q" class="form-control form-control-lg" placeholder="Search unit code, topic, or subject..." value="{{ query }}">
                        <button type="submit" class="btn btn-light btn-lg">Search</button>
                    </form>
                </div>
                <div class="col-lg-4">
                    <div class="bg-white text-dark rounded-4 p-4 shadow-sm">
                        <h5 class="fw-bold">Popular units</h5>
                        <div class="d-grid gap-2 mt-3">
                            <a href="/teach/SMA2104" class="btn btn-outline-primary">SMA2104</a>
                            <a href="/teach/ICS1101" class="btn btn-outline-primary">ICS1101</a>
                            <a href="/teach/PHY1201" class="btn btn-outline-primary">PHY1201</a>
                            <a href="/teach/ACC2201" class="btn btn-outline-primary">ACC2201</a>
                            <a href="/teach/SMA2160" class="btn btn-outline-primary">JKUAT SMA2160</a>
                            <a href="/teach/SMA2100" class="btn btn-outline-primary">JKUAT SMA2100</a>
                        </div>
                    </div>
                </div>
            </div>
        </div>

        {% with messages = get_flashed_messages(with_categories=true) %}
            {% if messages %}
                {% for category, message in messages %}
                    <div class="alert alert-{{ category }}">{{ message }}</div>
                {% endfor %}
            {% endif %}
        {% endwith %}

        <div class="row g-4 mb-5">
            <div class="col-md-3">
                <div class="feature-card card h-100 shadow-sm border-0 p-3 bg-white"><div class="card-body"><h5 class="fw-bold text-primary">📘 Deep notes</h5><p class="text-secondary mb-0">Learn in-depth concept explanations and study guides.</p></div></div>
            </div>
            <div class="col-md-3">
                <div class="feature-card card h-100 shadow-sm border-0 p-3 bg-white"><div class="card-body"><h5 class="fw-bold text-primary">🧪 CATs</h5><p class="text-secondary mb-0">Short classroom tests to check understanding quickly.</p></div></div>
            </div>
            <div class="col-md-3">
                <div class="feature-card card h-100 shadow-sm border-0 p-3 bg-white"><div class="card-body"><h5 class="fw-bold text-primary">🧠 Exams</h5><p class="text-secondary mb-0">Exam-style practice questions for revision and confidence.</p></div></div>
            </div>
            <div class="col-md-3">
                <div class="feature-card card h-100 shadow-sm border-0 p-3 bg-white"><div class="card-body"><h5 class="fw-bold text-primary">🎥 Videos</h5><p class="text-secondary mb-0">Open lecture and tutorial video resources for each unit.</p></div></div>
            </div>
        </div>

        {{ guide_panel|safe }}

        <h3 class="fw-bold mb-3">Available study units</h3>
        <div class="row g-4">
            {% for unit in units %}
            <div class="col-md-4">
                <div class="unit-card card h-100 shadow-sm border-0 bg-white p-3">
                    <div class="card-body">
                        <span class="badge bg-primary mb-2">{{ unit.code }}</span>
                        <h5 class="fw-bold">{{ unit.title }}</h5>
                        <p class="text-muted mb-2">{{ unit.university }}</p>
                        <p class="text-secondary">{{ unit.description }}</p>
                    </div>
                    <div class="card-footer bg-white border-0 pt-0">
                        <a href="/teach/{{ unit.code }}" class="btn btn-primary w-100">Open teaching page</a>
                    </div>
                </div>
            </div>
            {% else %}
            <div class="col-12"><div class="alert alert-info">No unit matched your search. Try SMA2160, SMA2100, DIT0106, ICS1101, PHY1201, or ACC2201.</div></div>
            {% endfor %}
        </div>

        <div class="mt-5">
            <h3 class="fw-bold">Uploaded academic resources</h3>
            <div class="row mt-3 g-4">
                {% for doc in documents %}
                <div class="col-md-4">
                    <div class="card h-100 shadow-sm border-0"><div class="card-body"><span class="badge bg-secondary mb-2">{{ doc.university }}</span><h5><a href="/document/{{ doc.id }}" class="text-decoration-none text-dark">{{ doc.title }}</a></h5><p class="text-muted mb-1"><strong>Course:</strong> {{ doc.course }}</p><p class="text-secondary">{{ doc.description[:100] if doc.description else 'No description.' }}...</p></div><div class="card-footer bg-white border-0"><a href="/document/{{ doc.id }}" class="btn btn-outline-primary btn-sm">View</a><a href="/download/{{ doc.filename }}" class="btn btn-success btn-sm">Download</a></div></div>
                </div>
                {% else %}
                <div class="col-12"><div class="alert alert-light border">No uploaded documents yet.</div></div>
                {% endfor %}
            </div>
        </div>
        <footer class="border-top mt-5 py-4 text-center text-secondary">PROF AMON by EDIS is created and developed by <strong>PETER EDIS</strong>.</footer>
    </div>
</body>
</html>
"""

STUDY_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Study Space | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
    <nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm">
        <div class="container">
            <a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a>
            <div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/">Home</a><a class="btn btn-outline-light" href="/#edis-guide">Guide</a><a class="btn btn-outline-light" href="/upload">Upload</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div>
        </div>
    </nav>
    <div class="container py-5">
        <h2 class="fw-bold mb-4">Study Space</h2>
        {{ guide_panel|safe }}
        <form method="GET" action="/study" class="row g-2 mb-4">
            <div class="col-md-10"><input type="text" name="q" class="form-control" value="{{ query }}" placeholder="Search unit code or title..."></div>
            <div class="col-md-2"><button type="submit" class="btn btn-primary w-100">Search</button></div>
        </form>
        <div class="row g-4">
            {% for unit in units %}
            <div class="col-md-4">
                <div class="card shadow-sm h-100 border-0"><div class="card-body"><span class="badge bg-primary mb-2">{{ unit.code }}</span><h5 class="fw-bold">{{ unit.title }}</h5><p class="text-muted">{{ unit.university }}</p><p class="small text-muted">{% if unit.programme %}{{ unit.programme }} · Year {{ unit.year }}, Semester {{ unit.semester }}{% else %}{{ unit.level }}{% endif %}</p><p class="text-secondary">{{ unit.description }}</p></div><div class="card-footer bg-white border-0"><a href="/teach/{{ unit.code }}" class="btn btn-primary w-100">Open lesson</a></div></div>
            </div>
            {% endfor %}
        </div>
    </div>
</body>
</html>
"""

GUIDE_PANEL_HTML = """
<section id="edis-guide" class="mb-5 p-4 bg-white border rounded-3 shadow-sm" aria-labelledby="edis-guide-title">
    <header class="mb-4"><span class="badge bg-primary mb-2">EDIS</span><h2 id="edis-guide-title" class="fw-bold">Your study guide</h2><p class="text-secondary">Check off each step as you set up your account and study workflow.</p></header>
    <div class="d-flex align-items-center gap-3 mb-3"><progress id="guide-progress" value="0" max="6" class="flex-grow-1"></progress><span id="guide-progress-label" aria-live="polite">0 of 6 complete</span><button id="reset-guide" type="button" class="btn btn-outline-secondary btn-sm">Reset</button></div>
    <ol class="list-group list-group-numbered mb-4">
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="account"><strong>Create an account or sign in.</strong><p class="mb-2">An account saves CAT/exam results and remembers assessment questions you have already seen.</p><a href="/register">Create account</a> · <a href="/login">Sign in</a></li>
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="course"><strong>Find your course.</strong><p class="mb-2">Search by unit code, unit title, or university. Agribusiness students can search EAG to browse the curriculum in year and semester order.</p><a href="/study">Browse study units</a></li>
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="lesson"><strong>Study the lesson.</strong><p class="mb-2">Read the learning objectives first, then work through topics from foundations to application. Use key terms and practice prompts to check understanding.</p></li>
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="assessment"><strong>Take a CAT or exam.</strong><p class="mb-2">Start a 10-question online CAT or a 15-question exam. Submit to see your score, corrections, topic, and an explanation. Sign in first to save attempts and continue with fresh questions.</p></li>
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="upload"><strong>Upload notes or an assignment.</strong><p class="mb-2">Upload readable TXT, PDF, or DOCX files. Notes become an organized study pack with generated assessments; assignments get draft answers based on the uploaded text. Review all generated work against your source.</p><a href="/upload">Upload a document</a></li>
        <li class="list-group-item"><input class="form-check-input me-2" type="checkbox" data-guide-step="history"><strong>Review your learning history.</strong><p class="mb-2">Use your profile to see scores and resume the next CAT or exam for each unit.</p><a href="/profile">Open student profile</a></li>
    </ol>
    <p class="small text-muted">Scanned PDFs need OCR before text can be analyzed. Gemini-generated explanations require a server-side GEMINI_API_KEY; without it, the app drafts content from the uploaded text.</p>
</section>
<script>
(() => {
    const storageKey = 'edis-guide-progress-v1';
    const checkboxes = Array.from(document.querySelectorAll('[data-guide-step]'));
    const progress = document.getElementById('guide-progress');
    const label = document.getElementById('guide-progress-label');
    const update = () => {
        const completed = checkboxes.filter((checkbox) => checkbox.checked).length;
        progress.value = completed;
        label.textContent = `${completed} of ${checkboxes.length} complete`;
        localStorage.setItem(storageKey, JSON.stringify(checkboxes.filter((checkbox) => checkbox.checked).map((checkbox) => checkbox.dataset.guideStep)));
    };
    try {
        const saved = new Set(JSON.parse(localStorage.getItem(storageKey) || '[]'));
        checkboxes.forEach((checkbox) => { checkbox.checked = saved.has(checkbox.dataset.guideStep); });
    } catch (error) {
        localStorage.removeItem(storageKey);
    }
    checkboxes.forEach((checkbox) => checkbox.addEventListener('change', update));
    document.getElementById('reset-guide').addEventListener('click', () => {
        checkboxes.forEach((checkbox) => { checkbox.checked = false; });
        update();
    });
    update();
})();
</script>
"""

UNIT_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ unit.title }} | PROF AMON by EDIS</title>
    <meta name="description" content="Study {{ unit.code }} {{ unit.title }} at JKUAT with structured unit notes, objectives, practice questions, CATs, and exam preparation.">
    <meta name="author" content="PETER EDIS">
    <meta name="robots" content="index, follow">
    <link rel="canonical" href="{{ canonical_url }}">
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        .topic-card { border-radius: 18px; }
        .pill { background: #edf3ff; color: #1e3a8a; border-radius: 999px; padding: 6px 12px; margin: 5px; display: inline-block; }
    </style>
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/study">Study</a><a class="btn btn-outline-light" href="/#edis-guide">Guide</a><a class="btn btn-outline-light" href="/upload">Upload</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
<div class="container py-5">
    <div class="card shadow-sm border-0 mb-4"><div class="card-body p-4"><span class="badge bg-primary mb-3">{{ unit.code }}</span><h1 class="fw-bold">{{ unit.title }}</h1><p class="text-muted mb-3">{{ unit.university }} • {{ unit.level }}</p><p>{{ unit.description }}</p>{% for objective in unit.objectives %}<div class="alert alert-light border mb-2">✅ {{ objective }}</div>{% endfor %}</div></div>
    <div class="row g-4 mb-4">
        {% for topic in unit.topics %}
        <div class="col-lg-6"><div class="card topic-card h-100 shadow-sm border-0 p-3"><div class="card-body"><h4 class="fw-bold">{{ topic.title }}</h4><p class="text-secondary">{{ topic.summary }}</p><div class="mb-3">{% for term in topic.key_terms %}<span class="pill">{{ term }}</span>{% endfor %}</div><ul>{% for note in topic.notes %}<li>{{ note }}</li>{% endfor %}</ul></div></div></div>
        {% endfor %}
    </div>

    <div class="row g-4 mb-4">
        <div class="col-lg-4"><div class="card shadow-sm border-0 h-100"><div class="card-body"><h4 class="fw-bold">Practice questions</h4><ol>{% for q in unit.practice_questions %}<li class="mb-2">{{ q }}</li>{% endfor %}</ol></div></div></div>
        <div class="col-lg-4"><div class="card shadow-sm border-0 h-100"><div class="card-body"><h4 class="fw-bold">CAT questions</h4><p class="text-secondary">Practice these written questions, or take the auto-marked online CAT.</p><ol>{% for q in cat_questions %}<li class="mb-2">{{ q }}</li>{% endfor %}<li class="mb-2">Explain how a key concept from these notes applies to a new example.</li></ol><a href="/assessment/{{ unit.code }}/cat" class="btn btn-primary w-100">Start online CAT</a></div></div></div>
        <div class="col-lg-4"><div class="card shadow-sm border-0 h-100"><div class="card-body"><h4 class="fw-bold">Exam questions</h4><p class="text-secondary">Work through these longer questions, then check your understanding with an online exam.</p><ol>{% for q in exam_questions %}<li class="mb-2">{{ q }}</li>{% endfor %}<li class="mb-2">Justify your method and check whether your final result is reasonable.</li></ol><a href="/assessment/{{ unit.code }}/exam" class="btn btn-outline-primary w-100">Start online exam</a></div></div></div>
    </div>

    <div class="card shadow-sm border-0 mb-4"><div class="card-body p-4"><h3 class="fw-bold">Ask the tutor</h3><form method="POST" action="/ask/{{ unit.code }}"><textarea name="question" class="form-control mb-3" rows="4" placeholder="Ask about definitions, examples, difference between concepts, formulas, or study strategy..."></textarea><button type="submit" class="btn btn-primary">Ask question</button></form>{% if answer %}<div class="alert alert-info mt-3 mb-0"><strong>Answer:</strong><br>{{ answer }}</div>{% endif %}</div></div>

    <div class="card shadow-sm border-0 mb-4"><div class="card-body p-4"><h3 class="fw-bold">Video learning</h3><p class="text-secondary">Search YouTube for lecture and tutorial videos about {{ unit.title }}.</p><div class="d-grid gap-2">{% for link in video_links %}<a href="{{ link }}" target="_blank" rel="noopener noreferrer" class="btn btn-outline-primary">Search YouTube for {{ unit.title }} videos</a>{% endfor %}</div></div></div>

    {% if external_resources %}
    <section class="mb-4" aria-labelledby="resource-heading">
        <h3 id="resource-heading" class="fw-bold">JKUAT library and study resources</h3>
        <p class="text-secondary">These links lead to independent external services. Library access may require a JKUAT account; student-contributed material may be incomplete, paid, or governed by separate academic-integrity rules. These notes are original study explanations, not copied course documents.</p>
        <div class="row g-3">
            {% for resource in external_resources %}
            <div class="col-md-6 col-xl-4"><article class="card h-100 border-0 shadow-sm"><div class="card-body"><h4 class="h5 fw-bold">{{ resource.name }}</h4><p class="text-secondary">{{ resource.description }}</p><a href="{{ resource.url }}" target="_blank" rel="noopener noreferrer" class="btn btn-outline-primary">Open resource</a></div></article></div>
            {% endfor %}
        </div>
    </section>
    {% endif %}

    <div class="card shadow-sm border-0"><div class="card-body p-4"><h3 class="fw-bold">Study strategy</h3><ul><li>Read the learning objectives before each topic.</li><li>Study one concept at a time and explain it in your own words.</li><li>Solve CAT and exam questions without looking at your notes.</li><li>Review your mistakes and practise again until the process feels natural.</li><li>Use uploaded notes and video resources to reinforce weak areas.</li></ul><a href="/study" class="btn btn-primary">Browse more units</a></div></div>
    <footer class="border-top mt-5 py-4 text-center text-secondary">PROF AMON by EDIS is created and developed by <strong>PETER EDIS</strong>.</footer>
</div>
</body>
</html>
"""

UPLOAD_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Upload Notes | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
    <nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/">Home</a><a class="btn btn-outline-light" href="/study">Study</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
    <div class="container py-5"><div class="row justify-content-center"><div class="col-lg-7"><div class="card shadow-sm border-0 p-4"><h3 class="fw-bold mb-3">Upload notes or an assignment</h3><p class="text-secondary">Readable text in the file is organized into original study notes. Assignment uploads receive draft answers based only on evidence extracted from your document.</p><form action="/upload" method="POST" enctype="multipart/form-data"><div class="mb-3"><label class="form-label">Document title</label><input type="text" name="title" class="form-control" required></div><div class="mb-3"><label class="form-label">University / institution</label><input type="text" name="university" class="form-control" required></div><div class="mb-3"><label class="form-label">Course / unit code</label><input type="text" name="course" class="form-control" required></div><div class="mb-3"><label class="form-label">What are you uploading?</label><select name="document_type" class="form-select"><option value="notes">Lecture notes / reading</option><option value="assignment">Assignment questions</option></select></div><div class="mb-3"><label class="form-label">Description</label><textarea name="description" class="form-control" rows="3"></textarea></div><div class="mb-3"><label class="form-label">PDF, DOCX, or TXT file</label><input type="file" name="file" class="form-control" accept=".pdf,.docx,.txt" required></div><p class="small text-muted">Scanned PDFs need OCR before their text can be analyzed. Assignment drafts should be reviewed and completed in your own words.</p><button type="submit" class="btn btn-primary w-100">Upload and analyze</button></form></div></div></div></div>
</body>
</html>
"""

DOCUMENT_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ doc.title }} | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/">Home</a><a class="btn btn-outline-light" href="/upload">Upload</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
<main class="container py-5"><div class="card shadow-sm border-0 p-4 mb-4"><span class="badge bg-secondary mb-2">{{ doc.university }} · {{ doc.document_type|title }}</span><h1 class="h2 fw-bold">{{ doc.title }}</h1><p class="text-muted mb-2">Course: {{ doc.course }}</p><p>{{ doc.description if doc.description else 'No description provided.' }}</p><p class="small text-secondary">Analyzed {{ doc.source_word_count }} words from the uploaded file. Generated content is a study draft; check it against the original document.</p><div class="d-flex gap-2 mt-3"><a href="/" class="btn btn-outline-secondary">Back</a><a href="/download/{{ doc.filename }}" class="btn btn-success">Download source file</a></div></div>
{% if doc.generated_notes %}<section class="card shadow-sm border-0 p-4 mb-4"><h2 class="h4 fw-bold">Organized study notes</h2><pre class="text-wrap mb-3" style="white-space: pre-wrap; font: inherit">{{ doc.generated_notes }}</pre><div class="d-flex flex-wrap gap-2"><a href="/assessment/DOC{{ doc.id }}/cat" class="btn btn-primary">Generate 10-question CAT</a><a href="/assessment/DOC{{ doc.id }}/exam" class="btn btn-outline-primary">Generate 15-question exam</a></div></section>{% endif %}
{% if answers_text or answers %}<section class="card shadow-sm border-0 p-4 mb-4"><h2 class="h4 fw-bold">Assignment draft answers</h2>{% if answers_text %}<pre class="text-wrap mb-0" style="white-space: pre-wrap; font: inherit">{{ answers_text }}</pre>{% endif %}{% for answer in answers %}<article class="border-top mt-3 pt-3"><h3 class="h6 fw-bold">{{ answer.question }}</h3><p>{{ answer.answer }}</p><p class="small text-muted">{{ answer.review }}</p></article>{% endfor %}</section>{% endif %}</main>
</body>
</html>
"""

LOGIN_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Login | PROF AMON by EDIS</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
        <div class="container py-5">
                <div class="row justify-content-center">
                        <div class="col-lg-5">
                                <div class="card shadow-sm border-0 p-4">
                                        <h2 class="fw-bold mb-3">Student login</h2>
                                        <form method="POST" action="/login">
                                                <div class="mb-3"><label class="form-label">Email</label><input type="email" name="email" class="form-control" required></div>
                                                <div class="mb-3"><label class="form-label">Password</label><input type="password" name="password" class="form-control" required></div>
                                                <button type="submit" class="btn btn-primary w-100">Log in</button>
                                        </form>
                                        <p class="mt-3 text-center mb-0">Need an account? <a href="/register">Register here</a></p>
                                </div>
                        </div>
                </div>
        </div>
</body>
</html>
"""

REGISTER_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Register | PROF AMON by EDIS</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
        <div class="container py-5">
                <div class="row justify-content-center">
                        <div class="col-lg-6">
                                <div class="card shadow-sm border-0 p-4">
                                        <h2 class="fw-bold mb-3">Create student profile</h2>
                                        <form method="POST" action="/register">
                                                <div class="mb-3"><label class="form-label">Full name</label><input type="text" name="full_name" class="form-control" required></div>
                                                <div class="mb-3"><label class="form-label">Email</label><input type="email" name="email" class="form-control" required></div>
                                                <div class="mb-3"><label class="form-label">Password</label><input type="password" name="password" class="form-control" required></div>
                                                <div class="mb-3"><label class="form-label">Major / area of study</label><input type="text" name="major" class="form-control" placeholder="Computer Science"></div>
                                                <button type="submit" class="btn btn-primary w-100">Register</button>
                                        </form>
                                        <p class="mt-3 text-center mb-0">Already a member? <a href="/login">Login</a></p>
                                </div>
                        </div>
                </div>
        </div>
</body>
</html>
"""

PROFILE_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>{{ user.full_name }} | Student Profile</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex gap-2"><a class="btn btn-light" href="/">Home</a><a class="btn btn-outline-light" href="/study">Study</a></div></div></nav>
<div class="container py-5">
        <div class="card shadow-sm border-0 mb-4"><div class="card-body p-4"><h2 class="fw-bold">{{ user.full_name }}</h2><p class="mb-1"><strong>Email:</strong> {{ user.email }}</p><p class="mb-1"><strong>Major:</strong> {{ user.major }}</p><p class="mb-0"><strong>Profile:</strong> {{ user.profile_title }}</p></div></div>

        <h3 class="fw-bold mb-3">Quiz history</h3>
        {% if results %}
        <div class="table-responsive">
                <table class="table table-striped bg-white shadow-sm">
                        <thead><tr><th>Unit</th><th>Score</th><th>Total</th><th>%</th><th>Grade</th><th>Date</th></tr></thead>
                        <tbody>
                        {% for result in results %}
                                <tr><td>{{ result.unit_code }}</td><td>{{ result.score }}</td><td>{{ result.total }}</td><td>{{ result.percent }}%</td><td>{{ result.grade }}</td><td>{{ result.submitted_at.strftime('%d %b %Y') }}</td></tr>
                        {% endfor %}
                        </tbody>
                </table>
        </div>
        {% else %}
        <div class="alert alert-info">No quiz attempts yet. Try a quiz to begin tracking your progress.</div>
        {% endif %}
        {% if progress_records %}
        <h3 class="fw-bold mt-5 mb-3">Continue your units</h3>
        <div class="table-responsive">
            <table class="table table-striped bg-white shadow-sm">
                <thead><tr><th>Unit</th><th>Assessment</th><th>Attempts</th><th>Next</th><th></th></tr></thead>
                <tbody>
                {% for progress in progress_records %}
                    <tr><td>{{ progress.unit_code }}</td><td>{{ progress.assessment_type|upper }}</td><td>{{ progress.attempt_count }}</td><td>{{ progress.updated_at.strftime('%d %b %Y') if progress.updated_at else 'Just now' }}</td><td><a class="btn btn-sm btn-primary" href="/assessment/{{ progress.unit_code }}/{{ progress.assessment_type }}">Continue</a></td></tr>
                {% endfor %}
                </tbody>
            </table>
        </div>
        {% endif %}
</div>
</body>
</html>
"""

QUIZ_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ unit.title }} Quiz | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/teach/{{ unit.code }}">Unit</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
<div class="container py-5">
    <h2 class="fw-bold mb-3">{{ unit.title }} quiz</h2>
    <p class="text-muted">Answer each question carefully. Your score is saved to your profile.</p>
    <form method="POST" action="/quiz/{{ unit.code }}">
        {% for item in questions %}
        {% set question_index = loop.index0 %}
        <div class="card shadow-sm border-0 mb-3">
            <div class="card-body">
                <p class="fw-bold mb-3">{{ loop.index }}. {{ item.question }}</p>
                {% for choice in item.choices %}
                <div class="form-check mb-2">
                    <input class="form-check-input" type="radio" name="q{{ question_index }}" value="{{ choice }}" id="q{{ question_index }}_{{ loop.index0 }}" required>
                    <label class="form-check-label" for="q{{ question_index }}_{{ loop.index0 }}">{{ choice }}</label>
                </div>
                {% endfor %}
            </div>
        </div>
        {% endfor %}
        <button type="submit" class="btn btn-primary btn-lg">Submit quiz</button>
    </form>
</div>
</body>
</html>
"""

ASSESSMENT_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ assessment_title }} | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/teach/{{ unit.code }}">Unit notes</a><a class="btn btn-outline-light" href="/study">Study</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
<main class="container py-5">
    <h1 class="fw-bold">{{ assessment_title }}</h1>
    <p class="text-secondary">{{ unit.code }}: {{ unit.title }}. Submit your answers to see your score and corrections immediately.</p>
    <p class="text-muted">Attempt {{ attempt_number }} · Questions progress from introductions to core concepts and advanced application.</p>
    {% if result %}
    <div class="alert alert-success"><strong>Score: {{ result.score }}/{{ result.total }} ({{ result.percent }}%) · Grade {{ result.grade }}</strong>{% if saved %}<div>Attempt saved. The next attempt will continue with questions you have not seen yet.</div>{% else %}<div>Log in before submitting to save progress and continue the unit across sessions.</div>{% endif %}</div>
    {% endif %}
    <form method="POST" action="/assessment/{{ unit.code }}/{{ assessment_type }}">
        {% for item in questions %}
        {% set question_index = loop.index0 %}
        <section class="card shadow-sm border-0 mb-3">
            <div class="card-body">
                <h2 class="h5 fw-bold">{{ loop.index }}. {{ item.question }}</h2>
                {% for choice in item.choices %}
                {% set choice_id = 'a' ~ question_index ~ '_' ~ loop.index0 %}
                <div class="form-check mb-2">
                    <input class="form-check-input" type="radio" name="q{{ question_index }}" value="{{ choice }}" id="{{ choice_id }}" {% if not result %}required{% endif %} {% if result and result.details[question_index].selected == choice %}checked{% endif %}>
                    <label class="form-check-label" for="{{ choice_id }}">{{ choice }}</label>
                </div>
                {% endfor %}
                {% if result %}
                {% set detail = result.details[question_index] %}
                <div class="alert {% if detail.is_correct %}alert-success{% else %}alert-warning{% endif %} mt-3 mb-0">
                    {% if detail.is_correct %}Correct.{% else %}Not quite. Correct answer: <strong>{{ detail.correct }}</strong>.{% endif %}
                    <div class="small mt-2"><strong>Topic:</strong> {{ detail.topic_title }}{% if detail.explanation %}<br><strong>From your notes:</strong> {{ detail.explanation }}{% endif %}</div>
                </div>
                {% endif %}
            </div>
        </section>
        {% endfor %}
        <button type="submit" class="btn btn-primary btn-lg">{% if result %}Try again{% else %}Submit and mark{% endif %}</button>
        <a href="/teach/{{ unit.code }}" class="btn btn-outline-secondary btn-lg ms-2">Review lesson notes</a>
    </form>
</main>
</body>
</html>
"""

QUIZ_RESULT_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Quiz Result | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light">
<nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/teach/{{ unit.code }}">Unit</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
<div class="container py-5">
    <div class="card shadow-sm border-0 p-4">
        <h2 class="fw-bold">{{ unit.title }} quiz result</h2>
        <div class="alert alert-success mt-3"><h4 class="mb-0">Score: {{ score }}/{{ total }} ({{ percent }}%) — Grade: {{ grade }}</h4></div>
        <div class="mt-4">
            {% for item in details %}
            <div class="border rounded p-3 mb-2 {% if item.is_correct %}border-success{% else %}border-danger{% endif %}">
                <strong>{{ loop.index }}.</strong> {{ item.question }}<br>
                <small>Your answer: {{ item.selected or 'No answer selected' }}</small><br>
                <small>Correct answer: {{ item.correct }}</small>
            </div>
            {% endfor %}
        </div>
        <div class="d-flex gap-2 mt-4"><a href="/quiz/{{ unit.code }}" class="btn btn-primary">Retake quiz</a><a href="/teach/{{ unit.code }}" class="btn btn-outline-primary">Review unit</a></div>
    </div>
</div>
</body>
</html>
"""

CHAT_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Tutor Chat | PROF AMON by EDIS</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        .chat-box { max-height: 500px; overflow-y: auto; }
        .bubble-user { background: #0d6efd; color: white; }
        .bubble-tutor { background: #e9ecef; color: #212529; }
    </style>
</head>
<body class="bg-light">
    <nav class="navbar navbar-expand-lg navbar-dark bg-primary shadow-sm"><div class="container"><a class="navbar-brand fw-bold" href="/">PROF AMON by EDIS</a><div class="d-flex flex-wrap gap-2"><a class="btn btn-light" href="/">Home</a>{% if current_user %}<a class="btn btn-outline-light" href="/profile">Profile</a><a class="btn btn-light" href="/logout">Log out</a>{% else %}<a class="btn btn-outline-light" href="/login">Sign in</a><a class="btn btn-light" href="/register">Create account</a>{% endif %}</div></div></nav>
    <div class="container py-5">
        <div class="card shadow-sm border-0 p-4">
            <h2 class="fw-bold mb-3">Live AI tutor</h2>
            <div class="chat-box mb-3">
                {% for message in messages %}
                <div class="mb-3">
                    <div class="d-inline-block p-3 rounded {% if message.role == 'user' %}bubble-user{% else %}bubble-tutor{% endif %}">
                        {{ message.content }}
                    </div>
                </div>
                {% endfor %}
            </div>
            <form method="POST" action="/chat">
                <div class="mb-3">
                    <label class="form-label">Ask your tutor</label>
                    <textarea name="message" class="form-control" rows="4" placeholder="Ask a concept, formula, or exam question..."></textarea>
                </div>
                <button type="submit" class="btn btn-primary">Send</button>
            </form>
        </div>
    </div>
</body>
</html>
"""


@app.route('/register', methods=['GET', 'POST'])
def register():
        if request.method == 'POST':
                full_name = request.form.get('full_name', '').strip()
                email = request.form.get('email', '').strip().lower()
                password = request.form.get('password', '')
                major = request.form.get('major', '').strip() or 'Undecided'

                if not full_name or not email or not password:
                        flash('Please fill in all required fields.', 'danger')
                        return redirect(url_for('register'))

                if User.query.filter_by(email=email).first():
                        flash('An account with this email already exists.', 'warning')
                        return redirect(url_for('login'))

                user = User(full_name=full_name, email=email, major=major)
                user.set_password(password)
                db.session.add(user)
                db.session.commit()
                session['user_id'] = user.id
                flash('Registration successful. Welcome to PROF AMON by EDIS.', 'success')
                return redirect(url_for('index'))

        return render_template_string(REGISTER_HTML)


@app.route('/login', methods=['GET', 'POST'])
def login():
        if request.method == 'POST':
                email = request.form.get('email', '').strip().lower()
                password = request.form.get('password', '')
                user = User.query.filter_by(email=email).first()

                if user and user.check_password(password):
                        session['user_id'] = user.id
                        flash('Login successful.', 'success')
                        return redirect(url_for('index'))

                flash('Invalid email or password.', 'danger')
                return redirect(url_for('login'))

        return render_template_string(LOGIN_HTML)


@app.route('/logout')
def logout():
        session.pop('user_id', None)
        flash('You have been logged out.', 'info')
        return redirect(url_for('index'))


@app.route('/profile')
def profile():
        user = get_current_user()
        if not user:
                flash('Please log in to view your profile.', 'warning')
                return redirect(url_for('login'))
        results = QuizResult.query.filter_by(user_id=user.id).order_by(QuizResult.submitted_at.desc()).all()
        progress_records = AssessmentProgress.query.filter_by(user_id=user.id).order_by(AssessmentProgress.updated_at.desc()).all()
        return render_template_string(PROFILE_HTML, user=user, results=results, progress_records=progress_records)


@app.route('/quiz/<unit_code>', methods=['GET', 'POST'])
def quiz(unit_code):
    unit = UNIT_LIBRARY.get(unit_code.strip().upper()) or find_unit(unit_code)
    questions = unit.get('quiz_questions', [])

    if request.method == 'GET':
        return render_template_string(QUIZ_HTML, unit=unit, questions=questions)

    user = get_current_user()
    if not user:
        flash('Please log in to submit a quiz.', 'warning')
        return redirect(url_for('login'))

    submitted = {}
    for index in range(len(questions)):
        submitted[f'q{index}'] = request.form.get(f'q{index}', '')
    score, total, percent, grade, details = score_quiz(unit, submitted)
    result = QuizResult(user_id=user.id, unit_code=unit['code'], score=score, total=total, percent=percent, grade=grade)
    db.session.add(result)
    db.session.commit()
    return render_template_string(QUIZ_RESULT_HTML, unit=unit, score=score, total=total, percent=percent, grade=grade, details=details)


@app.route('/assessment/<unit_code>/<assessment_type>', methods=['GET', 'POST'])
def assessment(unit_code, assessment_type):
    if assessment_type not in {'cat', 'exam'}:
        return 'Assessment not found', 404

    unit = resolve_assessment_unit(unit_code)
    assessment_title = 'Online CAT' if assessment_type == 'cat' else 'Online Exam Practice'
    result = None
    saved = False
    user = get_current_user()
    progress = None
    if user:
        progress = AssessmentProgress.query.filter_by(
            user_id=user.id,
            unit_code=unit['code'],
            assessment_type=assessment_type,
        ).first()

    used_ids = set()
    if progress:
        try:
            used_ids = set(json.loads(progress.used_question_ids or '[]'))
        except (TypeError, json.JSONDecodeError):
            used_ids = set()

    attempt_number = (progress.attempt_count + 1) if progress else 1
    session_key = f"active_assessment:{unit['code']}:{assessment_type}"
    questions = []

    if request.method == 'POST':
        active_ids = session.pop(session_key, [])
        bank = build_assessment_bank(unit)
        bank_by_id = {item['id']: item for items in bank.values() for item in items}
        questions = [dict(bank_by_id[question_id]) for question_id in active_ids if question_id in bank_by_id]
        if len(questions) != ASSESSMENT_SIZES[assessment_type]:
            questions = generate_assessment_questions(unit, assessment_type, used_ids)
        submitted = {f'q{index}': request.form.get(f'q{index}', '') for index in range(len(questions))}
        score, total, percent, grade, details = score_quiz(unit, submitted, questions)
        result = {
            'score': score,
            'total': total,
            'percent': percent,
            'grade': grade,
            'details': details,
        }

        if user:
            if progress is None:
                progress = AssessmentProgress(
                    user_id=user.id,
                    unit_code=unit['code'],
                    assessment_type=assessment_type,
                    used_question_ids='[]',
                    attempt_count=0,
                )
                db.session.add(progress)
            used_ids.update(question['id'] for question in questions)
            progress.used_question_ids = json.dumps(sorted(used_ids))
            progress.attempt_count += 1
            db.session.add(QuizResult(
                user_id=user.id,
                unit_code=f"{unit['code']} {assessment_type.upper()}",
                score=score,
                total=total,
                percent=percent,
                grade=grade,
            ))
            db.session.commit()
            saved = True
            attempt_number = progress.attempt_count
    else:
        questions = generate_assessment_questions(unit, assessment_type, used_ids)
        session[session_key] = [question['id'] for question in questions]

    return render_template_string(
        ASSESSMENT_HTML,
        unit=unit,
        questions=questions,
        assessment_type=assessment_type,
        assessment_title=assessment_title,
        result=result,
        saved=saved,
        attempt_number=attempt_number,
    )


@app.route('/chat', methods=['GET', 'POST'])
def chat():
        user = get_current_user()
        if not user:
                flash('Please log in to use the tutor.', 'warning')
                return redirect(url_for('login'))

        messages = ChatMessage.query.filter_by(user_id=user.id).order_by(ChatMessage.created_at.asc()).all()

        if request.method == 'POST':
                message = request.form.get('message', '').strip()
                if not message:
                        flash('Please type a valid question for the tutor.', 'warning')
                        return redirect(url_for('chat'))

                db.session.add(ChatMessage(user_id=user.id, role='user', content=message))

                unit_hint = None
                for unit in UNIT_LIBRARY.values():
                        if any(keyword.lower() in message.lower() for keyword in unit['title'].lower().split()[:3]):
                                unit_hint = unit
                                break
                if unit_hint is None:
                        unit_hint = UNIT_LIBRARY.get('ICS1101') or find_unit('ICS1101')

                tutor_answer = answer_question(unit_hint, message)
                db.session.add(ChatMessage(user_id=user.id, role='tutor', content=tutor_answer))
                db.session.commit()
                messages = ChatMessage.query.filter_by(user_id=user.id).order_by(ChatMessage.created_at.asc()).all()

        return render_template_string(CHAT_HTML, messages=messages)


@app.route('/')
def index():
    search_query = request.args.get('q', '').strip()
    if search_query:
        documents = Document.query.filter((Document.title.contains(search_query)) | (Document.university.contains(search_query)) | (Document.course.contains(search_query))).all()
    else:
        documents = Document.query.all()

    units = []
    if search_query:
        for key, unit in UNIT_LIBRARY.items():
            if unit_matches_query(search_query, key, unit):
                units.append(unit)
    else:
        units = list(UNIT_LIBRARY.values())[:6]

    return render_template_string(INDEX_HTML, documents=documents, units=units, query=search_query, canonical_url=public_url('index'), guide_panel=GUIDE_PANEL_HTML)


@app.route('/robots.txt')
def robots_txt():
    base_url = app.config['PUBLIC_BASE_URL'] or request.url_root.rstrip('/')
    body = f'User-agent: *\nAllow: /\nSitemap: {base_url}/sitemap.xml\n'
    return Response(body, mimetype='text/plain')


@app.route('/sitemap.xml')
def sitemap_xml():
    namespace = 'http://www.sitemaps.org/schemas/sitemap/0.9'
    root = Element('urlset', {'xmlns': namespace})
    base_url = app.config['PUBLIC_BASE_URL'] or request.url_root.rstrip('/')
    paths = [url_for('index'), url_for('study'), url_for('guide')]
    paths.extend(url_for('teach_unit', unit_code=unit['code']) for unit in UNIT_LIBRARY.values())

    for path in paths:
        entry = SubElement(root, f'{{{namespace}}}url')
        SubElement(entry, f'{{{namespace}}}loc').text = f'{base_url}{path}'

    return Response(tostring(root, encoding='utf-8', xml_declaration=True), mimetype='application/xml')


@app.route('/study')
def study():
    q = request.args.get('q', '').strip()
    units = []
    if q:
        for key, unit in UNIT_LIBRARY.items():
            if unit_matches_query(q, key, unit):
                units.append(unit)
        if not units:
            units = [find_unit(q)]
    else:
        units = list(UNIT_LIBRARY.values())
    units.sort(key=lambda unit: (unit.get('year', 99), unit.get('semester', 99), unit['code']))
    return render_template_string(STUDY_HTML, units=units, query=q, guide_panel=GUIDE_PANEL_HTML)


@app.route('/guide')
def guide():
    return redirect(url_for('index') + '#edis-guide')


@app.route('/teach/<unit_code>')
def teach_unit(unit_code):
    unit_code = unit_code.strip().upper()
    unit = UNIT_LIBRARY.get(unit_code) or find_unit(unit_code)
    video_links = build_video_links(unit)
    external_resources = build_external_resources(unit) if 'jkuat' in unit.get('university', '').lower() else []
    return render_template_string(UNIT_HTML, unit=unit, cat_questions=generate_cat(unit), exam_questions=generate_exam(unit), answer='', video_links=video_links, external_resources=external_resources, canonical_url=public_url('teach_unit', unit_code=unit['code']))


@app.route('/ask/<unit_code>', methods=['POST'])
def ask_question(unit_code):
    unit_code = unit_code.strip().upper()
    unit = UNIT_LIBRARY.get(unit_code) or find_unit(unit_code)
    user_question = request.form.get('question', '').strip()
    answer = answer_question(unit, user_question)
    video_links = build_video_links(unit)
    external_resources = build_external_resources(unit) if 'jkuat' in unit.get('university', '').lower() else []
    return render_template_string(UNIT_HTML, unit=unit, cat_questions=generate_cat(unit), exam_questions=generate_exam(unit), answer=answer, video_links=video_links, external_resources=external_resources, canonical_url=public_url('teach_unit', unit_code=unit['code']))


@app.route('/upload', methods=['GET', 'POST'])
def upload():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        university = request.form.get('university', '').strip()
        course = request.form.get('course', '').strip()
        description = request.form.get('description')
        document_type = request.form.get('document_type', 'notes')
        file = request.files.get('file')

        if document_type not in {'notes', 'assignment'}:
            flash('Choose either lecture notes or assignment questions.', 'danger')
            return redirect(request.url)

        if not file or file.filename == '':
            flash('No file selected.', 'danger')
            return redirect(request.url)

        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            base, ext = os.path.splitext(filename)
            counter = 1
            unique_filename = filename
            while os.path.exists(os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)):
                unique_filename = f'{base}_{counter}{ext}'
                counter += 1
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            file.save(filepath)
            try:
                extracted_text = extract_document_text(filepath)
                if len(extracted_text) > 250000:
                    extracted_text = extracted_text[:250000]
                pack = build_local_document_pack(extracted_text, document_type)
            except (ImportError, ValueError, OSError) as error:
                flash(str(error) or 'Could not read the uploaded document.', 'danger')
                return redirect(request.url)

            generated_text = request_gemini_study_pack(extracted_text, title, document_type)
            generated_notes = (generated_text if document_type == 'notes' and generated_text else pack['notes'])
            generated_answers = None
            if document_type == 'assignment':
                generated_answers = generated_text or json.dumps(pack['answers'], ensure_ascii=False)

            doc = Document(
                title=title,
                university=university,
                course=course,
                filename=unique_filename,
                description=description,
                document_type=document_type,
                extracted_text=extracted_text,
                generated_notes=generated_notes,
                generated_answers=generated_answers,
                source_word_count=pack['source_word_count'],
            )
            db.session.add(doc)
            db.session.commit()
            flash('Document analyzed. Review the generated study material before using it.', 'success')
            return redirect(url_for('view_document', doc_id=doc.id))

        flash('Invalid file type. Allowed: PDF, TXT, DOC, DOCX.', 'danger')

    return render_template_string(UPLOAD_HTML)


@app.route('/document/<int:doc_id>')
def view_document(doc_id):
    doc = Document.query.get_or_404(doc_id)
    answers = []
    answers_text = None
    if doc.generated_answers:
        try:
            answers = json.loads(doc.generated_answers)
        except json.JSONDecodeError:
            answers_text = doc.generated_answers
    return render_template_string(DOCUMENT_HTML, doc=doc, answers=answers, answers_text=answers_text)


@app.route('/download/<filename>')
def download_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)


class Document(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(150), nullable=False)
    university = db.Column(db.String(100), nullable=False)
    course = db.Column(db.String(100), nullable=False)
    filename = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text, nullable=True)
    document_type = db.Column(db.String(30), nullable=False, default='notes')
    extracted_text = db.Column(db.Text, nullable=True)
    generated_notes = db.Column(db.Text, nullable=True)
    generated_answers = db.Column(db.Text, nullable=True)
    generated_questions = db.Column(db.Text, nullable=True)
    source_word_count = db.Column(db.Integer, nullable=False, default=0)

    def __repr__(self):
        return f'<Document {self.title}>'


with app.app_context():
    db.create_all()
    existing_columns = {column['name'] for column in inspect(db.engine).get_columns('document')}
    new_columns = {
        'document_type': "VARCHAR(30) NOT NULL DEFAULT 'notes'",
        'extracted_text': 'TEXT',
        'generated_notes': 'TEXT',
        'generated_answers': 'TEXT',
        'generated_questions': 'TEXT',
        'source_word_count': 'INTEGER NOT NULL DEFAULT 0',
    }
    with db.engine.begin() as connection:
        for column_name, column_definition in new_columns.items():
            if column_name not in existing_columns:
                connection.exec_driver_sql(f'ALTER TABLE document ADD COLUMN {column_name} {column_definition}')


os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(
        debug=os.environ.get('APP_ENV') != 'production' and os.environ.get('FLASK_DEBUG', '').lower() == 'true',
        host='0.0.0.0',
        port=int(os.environ.get('PORT', '5000')),
    )
